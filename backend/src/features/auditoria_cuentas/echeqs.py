"""Cheques entregados (e-cheqs emitidos, endosos y cheques en papel) — 035.

El banco debita un cheque el día que se cobra, pero el pago real al proveedor es el día que se le entrega. Para un
documento en dólares pagado con cheques en pesos, la fecha que fija el valor del dólar es la de entrega (la
"fecha de emisión" del e-cheq). Estas funciones leen los comprobantes (PDF del home banking) y la planilla
"Pagos Cheques Diferidos.xlsx" y arman filas listas para cargar y cruzar con los débitos de Galicia.

Funciones puras: `leer_pdf`, `leer_planilla`, `fusionar`, `cruzar_con_galicia` (no tocan la base).
"""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from pathlib import Path

VENTANA_DIAS = 10  # un cheque se debita pocos días después de su fecha de pago


def _num(s: str) -> float:
    s = s.replace("$", "").strip()
    return float(s.replace(".", "").replace(",", ".")) if "," in s else float(s)


def _fecha(s: str) -> str:
    d, m, a = s.split("/")
    a = int(a)
    a += 2000 if a < 100 else 0
    return f"{a:04d}-{int(m):02d}-{int(d):02d}"


def numero_normalizado(valor) -> str | None:
    """El número del cheque sin ceros a la izquierda ni espacios: '00000065' y '65' son el mismo cheque."""
    solo = re.sub(r"\D", "", str(valor or ""))
    return solo.lstrip("0") or None if solo else None


def leer_pdf(texto: str, nombre_archivo: str = "") -> dict:
    """Interpreta el texto de un comprobante. Devuelve {} si el formato no se reconoce (PDF escaneado, por ejemplo)."""
    t = texto or ""
    r: dict = {}
    if "Detalle de cheque electr" in t:
        m = re.search(r"Fecha de pago Fecha de [Ee]misi[oó]n\s+(\d\d/\d\d/\d\d+)\s+(\d\d/\d\d/\d\d+)", t)
        if m:
            r["fechaPago"], r["fechaEntrega"] = _fecha(m.group(1)), _fecha(m.group(2))
        m = re.search(r"Importe N[°º] de [Cc]heque\s*\n\s*\$?\s*([\d.,]+)\s+(\d+)", t)
        if m:
            r["importe"], r["numero"] = _num(m.group(1)), m.group(2)
        m = re.search(r"Estado\s*\n\s*(.+)", t)
        r["estado"] = m.group(1).strip() if m else None
        m = re.search(r"A la orden\s+(.+)", t)
        r["descripcion"] = m.group(1).strip() if m else None
        m = re.search(r"Emitido a CUIT/CUIL/CDI\s*\n(.+?)\s+(\d{11})", t)
        if m:
            r["beneficiario"], r["cuitBeneficiario"] = m.group(1).strip(), m.group(2)
        m = re.search(r"Raz[oó]n social CUIT/CUIL/CDI\s*\n(.+?)\s+(\d{11})", t)
        librador = m.group(2) if m else None
        # un cheque que figura "emitido a" la empresa y librado por un tercero es un cheque recibido y endosado
        r["tipo"] = "endoso" if (librador and librador != "30712114602") or "Endoso" in (r.get("estado") or "") else "emitido"
        if r["tipo"] == "endoso":
            # la fecha de emisión es la del tercero que libró el cheque, no la de nuestra entrega al proveedor
            r["fechaEmisionTercero"] = r.pop("fechaEntrega", None)
    elif "Consulta de Operaciones" in t:
        for k, rx in {"beneficiario": r"Emitido a\s+(.+)", "cuitBeneficiario": r"Documento\s+([\d-]+)", "descripcion": r"Referencia\s+(.+)"}.items():
            m = re.search(rx, t)
            if m:
                r[k] = m.group(1).strip()
        m = re.search(r"Importe Total\s+\$\s*([\d.,]+)", t)
        if m:
            r["importe"] = _num(m.group(1))
        m = re.search(r"Fecha de Pago\s+(\d\d/\d\d/\d{4})", t)
        if m:
            r["fechaPago"] = _fecha(m.group(1))
        m = re.search(r"(\d{8}) (\d{6}) [ap]\.?m", t)
        if m:
            r["fechaEntrega"] = f"{m.group(1)[:4]}-{m.group(1)[4:6]}-{m.group(1)[6:8]}"
        m = re.search(r"Estado\s+(\w+)", t)
        r["estado"] = m.group(1) if m else None
        m = re.search(r"[Cc]heq\s+(\d+)", nombre_archivo)
        r["numero"] = m.group(1) if m else None
        r["tipo"] = "emitido"
    elif "Datos de la operaci" in t and "Datos del cheque" in t:
        m = re.search(r"Fecha de ejecuci[oó]n Tipo de operaci[oó]n\s*\n\s*(\d\d/\d\d/\d{4})\s+(\w+)", t)
        if m:
            r["fechaEntrega"], operacion = _fecha(m.group(1)), m.group(2)
            r["estado"] = operacion
            r["tipo"] = "endoso" if operacion == "Endoso" else "emitido"
        m = re.search(r"N[uú]mero de cheque(?: Cl[aá]usula)?\s*\n\s*(\d+)", t)
        r["numero"] = m.group(1) if m else None
        m = re.search(r"Fecha de pago Monto\s*\n\s*(\d\d/\d\d/\d{4})\s+\$\s*([\d.,]+)", t)
        if m:
            r["fechaPago"], r["importe"] = _fecha(m.group(1)), _num(m.group(2))
        m = re.search(r"Raz[oó]n social\s*\n\s*(.+?)\s*\n\s*CUIT/CUIL\s*\n\s*([\d-]+)", t.split("Datos del beneficiario")[-1])
        if m:
            r["beneficiario"], r["cuitBeneficiario"] = m.group(1).strip(), m.group(2)
        m = re.search(r"Motivo de la emisi[oó]n\s*\n\s*(.+)", t)
        r["descripcion"] = m.group(1).strip() if m else None
    return r


def completar_fecha_con_el_archivo(fila: dict, mtime: float | None) -> dict:
    """Si el comprobante no trae la fecha de entrega, se usa la del archivo y se deja marcada como estimada."""
    if not fila.get("fechaEntrega") and mtime:
        fila["fechaEntrega"] = datetime.fromtimestamp(mtime).date().isoformat()
        fila["fechaEstimada"] = True
    return fila


def leer_carpeta(raiz: Path) -> tuple[list[dict], list[str]]:
    """Todos los PDF de la carpeta: (filas leídas, archivos que no se pudieron leer por ser imágenes escaneadas)."""
    import pdfplumber

    filas, ilegibles = [], []
    for p in sorted(raiz.rglob("*.pdf")):
        try:
            with pdfplumber.open(p) as pdf:
                texto = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
            r = leer_pdf(texto, p.name)
        except Exception:
            r = {}
        if not r.get("importe"):
            ilegibles.append(str(p.relative_to(raiz)))
            continue
        r.update(archivo=str(p.relative_to(raiz)), carpeta=p.parent.name, fuente="pdf")
        completar_fecha_con_el_archivo(r, p.stat().st_mtime)
        filas.append(r)
    return filas, ilegibles


def leer_planilla(ruta: Path) -> list[dict]:
    """Filas de las solapas 'E - cheqs' y 'Fisicos' de Pagos Cheques Diferidos.xlsx."""
    import openpyxl

    wb = openpyxl.load_workbook(ruta, data_only=True)
    filas = []
    for hoja, tipo in (("E - cheqs", "emitido"), ("Fisicos", "papel")):
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        cab = [str(c or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        for v in ws.iter_rows(min_row=2, values_only=True):
            d = dict(zip(cab, v))
            fecha = d.get("Fecha")
            total = d.get("Total a pagar")
            numero = d.get("Cheque numero")
            if not isinstance(fecha, (date, datetime)) or not isinstance(total, (int, float)) or not numero_normalizado(numero):
                continue
            venc = d.get("Vencimiento")
            usd, tc = d.get("Importe en u$s"), d.get("TC")
            filas.append({
                "tipo": tipo, "numero": str(numero), "importe": round(float(total), 2),
                "fechaEntrega": fecha.date().isoformat() if isinstance(fecha, datetime) else fecha.isoformat(),
                "fechaPago": venc.date().isoformat() if isinstance(venc, datetime) else None,
                "beneficiario": d.get("Proveedor"), "descripcion": d.get("Documento"),
                "documento": d.get("Documento"), "estado": d.get("Estado"),
                "importeUsd": float(usd) if isinstance(usd, (int, float)) and (tc or 0) > 1 else None,
                "tcEmision": float(tc) if isinstance(tc, (int, float)) and tc > 1 else None, "fuente": "excel",
            })
    return filas


def fusionar(pdfs: list[dict], planilla: list[dict]) -> list[dict]:
    """Un solo registro por cheque: el comprobante manda y la planilla aporta documento, dólares y tipo de cambio."""
    por_numero: dict[tuple[str | None, str], list[dict]] = {}
    for x in planilla:
        por_numero.setdefault((x["tipo"] if x["tipo"] == "papel" else "emitido", numero_normalizado(x["numero"])), []).append(x)
    usados: set[int] = set()
    resultado = []
    for p in pdfs:
        n = numero_normalizado(p.get("numero"))
        cand = [x for x in por_numero.get(("emitido", n), []) if abs(x["importe"] - p["importe"]) < 1.0 and id(x) not in usados]
        r = dict(p)
        if cand:
            x = cand[0]
            usados.add(id(x))
            for k in ("documento", "importeUsd", "tcEmision"):
                r[k] = x.get(k)
            r["fuente"] = "pdf+excel"
        resultado.append(r)
    for x in planilla:
        if id(x) not in usados:
            resultado.append(dict(x))
    return resultado


def cruzar_con_galicia(cheques: list[dict], movimientos: list[dict]) -> dict[int, dict]:
    """Índice del cheque (posición en `cheques`) -> movimiento de Galicia que lo debitó.

    `movimientos`: id, fecha (date), descripcion, debito. Se cruza por importe exacto (±$0,01), una ventana de fechas
    alrededor de la fecha de pago y el número en la descripción o, si no hay otro candidato, por ser el único."""
    usados: set[int] = set()
    cruces: dict[int, dict] = {}
    for i, c in enumerate(cheques):
        if c.get("tipo") == "endoso" or not c.get("fechaPago"):
            continue
        pago = date.fromisoformat(c["fechaPago"])
        cand = [m for m in movimientos if m["id"] not in usados and m["debito"] is not None and abs(float(m["debito"]) - c["importe"]) <= 0.01
                and abs((m["fecha"] - pago).days) <= VENTANA_DIAS]
        n = numero_normalizado(c.get("numero"))
        con_numero = [m for m in cand if n and re.search(rf"(?<!\d)0*{re.escape(n)}(?!\d)", re.sub(r"\s+", " ", m["descripcion"] or ""))]
        elegido = (con_numero or (cand if len(cand) == 1 else []))[:1]
        if elegido:
            m = elegido[0]
            usados.add(m["id"])
            cruces[i] = m
    return cruces
