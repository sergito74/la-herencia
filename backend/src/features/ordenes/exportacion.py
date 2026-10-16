"""Exportaciones a planilla (.xlsx) de Órdenes de Trabajo (FR-019)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.features.ordenes import repository, resultado

_PESOS = '"$" #,##0.00;[Red]-"$" #,##0.00'
_CANT = "#,##0.00"
_FECHA = "dd/mm/yyyy"


def _hoja(wb: Workbook, titulo: str, columnas: list[tuple[str, float]]):
    ws = wb.active
    ws.title = titulo
    ws.append([c for c, _ in columnas])
    for i, (_, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3D2B")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    return ws


def _cerrar(ws, formatos: dict[int, str]) -> None:
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for col, fmt in formatos.items():
            fila[col].number_format = fmt
    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _bytes(wb: Workbook) -> bytes:
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


def ordenes_xlsx(**filtros) -> bytes:
    filtros.setdefault("page", 1)
    filtros.setdefault("page_size", 5000)
    items = repository.listar_ordenes(**filtros)["items"]
    wb = Workbook()
    ws = _hoja(wb, "Órdenes de trabajo", [("N°", 8), ("Fecha", 12), ("Labor", 22), ("Contratista", 26), ("Estado", 14)])
    for o in items:
        ws.append([o["idOrden"], o["fechaPedido"], o["tipoLabor"], o["contratista"], o["estado"]])
    _cerrar(ws, {1: _FECHA})
    return _bytes(wb)


def _fecha(v):
    """Fecha sin hora (las celdas de fecha no deben mostrar la hora)."""
    from datetime import date

    return date.fromisoformat(str(v)[:10]) if v else None


def _ha(v: float) -> str:
    return f"{v:,.1f}".rstrip("0").rstrip(".").replace(",", "§").replace(".", ",").replace("§", ".")


def formulario_retiro_xlsx(id_orden: int) -> bytes:
    """Formulario de Retiro con el formato de la planilla del campo (Retiro_151): hoja A4
    apaisada, encabezado con datos de la orden, Ing. Activo · Marca Comercial · Cantidad ·
    Unidad ordenado por principio activo, y firmas de entrega y retiro."""
    from openpyxl.styles import Border, Side

    orden = repository.obtener_orden(id_orden)
    if orden is None or orden["formularioRetiro"] is None:
        raise ValueError([f"La orden {id_orden} no tiene Formulario de Retiro."])
    form = orden["formularioRetiro"]
    wb = Workbook()
    ws = wb.active
    ws.title = f"Retiro {form['idFormularioRetiro']}"
    for col, ancho in zip("ABCDEF", (5, 48, 30, 13, 10, 11)):
        ws.column_dimensions[col].width = ancho
    fino = Side(style="thin", color="808080")
    grueso = Side(style="medium", color="000000")

    ws.append(["La Herencia", None, "RETIRO DE AGROQUÍMICOS", None, "Formulario N°", form["idFormularioRetiro"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws["C1"].font = Font(bold=True, size=14)
    ws["E1"].font = Font(bold=True, color="666666")
    ws["F1"].font = Font(bold=True, size=14)
    ws.append(["Producción agrícola · Depósito de insumos", None, None, None, "Emitido", _fecha(form["fechaEmision"])])
    ws["F2"].number_format = _FECHA
    ws["A2"].font = Font(italic=True, color="666666", size=9)
    for c in range(1, 7):
        ws.cell(row=2, column=c).border = Border(bottom=grueso)
    ws.append([])
    ws.append(["Orden N°", None, "Fecha", None, "Labor", None])
    ws.append([orden["idOrden"], None, _fecha(orden["fechaPedido"]), None, orden.get("tipoLabor"), None])
    ws.append(["Contratista", None, orden.get("contratista") or "Maquinaria propia"])
    destinos: dict[str, dict[str, float]] = {}
    for r in orden["insumos"]:
        for d in r["distribuciones"]:
            if d["aplicar"] and d["cantidadAsignada"] > 0:
                destinos.setdefault(" ".join(x for x in (d["cultivo"], d["campania"]) if x), {})[d["lote"] or str(d["idLote"])] = d["superficie"]
    ws.append(["Cultivos y lotes", None, "; ".join(f"{c}: {', '.join(l)} ({_ha(sum(l.values()))} ha)" for c, l in destinos.items()) or "—"])
    for fila in (4, 6, 7):
        ws.cell(row=fila, column=1).font = Font(bold=True, color="666666", size=9)
    for col in (3, 5):
        ws.cell(row=4, column=col).font = Font(bold=True, color="666666", size=9)
    ws["A5"].font = Font(bold=True, size=12)
    ws["C5"].number_format = _FECHA
    ws.append([])

    ws.append(["#", "Ing. Activo", "Marca Comercial", "Cantidad", "Unidad", "Control"])
    enc = ws.max_row
    for c in range(1, 7):
        celda = ws.cell(row=enc, column=c)
        celda.font = Font(bold=True)
        celda.border = Border(top=grueso, bottom=grueso)
    ws.cell(row=enc, column=4).alignment = Alignment(horizontal="right")
    renglones = sorted(orden["insumos"], key=lambda r: ((r.get("ingredienteActivo") or "").strip(), r["producto"] or ""))
    for n, r in enumerate(renglones, start=1):
        ws.append([n, (r.get("ingredienteActivo") or "").strip() or "—", r["producto"], r["cantidadTotal"], r["unidad"], "☐"])
        f = ws.max_row
        ws.cell(row=f, column=4).number_format = "#,##0.0##"
        ws.cell(row=f, column=6).alignment = Alignment(horizontal="center")
        ws.cell(row=f, column=3).font = Font(bold=True)
        for c in range(1, 7):
            ws.cell(row=f, column=c).border = Border(bottom=fino)
    if orden.get("observaciones"):
        ws.append([])
        ws.append(["Observaciones:", None, orden["observaciones"]])
    for _ in range(4):
        ws.append([])
    ws.append([None, "Entregó (encargado de depósito)", "Retiró (contratista / operario)", None, "Fecha y hora de retiro"])
    f = ws.max_row
    for c in (2, 3, 5):
        ws.cell(row=f, column=c).border = Border(top=grueso)
        ws.cell(row=f, column=c).font = Font(size=9)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{enc}:{enc}"
    ws.oddFooter.right.text = f"Orden N° {orden['idOrden']} · Formulario N° {form['idFormularioRetiro']} · Página &P de &N"
    return _bytes(wb)


def resultado_cultivo_xlsx(id_cultivo: int | None = None, id_campania: int | None = None, id_lote: int | None = None) -> bytes:
    items = resultado.costo_por_cultivo_campania(id_cultivo, id_campania, id_lote)
    wb = Workbook()
    ws = _hoja(
        wb, "Costo por cultivo-campaña",
        [("Cultivo", 20), ("Campaña", 14), ("Lote", 12), ("Insumos ($)", 16), ("Maquinaria ($)", 16), ("Contratista ($)", 16), ("Total ($)", 16)],
    )
    for r in items:
        ws.append([r["cultivo"], r["campania"], r["lote"], r["costoInsumos"], r["costoMaquinaria"], r["costoContratista"], r["costoTotalOrdenes"]])
    _cerrar(ws, {3: _PESOS, 4: _PESOS, 5: _PESOS, 6: _PESOS})
    return _bytes(wb)
