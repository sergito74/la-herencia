"""Convivencia XLSX: vista previa auditable y decisiones explícitas por fila."""

from __future__ import annotations

from src.features.cajas_efectivo.alcance_condominio import evidencia_relacion

from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
import json
from typing import Literal
from uuid import uuid5, NAMESPACE_URL
from zipfile import ZipFile, BadZipFile

import openpyxl
from pydantic import Field

from src.db.connection import atomic_reconciliation
from src.features.cajas_efectivo import operaciones as db, service
from src.features.cajas_efectivo.schemas import Estricto, ConfirmarOperacion

SHEETS = {
    "CajasGiamigli": {"Caja Efectivo Pesos": "giamigli-sa", "Caja Efectivo Dolares": "giamigli-dolares", "Caja chica campo": "campo-chica"},
    "CondominioLSC": {"Caja Efectivo": "condominio-lsc"},
}


class DecisionFila(Estricto):
    idFila: int = Field(gt=0)
    accion: Literal["importar", "vincular", "excluir", "pendiente"]
    idMovimiento: int | None = Field(default=None, gt=0)
    idOperacion: int | None = Field(default=None, gt=0)
    motivo: str | None = Field(default=None, min_length=1, max_length=255)
    operacion: ConfirmarOperacion | None = None
    confirmarDistinto: bool = False


class ConfirmarImportacion(Estricto):
    resoluciones: list[DecisionFila] = Field(max_length=20000)
    aceptarSaldoNegativoLote: bool = False


def _money(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        n = Decimal(str(value))
        return n.quantize(Decimal(".01"), rounding=ROUND_HALF_UP) if n.is_finite() else None
    except InvalidOperation:
        return None


def parse(content, libro, cutoff):
    if libro not in SHEETS or not content or len(content) > 10 * 1024 * 1024:
        raise ValueError("Libro desconocido o archivo fuera del límite de 10 MiB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if sum(i.file_size for i in infos) > 50 * 1024 * 1024 or len(infos) > 5000:
                raise ValueError("El contenido descomprimido excede el límite.")
            if any(
                "vbaproject" in i.filename.lower()
                or "externallinks/" in i.filename.lower()
                or i.flag_bits & 1
                for i in infos
            ):
                raise ValueError("No se admiten macros, enlaces externos ni archivos cifrados.")
        wb = openpyxl.load_workbook(
            BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError("El archivo no es un XLSX válido.") from exc
    rows = []
    try:
        for name, box in SHEETS[libro].items():
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            if ws.max_row and ws.max_row > 20001:
                raise ValueError("La hoja supera las 20.000 filas de datos.")
            saldo_relacionado = Decimal(0)
            for index, values in enumerate(ws.iter_rows(min_row=2, max_col=9, values_only=True), 2):
                if index > 20001:
                    raise ValueError("La hoja supera el límite de filas.")
                v = list(values) + [None] * 9
                if not any(x is not None for x in values):
                    continue
                when = (
                    v[0].date()
                    if isinstance(v[0], datetime)
                    else v[0]
                    if isinstance(v[0], date)
                    else None
                )
                if when and when > cutoff:
                    continue
                if box in ("giamigli-sa", "giamigli-dolares"):
                    amount, account, balance = _money(v[6]), v[2], _money(v[7])
                    concept, detail = v[1], v[3]
                elif box == "campo-chica":
                    debit, credit = _money(v[3]), _money(v[4])
                    amount = (
                        (credit or Decimal(0)) - (debit or Decimal(0))
                        if debit is not None or credit is not None
                        else None
                    )
                    account, balance, concept, detail = "B", _money(v[6]), v[1], v[2]
                else:
                    amount, account, balance, concept, detail = (
                        _money(v[2]),
                        "B",
                        _money(v[3]),
                        v[1],
                        None,
                    )
                if box == "condominio-lsc":
                    if not evidencia_relacion(concept, when, amount, index):
                        continue
                    if amount is not None:
                        saldo_relacionado += amount
                    balance = saldo_relacionado if amount is not None else None
                normalized = {
                    "caja": box,
                    "fecha": str(when) if when else None,
                    "importe": str(amount) if amount is not None else None,
                    "concepto": str(concept or "")[:255],
                    "detalle": str(detail or "")[:255],
                    "numeroDocumento": str(v[5])
                    if box == "giamigli-sa" and v[5] is not None
                    else None,
                    "cuenta": {"White": "A", "W": "A", "Blue": "B"}.get(
                        str(account), str(account or "B")
                    ),
                }
                pending = not when or amount is None or amount == 0
                rows.append(
                    {
                        "hoja": name,
                        "fila": index,
                        "normalizada": normalized,
                        "huella": service.huella(normalized),
                        "saldoPlanilla": str(balance) if balance is not None else None,
                        "cruda": [str(x) if x is not None else None for x in values],
                        "estado": "Ambigua" if pending else "Nueva",
                        "motivo": "Fecha/importe faltante, cero o fórmula sin valor cacheado."
                        if pending
                        else None,
                    }
                )
    finally:
        wb.close()
    if not rows:
        raise ValueError("No hay filas en las hojas admitidas al corte elegido.")
    return rows


def _candidates(row):
    n = row["normalizada"]
    if not n["fecha"] or not n["importe"]:
        return []
    return db.read(
        "SELECT TOP 20 IdMovimiento idMovimiento,IdOperacionCaja idOperacion,FilaPlanilla filaPlanilla,"
        "Concepto concepto,Cuenta cuenta,Anulada anulada FROM dbo.vw_MovimientosCajaEfectivoOperativos "
        "WHERE Caja=? AND CAST(Fecha AS date)=? AND Importe=? ORDER BY IdMovimiento",
        (db.CAJAS[n["caja"]], n["fecha"], Decimal(n["importe"])),
    )


@atomic_reconciliation
def previsualizar(content, libro, cutoff, user):
    parsed = parse(content, libro, cutoff)
    fingerprint = sha256(content).hexdigest()
    previous = db.one(
        "SELECT TOP 1 IdLote FROM dbo.ImportacionesCaja "
        "WHERE Libro=? AND HuellaArchivo=? AND FechaCorte=? AND VersionParser=2 ORDER BY IdLote DESC",
        (libro, fingerprint, cutoff),
    )
    if previous:
        return detalle(previous["IdLote"])
    id_lote = db.insert(
        "INSERT INTO dbo.ImportacionesCaja "
        "(Libro,HuellaArchivo,FechaCorte,VersionParser,Usuario) OUTPUT INSERTED.IdLote VALUES (?,?,?,2,?)",
        (libro, fingerprint, cutoff, user),
    )
    for row in parsed:
        candidates = _candidates(row)
        # Only an existing stable historical row with equal content is identified
        # automatically; manual same-amount operations remain human decisions.
        exact = [
            x
            for x in candidates
            if x["filaPlanilla"] == row["fila"]
            and x["concepto"] == row["normalizada"]["concepto"]
            and not x["anulada"]
            and {"White": "A", "W": "A", "Blue": "B"}.get(
                x.get("cuenta"), x.get("cuenta") or row["normalizada"]["cuenta"]
            )
            == row["normalizada"]["cuenta"]
        ]
        linked = exact[0] if len(exact) == 1 and row["estado"] == "Nueva" else None
        prior_source = db.one(
            "SELECT TOP 1 f.Huella,f.IdMovimiento FROM dbo.FilasOrigenCaja f "
            "JOIN dbo.ImportacionesCaja l ON l.IdLote=f.IdLote "
            "WHERE l.Libro=? AND f.Hoja=? AND f.NumeroFila=? AND l.IdLote<>? "
            "ORDER BY l.IdLote DESC",
            (libro, row["hoja"], row["fila"], id_lote),
        )
        if prior_source and prior_source["Huella"] != row["huella"]:
            row["estado"] = "Modificada"
            row["motivo"] = (
                "Cambió el contenido de una fila ya vista; revisar modificación o desplazamiento."
            )
            linked = None
        if linked:
            row["estado"] = "Vinculada"
        elif candidates and row["estado"] == "Nueva":
            row["estado"] = "Ambigua"
            row["motivo"] = "Coincidencias requieren vinculación explícita."
        elif row["estado"] == "Nueva":
            position = db.one(
                "SELECT TOP 1 IdMovimiento FROM dbo.vw_MovimientosCajaEfectivoOperativos "
                "WHERE Caja=? AND FilaPlanilla=?",
                (db.CAJAS[row["normalizada"]["caja"]], row["fila"]),
            )
            if position:
                row["estado"] = "Modificada"
                row["motivo"] = (
                    "La fila histórica en esa posición tiene otro contenido; revisar desplazamiento/modificación."
                )
        row["candidatas"] = candidates
        db.insert(
            "INSERT INTO dbo.FilasOrigenCaja "
            "(IdLote,Hoja,NumeroFila,Huella,DatosOriginales,Estado,IdMovimiento,IdOperacion,Motivo) "
            "OUTPUT INSERTED.IdFila VALUES (?,?,?,?,?,?,?,?,?)",
            (
                id_lote,
                row["hoja"],
                row["fila"],
                row["huella"],
                db.dumps(row),
                row["estado"],
                linked["idMovimiento"] if linked else None,
                linked["idOperacion"] if linked else None,
                row["motivo"],
            ),
        )
    return detalle(id_lote)


def detalle(id_lote):
    lot = db.one("SELECT * FROM dbo.ImportacionesCaja WHERE IdLote=?", (id_lote,))
    if not lot:
        raise service.NoExisteCaja("Importación inexistente.")
    rows = db.read(
        "SELECT * FROM dbo.FilasOrigenCaja WHERE IdLote=? ORDER BY Hoja,NumeroFila", (id_lote,)
    )
    result = []
    comparison = {}
    for row in rows:
        original = json.loads(row["DatosOriginales"])
        n = original["normalizada"]
        candidates = _candidates(original)
        linked_candidate = next((c for c in candidates if c['idMovimiento']==row['IdMovimiento']),None)
        differences = []
        if linked_candidate and n['caja'] in ('giamigli-sa','giamigli-dolares'):
            actual_account = {'White':'A','W':'A','Blue':'B'}.get(linked_candidate.get('cuenta'),linked_candidate.get('cuenta'))
            if actual_account != n['cuenta']:
                differences.append('Cuenta A/B')
            if linked_candidate['concepto'] != n['concepto']:
                differences.append('Concepto')
        result.append(
            {
                "idFila": row["IdFila"],
                "hoja": row["Hoja"],
                "fila": row["NumeroFila"],
                "estado": row["Estado"],
                "datos": n,
                "motivo": row["Motivo"],
                "idMovimiento": row["IdMovimiento"],
                "idOperacion": row["IdOperacion"],
                "candidatas": candidates,
                "diferenciasVinculo": differences,
            }
        )
        comp = comparison.setdefault(
            n["caja"], {"sumaFilas": Decimal(0), "saldoPlanilla": None, "pendientes": 0}
        )
        if n["importe"]:
            comp["sumaFilas"] += Decimal(n["importe"])
        if original["saldoPlanilla"] is not None:
            comp["saldoPlanilla"] = Decimal(original["saldoPlanilla"])
        if row["Estado"] in ("Nueva", "Ambigua", "Modificada"):
            comp["pendientes"] += 1
    for box, comp in comparison.items():
        system = db.one(
            "SELECT ISNULL(SUM(Importe),0) saldo FROM dbo.vw_MovimientosCajaEfectivoOperativos "
            "WHERE Caja=? AND Anulada=0 AND CAST(Fecha AS date)<=?",
            (db.CAJAS[box], lot["FechaCorte"]),
        )["saldo"]
        comp["saldoSistema"] = f"{system:.2f}"
        box_rows = [x for x in result if x["datos"]["caja"] == box]
        comp["faltantesSistema"] = [
            x["idFila"] for x in box_rows if x["estado"] == "Nueva" and not x["candidatas"]
        ]
        comp["ambiguas"] = [x["idFila"] for x in box_rows if x["estado"] == "Ambigua"]
        comp["modificadas"] = [x["idFila"] for x in box_rows if x["estado"] == "Modificada" or x['diferenciasVinculo']]
        comp["excluidas"] = [x["idFila"] for x in box_rows if x["estado"] == "Excluida"]
        linked = {x["idMovimiento"] for x in box_rows if x["idMovimiento"] is not None}
        all_movements = db.read(
            "SELECT TOP 20001 IdMovimiento id FROM dbo.vw_MovimientosCajaEfectivoOperativos "
            "WHERE Caja=? AND Anulada=0 AND CAST(Fecha AS date)<=? ORDER BY IdMovimiento",
            (db.CAJAS[box], lot["FechaCorte"]),
        )
        comp["soloSistema"] = [x["id"] for x in all_movements[:20000] if x["id"] not in linked]
        comp["comparacionTruncada"] = len(all_movements) > 20000
        comp["notaVinculos"] = (
            "Sin vínculo confirmado; revisar también las filas ambiguas antes de atribuir omisión."
        )
        comp["cobertura"] = (
            "Solo movimientos documentados con Giamigli; saldo de fuente = suma del subconjunto, no efectivo total del Condominio."
            if box == "condominio-lsc"
            else "Libro al corte"
        )
        comp["diferencia"] = (
            f"{system - comp['saldoPlanilla']:.2f}" if comp["saldoPlanilla"] is not None else None
        )
        comp["sumaFilas"] = f"{comp['sumaFilas']:.2f}"
        comp["saldoPlanilla"] = (
            f"{comp['saldoPlanilla']:.2f}" if comp["saldoPlanilla"] is not None else None
        )
    return {
        "idLote": id_lote,
        "libro": lot["Libro"],
        "estado": lot["Estado"],
        "fechaCorte": str(lot["FechaCorte"])[:10],
        "filas": result,
        "comparacion": comparison,
    }


@atomic_reconciliation
def confirmar(id_lote, body, user):
    lot = db.one("SELECT * FROM dbo.ImportacionesCaja WHERE IdLote=?", (id_lote,))
    if not lot:
        raise service.NoExisteCaja("Importación inexistente.")
    serialized = db.dumps(body.model_dump(mode="json"))
    if lot["Resoluciones"] == serialized:
        return detalle(id_lote)
    if lot["Estado"] == "Confirmado":
        if lot["Resoluciones"] != serialized:
            raise service.ConflictoCaja("El lote ya se confirmó con otras decisiones.")
        return detalle(id_lote)
    rows = {
        r["IdFila"]: r
        for r in db.read("SELECT * FROM dbo.FilasOrigenCaja WHERE IdLote=?", (id_lote,))
    }
    seen = set()
    # Validate every user preview against the original state before applying any
    # row. Later previews may be rebased only for changes caused by this batch,
    # under the same locks; all capacities are checked again by crear().
    for decision in body.resoluciones:
        if decision.accion == "importar" and decision.operacion:
            service._bloquear(decision.operacion.operacion)
            before = service.previsualizar(decision.operacion.operacion)
            if before["previewVersion"] != decision.operacion.previewVersion:
                raise service.ConflictoCaja(
                    "Cambió el contexto de una fila. Actualizá sus vistas previas."
                )
            if not set(before["advertencias"]).issubset(decision.operacion.advertenciasAceptadas):
                raise service.ConflictoCaja("Faltan advertencias aceptadas en una fila.")
    for decision in body.resoluciones:
        if decision.idFila not in rows or decision.idFila in seen:
            raise ValueError("Fila ajena al lote o resolución repetida.")
        seen.add(decision.idFila)
        row = rows[decision.idFila]
        original = json.loads(row["DatosOriginales"])
        n = original["normalizada"]
        movement, operation = None, None
        state = row["Estado"]
        if decision.accion == "pendiente":
            continue
        if decision.accion == "excluir":
            if not decision.motivo:
                raise ValueError("Excluir requiere motivo.")
            state = "Excluida"
        elif decision.accion == "vincular":
            if bool(decision.idMovimiento) == bool(decision.idOperacion):
                raise ValueError("Elegir un movimiento o una operación para vincular.")
            candidates = _candidates(original)
            chosen = [
                x
                for x in candidates
                if (decision.idMovimiento and x["idMovimiento"] == decision.idMovimiento)
                or (decision.idOperacion and x["idOperacion"] == decision.idOperacion)
            ]
            if len(chosen) != 1 or chosen[0]["anulada"]:
                raise service.ConflictoCaja(
                    "El vínculo no coincide de forma única con caja/fecha/importe vigente."
                )
            if n["caja"] == "condominio-lsc" and not evidencia_relacion(n["concepto"], n["fecha"], n["importe"], row["NumeroFila"]):
                raise ValueError("La fila no tiene relación documental con Giamigli y no puede vincularse.")
            movement, operation = chosen[0]["idMovimiento"], chosen[0]["idOperacion"]
            actual_account = {'White':'A','W':'A','Blue':'B'}.get(chosen[0].get('cuenta'),chosen[0].get('cuenta'))
            if n['caja'] in ('giamigli-sa','giamigli-dolares') and actual_account != n['cuenta'] and not decision.motivo:
                raise ValueError('Vincular con otra cuenta A/B requiere motivo; la comparación conservará la diferencia.')
            duplicate = db.one(
                "SELECT IdFila FROM dbo.FilasOrigenCaja WHERE IdLote=? "
                "AND IdMovimiento=? AND IdFila<>?",
                (id_lote, movement, decision.idFila),
            )
            if duplicate:
                raise service.ConflictoCaja(
                    "Ese movimiento ya representa otra fila del mismo lote."
                )
            state = "Vinculada"
        else:
            if row["Estado"] in ("Vinculada", "Excluida"):
                raise service.ConflictoCaja(
                    "La fila ya fue resuelta; conservar su vínculo o exclusión."
                )
            if (row["Estado"] == "Modificada" or _candidates(original)) and not (
                decision.confirmarDistinto and decision.motivo
            ):
                raise service.ConflictoCaja(
                    "Hay coincidencias/cambios: vincular o confirmar operación distinta con motivo."
                )
            if not decision.operacion:
                raise ValueError(
                    "Importar requiere clasificar la operación y confirmar su vista previa."
                )
            op = decision.operacion.operacion
            if n["caja"] == "condominio-lsc" and not evidencia_relacion(n["concepto"], n["fecha"], n["importe"], row["NumeroFila"]):
                raise ValueError("La fila no tiene relación documental con Giamigli y no puede importarse.")
            if n['caja'] in ('giamigli-sa','giamigli-dolares') and n['cuenta'] in ('A','B') and op.cuenta != n['cuenta']:
                raise ValueError('La cuenta A/B debe conservar la marca de la fila de planilla.')
            if (
                op.caja != n["caja"]
                or str(op.fecha) != n["fecha"]
                or service.piernas(op)[op.caja] != Decimal(n["importe"])
            ):
                raise ValueError("La operación no coincide con caja/fecha/importe de la fila.")
            key = uuid5(NAMESPACE_URL, f"LaHerencia:038:{id_lote}:{decision.idFila}")
            fresh = service.previsualizar(op)
            accepted = set(decision.operacion.advertenciasAceptadas)
            if body.aceptarSaldoNegativoLote:
                accepted.add("SALDO_NEGATIVO")
            if not set(fresh["advertencias"]).issubset(accepted):
                raise service.ConflictoCaja(
                    "El conjunto del lote genera nuevas advertencias. Revisá el saldo conjunto y confirmalas."
                )
            rebased = decision.operacion.model_copy(
                update={
                    "previewVersion": fresh["previewVersion"],
                    "advertenciasAceptadas": sorted(accepted),
                }
            )
            result, _ = service.crear(rebased, key, user)
            operation = result["idOperacion"]
            actual = db.one(
                "SELECT IdMovimiento FROM dbo.vw_MovimientosCajaEfectivoOperativos WHERE IdOperacionCaja=? AND Caja=?",
                (operation, db.CAJAS[n["caja"]]),
            )
            movement = actual["IdMovimiento"]
            state = "Vinculada"
        db.write(
            "UPDATE dbo.FilasOrigenCaja SET Estado=?,IdMovimiento=?,IdOperacion=?,Motivo=?,"
            "UsuarioResolucion=?,FechaResolucion=SYSUTCDATETIME() WHERE IdFila=?",
            (state, movement, operation, decision.motivo, user, decision.idFila),
        )
    pending = db.one(
        "SELECT COUNT(*) n FROM dbo.FilasOrigenCaja WHERE IdLote=? "
        "AND Estado IN ('Nueva','Ambigua','Modificada')",
        (id_lote,),
    )["n"]
    db.write(
        "UPDATE dbo.ImportacionesCaja SET Estado=?,Resoluciones=?,FechaConfirmacion=SYSUTCDATETIME() "
        "WHERE IdLote=?",
        ("Preview" if pending else "Confirmado", serialized, id_lote),
    )
    return detalle(id_lote)
