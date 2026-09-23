"""Exportaciones a planilla (.xlsx) de Tesorería (014 US3).

Mismo patrón que `backend/src/features/ordenes/exportacion.py` (research.md
§2 de 014).
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.features.tesoreria import repository

_PESOS = '"$" #,##0.00;[Red]-"$" #,##0.00'
_FECHA = "dd/mm/yyyy"


def _hoja(wb: Workbook, titulo: str, columnas: list[tuple[str, float]]):
    ws = wb.active
    ws.title = titulo
    ws.append([c for c, _ in columnas])
    for i, (_, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3D2B")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    return ws


def _cerrar(ws, formatos: dict[int, str]) -> None:
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for col, fmt in formatos.items():
            fila[col].number_format = fmt
    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _bytes(wb: Workbook) -> bytes:
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


def valores_propios_xlsx(fecha_desde: date | None, fecha_hasta: date | None) -> bytes:
    """Listado completo de valores propios (cheques), US3 FR-005/FR-006."""
    rows, _ = repository.get_movimientos(
        "valores-propios", fecha_desde, fecha_hasta, page=1, page_size=5000
    )
    wb = Workbook()
    ws = _hoja(
        wb, "Valores propios",
        [
            ("N° Cheque", 14), ("Fecha emisión", 14), ("Fecha vencimiento", 16),
            ("Importe", 16), ("Cobrado", 10), ("Fecha cobro", 14),
            ("N° Cuenta", 16), ("Comentarios", 30),
        ],
    )
    for v in rows:
        ws.append([
            v["numeroCheque"], v["fechaEmision"], v["fechaVencimiento"],
            v["importe"], v["cobrado"], v["fechaCobro"], v["numeroCuenta"], v["comentarios"],
        ])
    _cerrar(ws, {1: _FECHA, 2: _FECHA, 3: _PESOS, 5: _FECHA})
    return _bytes(wb)
