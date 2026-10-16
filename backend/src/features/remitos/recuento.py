"""Carga del recuento físico de stock desde la planilla (exportacion.recuento_xlsx).

Lee la planilla completada en el campo y compara lo contado contra la existencia del
sistema a la fecha del recuento. La diferencia no es una pérdida aparte (regla de
Sergio, 2026-10-15): el recuento es aproximado y las aplicaciones no son exactas, así
que se reparte entre las órdenes de trabajo que usaron el producto desde el recuento
anterior (en el primero, desde la última vez que el stock quedó en cero),
proporcional a lo usado. Cada orden conserva lo cargado y recibe un ajuste vinculado
(Ordenes_Trabajo_AjustesRecuento); su consumo efectivo pasa a ser lo cargado más el
ajuste, y el costo llega a los lotes y cultivos por la distribución de la orden. Solo
lo que no tiene órdenes en el período queda como faltante/sobrante sin imputar.

Los productos se identifican por la columna oculta "Id producto"; las planillas
anteriores, sin esa columna, se resuelven por el nombre exacto del producto.
"""

from __future__ import annotations

import io
import re
from datetime import date

from openpyxl import load_workbook

from src.db.connection import execute_write_transaction, fetch_all
from src.db.params import as_sql_datetime
from src.features.remitos import stock_repo

TOLERANCIA = 0.005
UMBRAL_REVISAR = 20  # % del uso del período


def _texto(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().upper()


def _numero(v) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(".", "").replace(",", ".") if "," in str(v) else str(v).strip()
    try:
        return float(s)
    except ValueError:
        raise ValueError([f"«{v}» no es una cantidad válida."]) from None


def leer_planilla(contenido: bytes) -> dict:
    """Filas contadas de la planilla y, si se completó, la fecha del recuento."""
    try:
        ws = load_workbook(io.BytesIO(contenido), data_only=True).active
    except Exception as exc:  # archivo que no es un .xlsx
        raise ValueError(["No se pudo leer el archivo: tiene que ser la planilla de recuento en Excel (.xlsx)."]) from exc
    encabezado, columnas, fecha = None, {}, None
    for fila in ws.iter_rows(min_row=1, max_row=15):
        for celda in fila:
            m = re.search(r"Fecha del recuento:\s*(\d{1,2})/(\d{1,2})/(\d{4})", str(celda.value or ""))
            if m:
                fecha = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        nombres = {_texto(c.value): c.column for c in fila if c.value}
        if "PRODUCTO" in nombres and "CANTIDAD CONTADA" in nombres:
            encabezado, columnas = fila[0].row, nombres
            break
    if encabezado is None:
        raise ValueError(["La planilla no tiene las columnas «Producto» y «Cantidad contada»."])
    filas = []
    for fila in ws.iter_rows(min_row=encabezado + 1):
        valores = {c.column: c.value for c in fila}
        nombre = valores.get(columnas["PRODUCTO"])
        n = valores.get(columnas.get("N°", 1))
        if not nombre or not isinstance(n, (int, float)):
            continue  # títulos de grupo, firma, filas vacías
        id_col = columnas.get("ID PRODUCTO")
        filas.append({
            "fila": fila[0].row,
            "idProducto": int(valores[id_col]) if id_col and valores.get(id_col) else None,
            "producto": str(nombre).strip(),
            "contada": _numero(valores.get(columnas["CANTIDAD CONTADA"])),
            "observaciones": (str(valores.get(columnas.get("OBSERVACIONES"))) if valores.get(columnas.get("OBSERVACIONES")) else None),
        })
    return {"fecha": fecha, "filas": filas}


def _kardex_hasta(id_producto: int, fecha: date) -> list[dict]:
    k = stock_repo.kardex(id_producto) or {"movimientos": []}
    return [m for m in k["movimientos"] if str(m["fecha"])[:10] <= fecha.isoformat()]  # orden cronológico, con saldo


def _desde(id_producto: int, fecha: date, movimientos: list[dict]) -> date | None:
    """Inicio del período cuyas órdenes absorben la diferencia: el recuento anterior del
    producto; en el primero, la última vez que su stock quedó en cero (None = desde siempre)."""
    previo = fetch_all(
        "SELECT MAX(r.Fecha) AS f FROM dbo.Stock_Recuentos r JOIN dbo.Stock_Recuentos_Detalle d ON d.IdRecuento = r.IdRecuento "
        "WHERE r.Anulado = 0 AND d.IdProducto = ? AND r.Fecha < ?",
        (id_producto, as_sql_datetime(fecha)),
    )[0]["f"]
    if previo:
        return date.fromisoformat(str(previo)[:10])
    # El ciclo actual del stock empieza en el último cero seguido de una entrada. No sirve
    # "el último cero" a secas: si la última orden dejó el stock en cero (DIFLU 50 RINDER,
    # orden 141), esa misma orden quedaría afuera y es la que hay que ajustar.
    cero = None
    for actual, siguiente in zip(movimientos, movimientos[1:]):
        if abs(actual["saldo"]) <= TOLERANCIA and siguiente["cantidad"] > 0:
            cero = str(actual["fecha"])[:10]
    return date.fromisoformat(cero) if cero else None


def _ordenes_del_periodo(id_producto: int, desde: date | None, hasta: date) -> list[dict]:
    """Renglones de órdenes de trabajo que usaron el producto en (desde, hasta], con su
    consumo efectivo (lo cargado más los ajustes de recuentos anteriores)."""
    filas = fetch_all(
        """
        SELECT oti.IdOrdenInsumo AS idOrdenInsumo, ot.IdOrdenTrabajo AS idOrdenTrabajo,
               COALESCE(ot.FechaEjecucion, ot.FechaPedido) AS fecha,
               oti.CantidadTotal + ISNULL((
                   SELECT SUM(ar.Cantidad) FROM dbo.Ordenes_Trabajo_AjustesRecuento ar
                   JOIN dbo.Stock_Recuentos_Detalle rd ON rd.IdRecuentoDetalle = ar.IdRecuentoDetalle
                   JOIN dbo.Stock_Recuentos r ON r.IdRecuento = rd.IdRecuento
                   WHERE r.Anulado = 0 AND ar.IdOrdenInsumo = oti.IdOrdenInsumo), 0) AS cantidad
        FROM dbo.Ordenes_Trabajo_Insumos oti JOIN dbo.Ordenes_Trabajo ot ON ot.IdOrdenTrabajo = oti.IdOrdenTrabajo
        WHERE ot.Estado <> 'Anulada' AND oti.IdProducto = ?
          AND COALESCE(ot.FechaEjecucion, ot.FechaPedido) <= ?
          AND (? IS NULL OR COALESCE(ot.FechaEjecucion, ot.FechaPedido) > ?)
        ORDER BY 3, 1
        """,
        (id_producto, as_sql_datetime(hasta), as_sql_datetime(desde), as_sql_datetime(desde)),
    )
    return [{**f, "fecha": str(f["fecha"])[:10], "cantidad": float(f["cantidad"])} for f in filas if float(f["cantidad"]) > TOLERANCIA]


def _repartir(diferencia: float, ordenes: list[dict]) -> tuple[list[dict], float]:
    """Reparte la diferencia (contado − sistema) entre las órdenes, proporcional a lo que
    usó cada una. Devuelve los ajustes por orden (+ = consumió más) y lo que no se pudo
    repartir, en signo de stock (sin órdenes, o un sobrante mayor que todo lo usado)."""
    consumo_extra = -diferencia
    total = sum(o["cantidad"] for o in ordenes)
    if total <= TOLERANCIA:
        return [], diferencia
    imputable = max(consumo_extra, -total)  # un sobrante no puede dejar una orden en negativo
    ajustes, acumulado = [], 0.0
    for i, o in enumerate(ordenes):
        cant = round(imputable - acumulado, 4) if i == len(ordenes) - 1 else round(imputable * o["cantidad"] / total, 4)
        acumulado += cant
        ajustes.append({**o, "ajuste": cant})
    resto = round(-(consumo_extra - imputable), 4)
    return ajustes, resto if abs(resto) > TOLERANCIA else 0.0


def previsualizar(contenido: bytes, fecha: date | None, asignaciones: dict[int, int] | None = None) -> dict:
    """`asignaciones`: fila de la planilla -> IdProducto elegido a mano en la vista previa,
    para los productos que no se pudieron identificar por nombre."""
    planilla = leer_planilla(contenido)
    asignaciones = asignaciones or {}
    fecha = fecha or planilla["fecha"]
    if fecha is None:
        raise ValueError(["Indicá la fecha del recuento."])
    items = {i["idProducto"]: i for i in stock_repo.existencias(estado="todos")["items"]}
    por_nombre: dict[str, list[int]] = {}
    for i in items.values():
        por_nombre.setdefault(_texto(i["producto"]), []).append(i["idProducto"])

    renglones, sin_contar = [], 0
    for f in planilla["filas"]:
        if f["contada"] is None:
            sin_contar += 1
            continue
        pid = asignaciones.get(f["fila"]) or f["idProducto"]
        if pid is None:
            candidatos = por_nombre.get(_texto(f["producto"]), [])
            if not candidatos:
                candidatos = [int(r["id"]) for r in fetch_all(
                    "SELECT CAST(IdProducto AS int) AS id FROM dbo.vw_ProductosBase WHERE UPPER(LTRIM(RTRIM(Producto))) = ?", (_texto(f["producto"]),))]
            if len(candidatos) != 1:
                renglones.append({**f, "estado": "noEncontrado" if not candidatos else "ambiguo", "sistema": None, "diferencia": None,
                                  "unidadBase": None, "ordenes": [], "sinImputar": 0.0, "desde": None,
                                  "usadoPeriodo": 0.0, "porcentajeAjuste": None, "revisar": False})
                continue
            pid = candidatos[0]
        info = items.get(pid, {})
        movimientos = _kardex_hasta(pid, fecha)
        sistema = round(movimientos[-1]["saldo"], 4) if movimientos else 0.0
        diferencia = round(f["contada"] - sistema, 4)
        estado = "sinDiferencia" if abs(diferencia) < TOLERANCIA else ("faltante" if diferencia < 0 else "sobrante")
        desde = _desde(pid, fecha, movimientos)
        ordenes, sin_imputar = ([], 0.0) if estado == "sinDiferencia" else _repartir(diferencia, _ordenes_del_periodo(pid, desde, fecha))
        usado = sum(o["cantidad"] for o in ordenes)
        # Peso del ajuste sobre lo usado en el período: más de UMBRAL_REVISAR sugiere una
        # orden sin cargar o un error de recuento más que imprecisión de las aplicaciones.
        porcentaje = round(sum(o["ajuste"] for o in ordenes) / usado * 100, 1) if usado > TOLERANCIA else None
        renglones.append({
            **f, "idProducto": pid, "producto": info.get("producto") or f["producto"], "unidadBase": info.get("unidadBase"),
            "sistema": sistema, "diferencia": diferencia, "estado": estado, "costoPromedio": info.get("costoPromedio"),
            "desde": desde.isoformat() if desde else None, "ordenes": ordenes, "sinImputar": sin_imputar,
            "usadoPeriodo": round(usado, 4), "porcentajeAjuste": porcentaje,
            "revisar": abs(sin_imputar) > TOLERANCIA or (porcentaje is not None and abs(porcentaje) > UMBRAL_REVISAR),
        })
    return {
        "fecha": fecha.isoformat(),
        "renglones": renglones,
        "sinContar": sin_contar,
        "resumen": {e: sum(1 for r in renglones if r["estado"] == e) for e in ("faltante", "sobrante", "sinDiferencia", "noEncontrado", "ambiguo")},
        "conSinImputar": sum(1 for r in renglones if abs(r["sinImputar"]) > TOLERANCIA),
        "aRevisar": sum(1 for r in renglones if r.get("revisar")),
    }


def aplicar(contenido: bytes, fecha: date | None, motivo: str, asignaciones: dict[int, int] | None = None) -> dict:
    """Recalcula la vista previa (no confía en lo que mande el navegador) y graba el
    recuento en una sola transacción: cada producto contado queda como detalle (marca el
    inicio del próximo período), las diferencias como ajustes de las órdenes del período
    y lo que no se puede repartir como faltante/sobrante sin imputar."""
    if not motivo.strip():
        raise ValueError(["Indicá el motivo del recuento."])
    prev = previsualizar(contenido, fecha, asignaciones)
    problemas = [f"Fila {r['fila']}: «{r['producto']}» no se encontró en el catálogo." if r["estado"] == "noEncontrado"
                 else f"Fila {r['fila']}: «{r['producto']}» coincide con más de un producto." for r in prev["renglones"] if r["estado"] in ("noEncontrado", "ambiguo")]
    if problemas:
        raise ValueError(problemas)
    f_rec = as_sql_datetime(date.fromisoformat(prev["fecha"]))
    stmts: list = [("INSERT INTO dbo.Stock_Recuentos (Fecha, Motivo) OUTPUT INSERTED.IdRecuento VALUES (?, ?)", (f_rec, motivo.strip()[:255]))]
    ajustes_orden = sin_imputar = 0
    for r in prev["renglones"]:
        idx_ajuste = None
        if abs(r["sinImputar"]) > TOLERANCIA:
            sin_imputar += 1
            idx_ajuste = len(stmts)
            sufijo = f" — {r['observaciones']}" if r.get("observaciones") else ""
            stmts.append((
                "INSERT INTO dbo.Stock_Ajustes (Fecha, IdProducto, Cantidad, CostoUnitario, Motivo) OUTPUT INSERTED.IdAjuste VALUES (?, ?, ?, ?, ?)",
                (f_rec, r["idProducto"], r["sinImputar"], r["costoPromedio"] if r["sinImputar"] > 0 else None,
                 f"{motivo.strip()} (sin órdenes para repartir){sufijo}"[:255]),
            ))
        idx_detalle = len(stmts)

        def detalle(res, r=r, idx_ajuste=idx_ajuste):
            return (
                "INSERT INTO dbo.Stock_Recuentos_Detalle (IdRecuento, IdProducto, Sistema, Contado, DesdeFecha, IdAjusteSinImputar, Observaciones) "
                "OUTPUT INSERTED.IdRecuentoDetalle VALUES (?, ?, ?, ?, ?, ?, ?)",
                (res[0], r["idProducto"], r["sistema"], r["contada"],
                 as_sql_datetime(date.fromisoformat(r["desde"])) if r["desde"] else None,
                 res[idx_ajuste] if idx_ajuste is not None else None, (r.get("observaciones") or "")[:255] or None),
            )

        stmts.append(detalle)
        for o in r["ordenes"]:
            if abs(o["ajuste"]) > 0.00005:
                ajustes_orden += 1

                def ajuste(res, o=o, idx_detalle=idx_detalle):
                    return ("INSERT INTO dbo.Ordenes_Trabajo_AjustesRecuento (IdRecuentoDetalle, IdOrdenInsumo, Cantidad) VALUES (?, ?, ?)",
                            (res[idx_detalle], o["idOrdenInsumo"], o["ajuste"]))

                stmts.append(ajuste)
    resultado = execute_write_transaction(stmts)
    return {"idRecuento": resultado[0], "ajustesOrden": ajustes_orden, "sinImputar": sin_imputar, **prev["resumen"]}


def listar() -> list[dict]:
    return fetch_all(
        """
        SELECT r.IdRecuento AS idRecuento, r.Fecha AS fecha, r.Motivo AS motivo, r.Anulado AS anulado, r.MotivoAnulacion AS motivoAnulacion,
               COUNT(d.IdRecuentoDetalle) AS productos,
               SUM(CASE WHEN ABS(d.Contado - d.Sistema) > 0.005 THEN 1 ELSE 0 END) AS conDiferencia,
               (SELECT COUNT(*) FROM dbo.Ordenes_Trabajo_AjustesRecuento ar JOIN dbo.Stock_Recuentos_Detalle x ON x.IdRecuentoDetalle = ar.IdRecuentoDetalle
                WHERE x.IdRecuento = r.IdRecuento) AS ajustesOrden,
               SUM(CASE WHEN d.IdAjusteSinImputar IS NOT NULL THEN 1 ELSE 0 END) AS sinImputar
        FROM dbo.Stock_Recuentos r LEFT JOIN dbo.Stock_Recuentos_Detalle d ON d.IdRecuento = r.IdRecuento
        GROUP BY r.IdRecuento, r.Fecha, r.Motivo, r.Anulado, r.MotivoAnulacion
        ORDER BY r.Fecha DESC, r.IdRecuento DESC
        """
    )


def anular(id_recuento: int, motivo: str) -> None:
    """Anula el recuento entero: sus ajustes por orden dejan de contar y se anulan los
    faltantes/sobrantes sin imputar que generó."""
    if not motivo.strip():
        raise ValueError(["Indicá el motivo de la anulación."])
    if not fetch_all("SELECT 1 AS x FROM dbo.Stock_Recuentos WHERE IdRecuento = ? AND Anulado = 0", (id_recuento,)):
        raise ValueError(["El recuento no existe o ya está anulado."])
    execute_write_transaction([
        ("UPDATE dbo.Stock_Recuentos SET Anulado = 1, MotivoAnulacion = ? WHERE IdRecuento = ?", (motivo.strip()[:255], id_recuento)),
        ("UPDATE dbo.Stock_Ajustes SET Anulado = 1, MotivoAnulacion = ? WHERE IdAjuste IN "
         "(SELECT IdAjusteSinImputar FROM dbo.Stock_Recuentos_Detalle WHERE IdRecuento = ? AND IdAjusteSinImputar IS NOT NULL)",
         (f"Recuento anulado: {motivo.strip()}"[:255], id_recuento)),
    ])
