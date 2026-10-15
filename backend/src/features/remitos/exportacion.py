"""Exportaciones a planilla (.xlsx) de remitos, existencias valorizadas y bajas de stock.

Los importes van como números con formato de celda (`$` / miles y decimales según la
configuración regional de quien abre el archivo) y las fechas como fechas reales.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.features.remitos import repository, stock_repo

_PESOS = '"$" #,##0.00;[Red]-"$" #,##0.00'
_CANT = "#,##0.###"
_FECHA = "dd/mm/yyyy"
_ESTADO_FACTURA = {"SinFactura": "Sin factura", "Facturado": "Facturado"}
_ESTADO_RENGLONES = {"SinVincular": "Sin vincular", "Parcial": "Parcial", "Completo": "Completo", "ConDiferencia": "Con diferencia"}


def _hoja(wb: Workbook, titulo: str, columnas: list[tuple[str, float]], primera: bool = False):
    ws = wb.active if primera else wb.create_sheet(titulo)
    ws.title = titulo
    ws.append([c for c, _ in columnas])
    for i, (_, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3D2B")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    return ws


def _cerrar(ws, formatos: dict[int, str]) -> None:
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for col, fmt in formatos.items():
            fila[col].number_format = fmt
    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _bytes(wb: Workbook) -> bytes:
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


def remitos_xlsx(**filtros) -> bytes:
    lista = repository.listar_remitos(page=1, page_size=100000, **filtros)["items"]
    wb = Workbook()
    ws = _hoja(wb, "Remitos", [("Fecha", 12), ("Proveedor", 34), ("N° de remito", 18), ("Productos", 46), ("Renglones", 10),
                               ("Facturas", 30), ("Factura", 13), ("Renglones vinculados", 16), ("Días sin factura", 12), ("Anulado", 9), ("A revisar", 9)], primera=True)
    for r in lista:
        ws.append([r["fecha"], r["proveedor"], r["nroRemito"], r["productos"], r["renglones"], ", ".join(r["facturas"]) or None,
                   _ESTADO_FACTURA[r["estadoFactura"]], _ESTADO_RENGLONES[r["estadoRenglones"]], r["diasSinFactura"],
                   "Sí" if r["anulado"] else None, "Sí" if r["revisarDuplicado"] else None])
    _cerrar(ws, {0: _FECHA})
    # detalle por renglón
    wd = _hoja(wb, "Renglones", [("Fecha", 12), ("Proveedor", 32), ("N° de remito", 18), ("Producto", 40), ("Tipo", 14), ("Cantidad", 12), ("Unidad", 9),
                                 ("Vencimiento", 12), ("Vinculado", 12), ("Estado", 14), ("Costo unitario", 15), ("Consumido", 12), ("Anulado", 9)])
    for r in lista:
        det = repository.get_remito(r["idRemito"])
        for l in det["renglones"]:
            wd.append([det["fecha"], det["proveedor"], det["nroRemito"], l["producto"], l["tipo"], l["cantidad"], l["unidad"], l["vencimiento"],
                       l["cantidadVinculada"], _ESTADO_RENGLONES[l["estadoVinculo"]], l["costoUnitario"], l["consumido"], "Sí" if det["anulado"] else None])
    _cerrar(wd, {0: _FECHA, 5: _CANT, 7: _FECHA, 8: _CANT, 10: _PESOS, 11: _CANT})
    return _bytes(wb)


def existencias_xlsx(q=None, tipo=None, estado=None) -> bytes:
    datos = stock_repo.existencias(q, tipo, estado)
    wb = Workbook()
    ws = _hoja(wb, "Existencias", [("Producto", 42), ("Tipo", 16), ("Unidad", 9), ("Existencia", 13), ("Valor (FIFO)", 16), ("Costo promedio", 15),
                                   ("Cantidad sin costo", 15), ("Observaciones", 40)], primera=True)
    for i in datos["items"]:
        obs = []
        if i["negativo"]:
            obs.append("Stock negativo")
        if i["cantidadCostoPendiente"] > 0.005:
            obs.append("Parte del stock sin costo (remito sin factura vinculada)")
        if i["equivalenciaPendiente"]:
            obs.append("Falta equivalencia de una presentación")
        ws.append([i["producto"], i["tipo"], i["unidadBase"], i["existencia"], i["valor"], i["costoPromedio"], i["cantidadCostoPendiente"] or None, "; ".join(obs) or None])
    if datos["items"]:
        n = ws.max_row
        ws.append(["Total", None, None, None, f"=SUM(E2:E{n})"])
        for c in range(1, 6):
            ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=5).number_format = _PESOS
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row - (1 if datos["items"] else 0)):
        fila[3].number_format = _CANT
        fila[4].number_format = _PESOS
        fila[5].number_format = _PESOS
        fila[6].number_format = _CANT
    if datos["items"]:
        ws.auto_filter.ref = f"A1:H{ws.max_row - 1}"
    return _bytes(wb)


def bajas_xlsx(**filtros) -> bytes:
    lista = stock_repo.listar_bajas(page=1, page_size=100000, **filtros)["items"]
    wb = Workbook()
    ws = _hoja(wb, "Bajas de stock", [("Fecha", 12), ("Motivo", 24), ("Detalle", 30), ("Rubro", 28), ("Centro de costos", 18), ("Producto", 40), ("Cantidad", 12),
                                      ("Unidad", 9), ("Gasto (FIFO)", 15), ("Costo pendiente", 14), ("Anulada", 9)], primera=True)
    for b in lista:
        for l in b["renglones"]:
            ws.append([b["fecha"], b["motivoLabel"], b["detalle"], b["rubro"], b["centro"], l["producto"], l["cantidad"], l["unidadBase"], l["costo"],
                       "Sí" if l["costoPendiente"] else None, "Sí" if b["anulada"] else None])
    _cerrar(ws, {0: _FECHA, 6: _CANT, 8: _PESOS})
    return _bytes(wb)


def recuento_xlsx(q=None, tipo=None, estado=None) -> bytes:
    """Planilla para el recuento físico en el campo: lista para imprimir, con la existencia
    del sistema y columnas en blanco para anotar lo contado. La diferencia se calcula sola
    cuando se carga la cantidad contada en la planilla."""
    from openpyxl.styles import Border, Side

    datos = stock_repo.existencias(q, tipo, estado)
    wb = Workbook()
    ws = wb.active
    ws.title = "Recuento físico"
    linea = Side(style="thin", color="808080")
    borde = Border(left=linea, right=linea, top=linea, bottom=linea)
    columnas = [("N°", 5), ("Producto", 36), ("Principio activo", 30), ("Unidad", 8), ("Stock sistema", 13), ("Cantidad contada", 14), ("Diferencia", 12), ("Observaciones", 26), ("Id producto", 10)]
    for i, (_, ancho) in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    ws.append(["Recuento físico de stock"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generado: {date.today().strftime('%d/%m/%Y')}", None, None, None, None, "Fecha del recuento: ____/____/______"])
    ws.append(["Responsable: ______________________________", None, None, None, None, "Lugar / depósito: ____________________"])
    ws.append([])
    ws.append([c for c, _ in columnas])
    encabezado = ws.max_row
    for c in range(1, len(columnas) + 1):
        celda = ws.cell(row=encabezado, column=c)
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde
    ws.row_dimensions[encabezado].height = 30

    grupos: dict[str, list] = {}
    for it in datos["items"]:
        grupos.setdefault(it["tipo"] or "Sin tipo", []).append(it)
    n = 0
    for nombre in sorted(grupos):
        ws.append([nombre])
        f = ws.max_row
        ws.merge_cells(start_row=f, start_column=1, end_row=f, end_column=len(columnas))
        ws.cell(row=f, column=1).font = Font(bold=True)
        ws.cell(row=f, column=1).fill = PatternFill("solid", fgColor="E3ECE6")
        for c in range(1, len(columnas) + 1):
            ws.cell(row=f, column=c).border = borde
        for it in grupos[nombre]:
            n += 1
            ws.append([n, it["producto"], it.get("ingredienteActivo"), it["unidadBase"], round(it["existencia"], 2), None, None, None, it["idProducto"]])
            f = ws.max_row
            ws.cell(row=f, column=7).value = f'=IF(F{f}="","",F{f}-E{f})'
            ws.row_dimensions[f].height = 22
            for c in range(1, len(columnas) + 1):
                celda = ws.cell(row=f, column=c)
                celda.border = borde
                celda.alignment = Alignment(vertical="center", wrap_text=c in (2, 3))
            for c in (5, 6, 7):
                ws.cell(row=f, column=c).number_format = "#,##0.00"
                ws.cell(row=f, column=c).alignment = Alignment(horizontal="right", vertical="center")

    # Filas en blanco para lo que aparezca en el depósito sin stock en el sistema: se
    # anota el nombre del producto tal como figura en el catálogo y lo contado.
    ws.append(["Otros productos encontrados (no figuran en el sistema): nombre exacto del producto"])
    f = ws.max_row
    ws.merge_cells(start_row=f, start_column=1, end_row=f, end_column=len(columnas))
    ws.cell(row=f, column=1).font = Font(bold=True)
    ws.cell(row=f, column=1).fill = PatternFill("solid", fgColor="E3ECE6")
    for _ in range(10):
        n += 1
        ws.append([n])
        f = ws.max_row
        ws.cell(row=f, column=7).value = f'=IF(F{f}="","",F{f}-E{f})'
        ws.row_dimensions[f].height = 22
        for c in range(1, len(columnas) + 1):
            ws.cell(row=f, column=c).border = borde

    # Columna oculta: identifica el producto al subir el recuento completado (recuento.py).
    ws.column_dimensions[get_column_letter(len(columnas))].hidden = True
    ws.append([])
    ws.append(["Firma del responsable: ______________________________"])
    ws.freeze_panes = ws.cell(row=encabezado + 1, column=1)
    ws.print_title_rows = f"{encabezado}:{encabezado}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "Página &P de &N"
    return _bytes(wb)
