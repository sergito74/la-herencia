"""Exportaciones del resultado de cultivo a Excel (012, FR-008)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.features.resultado_cultivo import resultado

_MONEDA = '"$" #,##0.00;[Red]-"$" #,##0.00'
_DOLARES = '"us$" #,##0.00;[Red]-"us$" #,##0.00'
_CANTIDAD = "#,##0.###"


def _crear_hoja(wb: Workbook, nombre: str, columnas: list[tuple[str, int]]):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet()
    ws.title = nombre
    ws.append([titulo for titulo, _ in columnas])
    for indice, (_, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=indice)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3D2B")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(indice)].width = ancho
    ws.freeze_panes = "A2"
    return ws


def _cerrar_hoja(ws, formatos: dict[int, str]) -> None:
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for columna, formato in formatos.items():
            fila[columna - 1].number_format = formato
    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _crear_libro(resultados: list[dict], id_campania: int) -> bytes:
    wb = Workbook()
    resumen = _crear_hoja(
        wb,
        "Resultado",
        [
            ("Cultivo", 24), ("Campaña", 16), ("Sup. sembrada (ha)", 18),
            ("Sup. cosechada (ha)", 18), ("Sup. picada (ha)", 16), ("Rinde (kg/ha)", 16),
            ("Costo ($)", 17), ("Costo (us$)", 17), ("Costo/ha sembrada ($)", 20),
            ("Costo/ha sembrada (us$)", 22), ("Costo/ha cosechada ($)", 22),
            ("Costo/ha cosechada (us$)", 24), ("Venta neta ($)", 18),
            ("Venta neta (us$)", 18), ("Margen bruto ($)", 18),
            ("Margen bruto (us$)", 18), ("Rentabilidad ($)", 18),
            ("Rentabilidad (us$)", 18), ("Costos y resultado en dólares parciales", 42),
        ],
    )
    detalle = _crear_hoja(
        wb,
        "Detalle de costos",
        [("Cultivo", 24), ("Campaña", 16), ("Concepto", 28), ("Rubro", 22),
         ("Pesos", 17), ("Dólares", 17), ("Origen", 18), ("Id compra", 15),
         ("Id detalle compra", 20), ("Id orden de trabajo", 22)],
    )

    for item in resultados:
        resumen.append([
            item["cultivo"], item["campania"], item["superficieSembrada"],
            item["superficieCosechada"], item["superficiePicada"], item["rinde"],
            item["costoTotalPesos"], item["costoTotalDolares"],
            item["costoPorHectareaSembradaPesos"], item["costoPorHectareaSembradaDolares"],
            item["costoPorHectareaCosechadaPesos"], item["costoPorHectareaCosechadaDolares"],
            item["ventaNetaPesos"], item["ventaNetaDolares"],
            item["margenBrutoPesos"], item["margenBrutoDolares"],
            item["rentabilidadPesos"], item["rentabilidadDolares"],
            "Sí: faltan importes históricos" if item.get("costeoDolaresIncompleto", False) else "No",
        ])
        for linea in resultado.detalle_costos(item["idCultivo"], id_campania):
            detalle.append([
                item["cultivo"], item["campania"], linea["concepto"], linea["rubro"],
                linea["montoPesos"], linea["montoDolares"], linea["origen"],
                linea["idCompra"], linea["idDetalleCompra"], linea["idOrdenTrabajo"],
            ])

    _cerrar_hoja(resumen, {3: _CANTIDAD, 4: _CANTIDAD, 5: _CANTIDAD, 6: _CANTIDAD,
                           7: _MONEDA, 8: _DOLARES, 9: _MONEDA, 10: _DOLARES,
                           11: _MONEDA, 12: _DOLARES, 13: _MONEDA, 14: _DOLARES,
                           15: _MONEDA, 16: _DOLARES,
                           17: "0.00%", 18: "0.00%"})
    _cerrar_hoja(detalle, {5: _MONEDA, 6: _DOLARES})

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def resultado_campania_xlsx(id_campania: int) -> bytes:
    resumen = resultado.resultado_campania(id_campania)
    items = [resultado.resultado_cultivo(c["idCultivo"], id_campania) for c in resumen["cultivos"]]
    return _crear_libro(items, id_campania)


def resultado_cultivo_xlsx(id_cultivo: int, id_campania: int) -> bytes:
    item = resultado.resultado_cultivo(id_cultivo, id_campania)
    return _crear_libro([item], id_campania)
