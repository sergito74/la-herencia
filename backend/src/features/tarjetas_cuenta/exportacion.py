"""Exportación a Excel de la cuenta de una tarjeta — 034 (FR-011, FR-023)."""

from __future__ import annotations

from openpyxl import Workbook

from src.features.cuentas_corrientes.exportacion import _FECHA, _PESOS, _bytes, _cerrar, _hoja
from src.features.tarjetas_cuenta.schemas import AVISO_SALDO


def cuenta_tarjeta_xlsx(cuenta: dict) -> bytes:
    wb = Workbook()
    ws = _hoja(wb, "Cuenta", [("Fecha", 12), ("Origen", 18), ("Resumen", 14), ("Detalle", 44), ("Proveedor", 32),
                              ("Deuda", 16), ("Crédito", 16), ("Saldo", 16), ("Vínculo", 22)])
    ws.append([None, "Saldo inicial", None, None, None, None, None, cuenta["saldoInicial"], None])
    for f in cuenta["filas"]:
        ws.append([f["fecha"], f["origen"], f.get("codigo"), f.get("detalle"), f.get("proveedor"),
                   f["deuda"], f["credito"], f["saldo"], f.get("estadoVinculo")])
    _cerrar(ws, {0: _FECHA, 5: _PESOS, 6: _PESOS, 7: _PESOS})
    ws.append([])
    ws.append([None, "Saldo final", None, None, None, None, None, cuenta["saldoFinal"]])
    d = cuenta["detalleSaldo"]
    ws.append([None, "Exigible", None, None, None, None, None, d["exigible"]])
    ws.append([None, "Consumo no resumido", None, None, None, None, None, d["noResumido"]])
    for fila in ws.iter_rows(min_row=ws.max_row - 5, max_row=ws.max_row):
        fila[7].number_format = _PESOS
    ws.append([])
    ws.append([AVISO_SALDO])
    return _bytes(wb)


def control_xlsx(control: dict) -> bytes:
    wb = Workbook()
    ws = _hoja(wb, "Control", [("Categoría", 36), ("Tarjeta", 20), ("Fecha", 12), ("Importe", 16), ("Medio", 10),
                               ("Movimiento", 12), ("Resumen", 10), ("Consumo", 10), ("Motivo", 70)])
    for h in control["hallazgos"]:
        ws.append([h["categoria"], h.get("tarjeta"), h.get("fecha"), h.get("importe"), h.get("medio"),
                   h.get("idMovimiento"), h.get("idResumen"), h.get("idLineaConsumo"), h["motivo"]])
    _cerrar(ws, {2: _FECHA, 3: _PESOS})
    return _bytes(wb)
