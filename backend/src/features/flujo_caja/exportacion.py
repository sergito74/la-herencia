"""Exportación a planilla del flujo de caja por rubro (030).

Mismo patrón que `cuentas_corrientes/exportacion.py`: la planilla reproduce
la vista (rango, granularidad, moneda) con importes numéricos, no texto.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.features.flujo_caja import repository

_FORMATO = {"ARS": '"$" #,##0.00;[Red]-"$" #,##0.00', "USD": '"us$" #,##0.00;[Red]-"us$" #,##0.00'}
_TITULO = Font(bold=True, color="FFFFFF")
_FONDO_TITULO = PatternFill("solid", fgColor="1F3D2B")
_FONDO_SECCION = PatternFill("solid", fgColor="E8EFE9")
_FONDO_SIN_APLICAR = PatternFill("solid", fgColor="FBEEDC")
_NEGRITA = Font(bold=True)


def flujo_por_rubro_xlsx(desde: date, hasta: date, granularidad: str, moneda: str) -> bytes:
    d = repository.flujo_por_rubro(desde, hasta, granularidad, moneda)
    periodos = d["periodos"]
    fmt = _FORMATO[moneda]

    wb = Workbook()
    ws = wb.active
    ws.title = "Flujo por rubro"
    ws.append([f"Flujo de caja por rubro — {desde:%d/%m/%Y} a {hasta:%d/%m/%Y} — {granularidad} — {moneda}"])
    ws.append([f"Generado {datetime.now():%d/%m/%Y %H:%M}"])
    ws.append([])
    ws.append(["", *periodos, "Total"])
    for c in ws[ws.max_row]:
        c.font, c.fill = _TITULO, _FONDO_TITULO

    def fila(etiqueta, valores: dict | None, total=None, estilo=None, fondo=None):
        vals = [valores.get(p) if valores else None for p in periodos]
        if total is None and valores:
            total = round(sum(v for v in vals if v is not None), 2)
        ws.append([etiqueta, *vals, total])
        for c in ws[ws.max_row][1:]:
            c.number_format = fmt
        if estilo:
            for c in ws[ws.max_row]:
                c.font = estilo
        if fondo:
            for c in ws[ws.max_row]:
                c.fill = fondo

    def seccion(titulo):
        ws.append([titulo])
        ws[ws.max_row][0].font = _NEGRITA
        ws[ws.max_row][0].fill = _FONDO_SECCION

    sin_aplicar = ("Pendiente de aplicar", "Histórico sin aplicar")
    seccion("Saldo inicial")
    for c in d["saldoInicial"]["cuentas"]:
        etiqueta = c["cuenta"] + (f" ({c['aclaracion']})" if c.get("aclaracion") else "")
        ws.append([etiqueta, c["importe"]])
        ws[ws.max_row][1].number_format = fmt
    ws.append(["Total saldo inicial", d["saldoInicial"]["total"]])
    ws[ws.max_row][0].font = _NEGRITA
    ws[ws.max_row][1].number_format = fmt

    seccion("Ingresos")
    for r in d["ingresos"]["rubros"]:
        fila(r["rubro"], r["valores"], r["total"], fondo=_FONDO_SIN_APLICAR if r["rubro"] in sin_aplicar else None)
    fila("Total ingresos", d["ingresos"]["totalPorPeriodo"], estilo=_NEGRITA)

    seccion("Egresos")
    for g in d["egresos"]["centrosCosto"]:
        ws.append([g["centroCosto"]])
        ws[ws.max_row][0].font = _NEGRITA
        for r in g["rubros"]:
            fila("   " + r["rubro"], r["valores"], r["total"], fondo=_FONDO_SIN_APLICAR if r["rubro"] in sin_aplicar else None)
        fila(f"Subtotal {g['centroCosto']}", g["subtotalPorPeriodo"], g["subtotal"], estilo=_NEGRITA)
    fila("Total egresos", d["egresos"]["totalPorPeriodo"], estilo=_NEGRITA)

    fila("Neto operativo", d["netoOperativoPorPeriodo"], estilo=_NEGRITA)

    seccion("Movimientos entre cuentas propias")
    for r in d["internos"]["rubros"]:
        fila(r["rubro"], r["valores"], r["total"])

    seccion("Saldo final")
    for cuenta, valores in d["saldoFinalPorCuenta"].items():
        ws.append([cuenta, *[valores.get(p) for p in periodos]])
        for c in ws[ws.max_row][1:]:
            c.number_format = fmt
    ws.append(["Total saldo final", *[d["saldoFinalPorPeriodo"].get(p) for p in periodos]])
    for c in ws[ws.max_row]:
        c.font = _NEGRITA
    for c in ws[ws.max_row][1:]:
        c.number_format = fmt

    ws.column_dimensions["A"].width = 46
    for i in range(2, len(periodos) + 3):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = "B5"

    if moneda == "USD" and d["sinTipoCambio"]:
        ws2 = wb.create_sheet("Sin tipo de cambio")
        ws2.append(["Período", "Sección", "Centro de costo", "Rubro", "Movimientos", "Importe ARS no convertido"])
        for c in ws2[1]:
            c.font, c.fill = _TITULO, _FONDO_TITULO
        for s in d["sinTipoCambio"]:
            ws2.append([s["periodo"], s["seccion"], s["centroCosto"], s["rubro"], s["cantidad"], s["importeArs"]])
            ws2[ws2.max_row][5].number_format = _FORMATO["ARS"]
        for col, ancho in zip("ABCDEF", (12, 12, 24, 36, 12, 22)):
            ws2.column_dimensions[col].width = ancho

    b = BytesIO()
    wb.save(b)
    return b.getvalue()
