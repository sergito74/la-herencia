"""Exportaciones a planilla (.xlsx) de Cuentas Corrientes (014 US1/US2).

Mismo patrón que `backend/src/features/ordenes/exportacion.py` (research.md
§2) — sin abstracción compartida entre features, cada módulo repite el
mismo patrón simple en vez de una capa de reporting genérica prematura.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.features.cuentas_corrientes import repository

_PESOS = '"$" #,##0.00;[Red]-"$" #,##0.00'
_FECHA = "dd/mm/yyyy"


def _hoja(wb: Workbook, titulo: str, columnas: list[tuple[str, float]]):
    ws = wb.active
    ws.title = titulo
    ws.append([c for c, _ in columnas])
    for i, (_, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3D2B")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    return ws


def _cerrar(ws, formatos: dict[int, str]) -> None:
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for col, fmt in formatos.items():
            fila[col].number_format = fmt
    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _bytes(wb: Workbook) -> bytes:
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


def cuenta_corriente_xlsx(
    id_contacto: int, fecha_desde: date | None, fecha_hasta: date | None
) -> bytes:
    """Cuenta corriente de un contacto puntual (US1, FR-001)."""
    rows, _ = repository.get_movimientos(id_contacto, fecha_desde, fecha_hasta, page=1, page_size=5000)
    saldo = repository.get_saldo(id_contacto)

    wb = Workbook()
    ws = _hoja(
        wb, "Movimientos",
        [("Fecha", 12), ("Documento", 16), ("N° Documento", 18), ("Deuda", 16), ("Crédito", 16), ("Saldo", 16)],
    )
    for m in rows:
        ws.append([m["fecha"], m["documento"], m["numeroDocumento"], m["deuda"], m["credito"], m.get("saldoParcial")])
    _cerrar(ws, {0: _FECHA, 3: _PESOS, 4: _PESOS, 5: _PESOS})

    if ws.max_row == 1:
        ws.append([None, None, None, None, "Saldo actual:", saldo["saldoParcial"] if saldo else None])
    else:
        ws.cell(row=ws.max_row + 2, column=5, value="Saldo actual:")
        celda_saldo = ws.cell(row=ws.max_row, column=6, value=saldo["saldoParcial"] if saldo else None)
        celda_saldo.number_format = _PESOS

    return _bytes(wb)


def saldos_xlsx(orden: str) -> bytes:
    """Saldo de todos los contactos con movimientos (US2, FR-003)."""
    items = repository.get_saldos_todos(orden)
    wb = Workbook()
    ws = _hoja(wb, "Saldos", [("Razón Social", 36), ("Saldo", 18)])
    for s in items:
        ws.append([s["razonSocial"], s["saldoParcial"]])
    _cerrar(ws, {1: _PESOS})
    return _bytes(wb)
