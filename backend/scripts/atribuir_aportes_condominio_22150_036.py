"""Dos depósitos documentados del Condominio. Inspección por defecto; --aplicar con respaldo.

No crea caja ni movimientos de socio ni modifica importes/fechas bancarias.
"""

import argparse
import hashlib
import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from src.db.connection import get_connection, _assert_target_is_wc
from src.features.vinculos.backup import backup_verificado

BASE = Path(r"C:\Users\Sergio\Documents")
SOURCES = [
    (
        BASE / "Alquileres urbanos" / "Condominio LSC" / "Movimientos LSC.xlsx",
        "Movs",
        [(335, 0, 8), (1227, 0, 8)],
    ),
    (
        BASE / "Alquileres urbanos" / "Condominio LSC" / "Movimientos Condominios.xlsx",
        "Aportes Campo",
        [(6, 1, 3), (27, 1, 3)],
    ),
    (
        BASE / "La Herencia" / "Administracion y gestion" / "Aportes de Capital Directores.xlsx",
        "Depositos 1",
        [(5, 1, 3), (26, 1, 3)],
    ),
]
EXPECTED = [
    {
        "id": 10310,
        "fecha": "2012-02-24",
        "fechaFuente": "2012-02-22",
        "importe": Decimal("18150"),
        "cuenta": 1,
    },
    {
        "id": 10824,
        "fecha": "2013-03-04",
        "fechaFuente": "2013-03-04",
        "importe": Decimal("4000"),
        "cuenta": 2,
    },
]
RULE = "condominio-aportes-bna-10310-10824-036"


def evidencias():
    result = []
    for path, sheet, positions in SOURCES:
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        w = load_workbook(path, read_only=True, data_only=True)
        try:
            for pos, expected in zip(positions, EXPECTED, strict=True):
                row, di, mi = pos
                v = next(w[sheet].iter_rows(min_row=row, max_row=row, values_only=True))
                if (
                    str(v[di])[:10] != expected["fechaFuente"]
                    or Decimal(str(v[mi])) != expected["importe"]
                ):
                    raise ValueError(f"Cambió la evidencia {sheet} fila {row}")
                if sheet == "Movs" and (
                    v[4] != "Retiro Giamigli" or v[5] != "EF" or v[7] != "Giamigli"
                ):
                    raise ValueError("La fuente no confirma relación con Giamigli")
                result.append(
                    {
                        "archivo": str(path),
                        "sha256": sha,
                        "hoja": sheet,
                        "fila": row,
                        "fecha": expected["fechaFuente"],
                        "importe": expected["importe"],
                        "idBNA": expected["id"],
                    }
                )
        finally:
            w.close()
        if hashlib.sha256(path.read_bytes()).hexdigest() != sha:
            raise RuntimeError("La fuente cambió durante la lectura")
    return result


def saldo(c):
    return c.execute(
        "SELECT ISNULL(SUM(Credito-Deuda),0) FROM dbo.vw_MovimientosCuenta_Base WHERE IdContacto=549"
    ).fetchone()[0]


def validar(c):
    state = c.execute(
        "SELECT Estado FROM dbo.RevisionFichas WITH (UPDLOCK,HOLDLOCK) WHERE IdContacto=549"
    ).fetchone()
    if state and state[0] in ("cerrada", "cerrada-con-excepcion"):
        raise RuntimeError("Cuenta549 cerrada; no modificar sin avisar a Sergio")
    for e in EXPECTED:
        r = c.execute(
            "SELECT [Fecha / Hora Mov#],Importe,IdCuentaBancaria,IdContacto FROM dbo.[Movimientos BNA] WITH (UPDLOCK,HOLDLOCK) WHERE IdMovimientoBNA=?",
            e["id"],
        ).fetchone()
        if (
            not r
            or str(r[0])[:10] != e["fecha"]
            or r[1] != e["importe"]
            or r[2] != e["cuenta"]
            or r[3] not in (None, 0, 549)
        ):
            raise RuntimeError("El depósito cambió o está atribuido a otra cuenta")
        for sql in [
            "SELECT 1 FROM dbo.ReasignacionesContacto WITH (UPDLOCK,HOLDLOCK) WHERE IdOrigen=? AND Origen IN ('bna','Banco Nacion','Banco Nación')",
            "SELECT 1 FROM dbo.AplicacionesPago WITH (UPDLOCK,HOLDLOCK) WHERE OrigenMovimiento='bna' AND IdMovimientoOrigen=? AND Anulada=0",
            "SELECT 1 FROM dbo.ConciliacionesTesoreria WITH (UPDLOCK,HOLDLOCK) WHERE Medio='bna' AND IdMovimiento=?",
        ]:
            if c.execute(sql, e["id"]).fetchone():
                raise RuntimeError("Hay efectos previos que requieren revisión")
        if c.execute(
            "SELECT 1 FROM dbo.TraspasosInternosTesoreria WITH (UPDLOCK,HOLDLOCK) WHERE (MedioA='bna' AND IdMovimientoA=?) OR (MedioB='bna' AND IdMovimientoB=?)",
            e["id"],
            e["id"],
        ).fetchone():
            raise RuntimeError("Hay traspasos previos")


def run(aplicar=False):
    _assert_target_is_wc()
    evidence = evidencias()
    with get_connection() as conn:
        c = conn.cursor()
        before = saldo(c)
        contacts = [
            c.execute(
                "SELECT IdContacto FROM dbo.[Movimientos BNA] WHERE IdMovimientoBNA=?", e["id"]
            ).fetchone()[0]
            for e in EXPECTED
        ]
    result = {
        "estado": "inspeccion",
        "evidencias": evidence,
        "movimientos": EXPECTED,
        "contactosAntes": contacts,
        "saldo549Antes": before,
    }
    if not aplicar:
        return result
    if contacts == [549, 549]:
        return {**result, "estado": "ya_atribuidos_sin_escrituras"}
    if any(x not in (0, None) for x in contacts):
        raise RuntimeError("Atribución parcial o distinta: revisar")
    result["backupVerificado"] = backup_verificado("condominio_aportes_22150_036")
    out = Path(__file__).resolve().parents[2] / "backups" / "condominio-ef-2011-2016-036"
    out.mkdir(parents=True, exist_ok=True)
    log = out / f"atribucion-{datetime.now():%Y%m%d-%H%M%S}.json"

    def save():
        log.write_text(
            json.dumps(result, default=str, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    result["estado"] = "respaldado"
    save()
    with get_connection(readonly=False) as conn:
        c = conn.cursor()
        try:
            c.execute("SET XACT_ABORT ON; BEGIN TRANSACTION")
            if c.execute("SELECT DB_NAME()").fetchone()[0] != "WC":
                raise RuntimeError("Solo WC")
            validar(c)
            if saldo(c) != before:
                raise RuntimeError("Cambió la cuenta desde inspección")
            if c.execute(
                "SELECT 1 FROM dbo.AuditoriaCorrecciones WHERE Regla=? AND Estado=?",
                RULE,
                "aplicada",
            ).fetchone():
                raise RuntimeError("Ya existe auditoría aplicada: revisar antes de repetir")
            for e in EXPECTED:
                c.execute(
                    "UPDATE dbo.[Movimientos BNA] SET IdContacto=549,Contacto='Condominio LSC' WHERE IdMovimientoBNA=? AND ISNULL(IdContacto,0)=0",
                    e["id"],
                )
                if c.rowcount != 1:
                    raise RuntimeError("Cambió un depósito concurrentemente")
            after = saldo(c)
            if after != before - Decimal("22150"):
                raise RuntimeError("El saldo no refleja el crédito exacto esperado")
            c.execute(
                "INSERT INTO dbo.AuditoriaCorrecciones (Regla,Estado,Parametros,Usuario,Fecha,UsuarioAplicacion,FechaAplicacion,Respaldo,Resumen) VALUES (?,'aplicada',?,? ,SYSUTCDATETIME(),?,SYSUTCDATETIME(),?,?)",
                RULE,
                json.dumps(
                    {
                        "movimientos": EXPECTED,
                        "evidencias": evidence,
                        "contactoAnterior": 0,
                        "contactoNuevo": 549,
                    },
                    default=str,
                    ensure_ascii=False,
                ),
                "Codex (036, Sergio 23/10/2026)",
                "Codex (036, Sergio 23/10/2026)",
                result["backupVerificado"],
                "Dos depósitos del Condominio documentados en Movs y planillas de aportes, $22150; no crear caja ni asientos de socio duplicados.",
            )
            c.execute("COMMIT TRANSACTION")
            result.update(estado="aplicado", saldo549Despues=after, aporteIdentificado="22150.00")
        except Exception as exc:
            c.execute("IF @@TRANCOUNT>0 ROLLBACK TRANSACTION")
            result.update(estado="revertido", error=str(exc))
            save()
            raise
    save()
    result["informe"] = str(log)
    from src.features.auditoria_cuentas.bimonetaria import invalidar
    from src.features.revision_cuentas.cache import invalidar as revision_invalidar

    invalidar()
    revision_invalidar()
    return result


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--aplicar", action="store_true")
    print(json.dumps(run(p.parse_args().aplicar), default=str, ensure_ascii=True))
