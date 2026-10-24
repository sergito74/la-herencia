"""036: atribución sistemática al Condominio (contacto 549) de los depósitos BNA que corresponden a aportes
de la planilla «Aportes de Capital Directores.xlsx / Depositos 1» (universo $733.781, aportes del Condominio
según decisión de Sergio). Inspección por defecto; --aplicar con respaldo verificado y auditoría.

Regla: por cada fila de la planilla (salvo efectivo, cheque entregado a terceros y préstamo/devolución) se
busca en Movimientos BNA un crédito de IMPORTE EXACTO entre la fecha de la fila y 3 días después. Solo se
atribuye si hay un único candidato, sin contacto (0/NULL) y sin aplicaciones, conciliaciones, traspasos ni
reasignaciones. Lo ambiguo o sin candidato queda listado y NO se toca. No crea caja, socios ni asientos.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from src.db.connection import _assert_target_is_wc, get_connection
from src.features.vinculos.backup import backup_verificado

FUENTE = Path(r"C:\Users\Sergio\Documents\La Herencia\Administracion y gestion\Aportes de Capital Directores.xlsx")
HOJA = "Depositos 1"
RULE = "condominio-aportes-bna-sistematico-036"
EXCLUIR = ("efectivo", "torrontegui", "avena", "prestamo", "préstamo", "devolucion", "devolución")
VENTANA = 3


def fuente() -> tuple[list[dict], str]:
    sha = hashlib.sha256(FUENTE.read_bytes()).hexdigest()
    wb = load_workbook(FUENTE, read_only=True, data_only=True)
    filas = []
    try:
        for i, r in enumerate(wb[HOJA].iter_rows(values_only=True), 1):
            # Columnas B (fecha), C (detalle), D (importe), como B14:D14.
            f, d, imp = (r[1], r[2], r[3]) if len(r) > 3 else (None, None, None)
            if isinstance(f, datetime) and isinstance(imp, (int, float)):
                filas.append({"fila": i, "fecha": f.date(), "detalle": str(d), "importe": Decimal(str(imp))})
    finally:
        wb.close()
    total = sum(f["importe"] for f in filas)
    if total != Decimal("733781"):
        raise RuntimeError(f"La planilla cambió: suma {total} (esperado 733781)")
    return filas, sha


def _efectos(c, i: int) -> bool:
    for sql in (
        "SELECT 1 FROM dbo.ReasignacionesContacto WHERE IdOrigen=? AND Origen IN ('bna','Banco Nacion','Banco Nación')",
        "SELECT 1 FROM dbo.AplicacionesPago WHERE OrigenMovimiento='bna' AND IdMovimientoOrigen=? AND Anulada=0",
        "SELECT 1 FROM dbo.ConciliacionesTesoreria WHERE Medio='bna' AND IdMovimiento=?",
    ):
        if c.execute(sql, i).fetchone():
            return True
    return bool(c.execute("SELECT 1 FROM dbo.TraspasosInternosTesoreria WHERE (MedioA='bna' AND IdMovimientoA=?) OR "
                          "(MedioB='bna' AND IdMovimientoB=?)", i, i).fetchone())


def matriz(c) -> list[dict]:
    filas, _ = fuente()
    usados: set[int] = set()
    out = []
    for f in filas:
        if any(x in f["detalle"].lower() for x in EXCLUIR) or f["importe"] <= 0:
            out.append({**f, "resultado": "excluida (no es depósito bancario)"})
            continue
        cands = c.execute(
            "SELECT IdMovimientoBNA, CONVERT(varchar(10),[Fecha / Hora Mov#],23), Importe, ISNULL(IdContacto,0), Concepto "
            "FROM dbo.[Movimientos BNA] WHERE Importe=? AND [Fecha / Hora Mov#] BETWEEN ? AND ?",
            f["importe"], f["fecha"].isoformat(), (f["fecha"] + timedelta(days=VENTANA)).isoformat()).fetchall()
        cands = [x for x in cands if x[0] not in usados]
        ya = [x for x in cands if x[3] == 549]
        libres = [x for x in cands if x[3] == 0]
        if ya:
            usados.add(ya[0][0])
            out.append({**f, "resultado": "ya atribuido", "idBNA": ya[0][0]})
        elif len(libres) == 1 and len(cands) == 1 and not _efectos(c, libres[0][0]):
            usados.add(libres[0][0])
            out.append({**f, "resultado": "atribuir", "idBNA": libres[0][0], "fechaBNA": libres[0][1], "concepto": libres[0][4]})
        else:
            out.append({**f, "resultado": "sin candidato único", "candidatos": [list(map(str, x)) for x in cands]})
    return out


def saldo549(c) -> Decimal:
    return Decimal(str(c.execute("SELECT ISNULL(SUM(Credito-Deuda),0) FROM dbo.vw_MovimientosCuenta_Base WHERE IdContacto=549").fetchone()[0]))


def run(aplicar=False):
    _assert_target_is_wc()
    _, sha = fuente()
    with get_connection() as conn:
        c = conn.cursor()
        m = matriz(c)
        antes = saldo549(c)
    plan = [x for x in m if x["resultado"] == "atribuir"]
    res = {"estado": "inspeccion", "sha256": sha, "saldo549": antes, "resumen": {}, "plan": plan,
           "noResueltos": [x for x in m if x["resultado"] == "sin candidato único"]}
    for x in m:
        res["resumen"][x["resultado"]] = res["resumen"].get(x["resultado"], 0) + 1
    res["importeAtribuir"] = sum(x["importe"] for x in plan)
    if not aplicar or not plan:
        return res
    res["backup"] = backup_verificado("condominio_aportes_bna_sistematico_036")
    out = Path(__file__).resolve().parents[2] / "backups" / "condominio-ef-2011-2016-036"
    out.mkdir(parents=True, exist_ok=True)
    log = out / f"aportes-bna-sistematico-{datetime.now():%Y%m%d-%H%M%S}.json"
    log.write_text(json.dumps(res, default=str, ensure_ascii=False, indent=2), encoding="utf-8")
    with get_connection(readonly=False) as conn:
        c = conn.cursor()
        try:
            c.execute("SET XACT_ABORT ON; BEGIN TRANSACTION")
            if c.execute("SELECT DB_NAME()").fetchone()[0] != "WC":
                raise RuntimeError("Solo WC")
            estado = c.execute("SELECT Estado FROM dbo.RevisionFichas WHERE IdContacto=549").fetchone()
            if estado and estado[0] in ("cerrada", "cerrada-con-excepcion"):
                raise RuntimeError("Cuenta 549 cerrada: avisar a Sergio antes de modificar")
            if [x for x in matriz(c) if x["resultado"] == "atribuir"] != plan or saldo549(c) != antes:
                raise RuntimeError("Cambió el estado desde la inspección")
            for x in plan:
                c.execute("UPDATE dbo.[Movimientos BNA] SET IdContacto=549, Contacto='Condominio LSC' "
                          "WHERE IdMovimientoBNA=? AND ISNULL(IdContacto,0)=0", x["idBNA"])
                if c.rowcount != 1:
                    raise RuntimeError(f"Cambió el depósito {x['idBNA']}")
            despues = saldo549(c)
            if despues != antes - res["importeAtribuir"]:
                raise RuntimeError(f"Saldo 549 inesperado: {antes} -> {despues}")
            c.execute("INSERT INTO dbo.AuditoriaCorrecciones (Regla,Estado,Parametros,Usuario,Fecha,UsuarioAplicacion,FechaAplicacion,Respaldo,Resumen) "
                      "VALUES (?,'aplicada',?,'Claude036',SYSUTCDATETIME(),'Claude036',SYSUTCDATETIME(),?,?)",
                      RULE, json.dumps({"informe": str(log), "ids": [x["idBNA"] for x in plan]}), res["backup"],
                      f"{len(plan)} depósitos BNA de aportes de la planilla Depositos 1 atribuidos a 549 por importe exacto y fecha (+{VENTANA} días), candidato único.")
            c.execute("COMMIT TRANSACTION")
        except Exception as exc:
            c.execute("IF @@TRANCOUNT>0 ROLLBACK TRANSACTION")
            res["estado"] = f"revertido: {exc}"
            log.write_text(json.dumps(res, default=str, ensure_ascii=False, indent=2), encoding="utf-8")
            raise
    res.update(estado="aplicado", saldo549Despues=despues, informe=str(log))
    log.write_text(json.dumps(res, default=str, ensure_ascii=False, indent=2), encoding="utf-8")
    return {k: res[k] for k in ("estado", "resumen", "importeAtribuir", "saldo549", "saldo549Despues", "backup", "informe")}


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--aplicar", action="store_true")
    r = run(p.parse_args().aplicar)
    print(json.dumps(r, default=str, ensure_ascii=False, indent=1))
