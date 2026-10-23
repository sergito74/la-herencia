"""Archiva por alcance, sin borrar ni modificar los movimientos originales. --aplicar opt-in."""

import argparse
import hashlib
import json
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook
from src.db.connection import get_connection, _assert_target_is_wc
from src.features.vinculos.backup import backup_verificado
from src.features.cajas_efectivo.alcance_condominio import evidencia_relacion

FUENTE = Path(
    r"C:\Users\Sergio\Documents\Alquileres urbanos\Condominio LSC\Movimientos LSC desde Abr 2016.xlsm"
)
OUT = Path(__file__).resolve().parents[2] / "backups" / "condominio-alcance-038"


def dumps(x):
    return json.dumps(x, default=str, ensure_ascii=False, sort_keys=True)


def snapshot(c):
    return [dict(zip([v[0] for v in c.description], r, strict=True)) for r in c.fetchall()]


def estado(c):
    result = {}
    for key, sql in {
        "otrasCajas": "SELECT Caja,COUNT_BIG(*) n,SUM(Importe) saldo FROM dbo.MovimientosCajaEfectivo WHERE Caja<>'CondominioLSC' GROUP BY Caja ORDER BY Caja",
        "socios": "SELECT IdSocio,COUNT_BIG(*) n,SUM(Importe) pesos,SUM(ImporteUSD) usd,SUM(ImporteKgCarne) kg FROM dbo.MovimientosCuentaSocio GROUP BY IdSocio ORDER BY IdSocio",
        "pagos": "SELECT COUNT_BIG(*) n,SUM([Importe imputado]) saldo FROM dbo.[Pagos efectivo]",
        "operaciones": "SELECT COUNT_BIG(*) n FROM dbo.OperacionesCajaEfectivo",
    }.items():
        c.execute(sql)
        result[key] = snapshot(c)
    return result


def preparar():
    _assert_target_is_wc()
    content = FUENTE.read_bytes()
    sha = hashlib.sha256(content).hexdigest()
    w = load_workbook(FUENTE, read_only=True, data_only=True)
    try:
        source = {
            i: v for i, v in enumerate(w["Caja Efectivo"].iter_rows(max_col=4, values_only=True), 1)
        }
        criado = list(w["Abr 16"].iter_rows(min_row=41, max_row=41, max_col=6, values_only=True))[0]
        if "dev. Giamigli" not in str(criado[0]) or Decimal(str(criado[5])) != Decimal("-3818"):
            raise ValueError("Cambió evidencia de Criado")
        with get_connection() as conn:
            c = conn.cursor()
            c.execute(
                "SELECT * FROM dbo.MovimientosCajaEfectivo WHERE Caja='CondominioLSC' ORDER BY IdMovimiento"
            )
            rows = snapshot(c)
            before = estado(c)
        plans = []
        for r in rows:
            if r["IdOperacionCaja"] is not None:
                continue
            v = source.get(r["FilaPlanilla"])
            amount = (
                Decimal(str(v[2])).quantize(Decimal(".01"), rounding=ROUND_HALF_UP)
                if v and isinstance(v[2], (int, float))
                else None
            )
            if (
                not v
                or str(v[0])[:10] != str(r["Fecha"])[:10]
                or str(v[1] or "")[:200] != r["Concepto"]
                or amount != Decimal(str(r["Importe"]))
            ):
                raise ValueError(f"No coincide la fuente de IdMovimiento {r['IdMovimiento']}")
            reason = evidencia_relacion(r["Concepto"], r["Fecha"], r["Importe"], r["FilaPlanilla"])
            plans.append(
                {
                    "id": r["IdMovimiento"],
                    "relacion": bool(reason),
                    "motivo": reason
                    or "Fuera del alcance operativo: sin relación documental comprobada con Giamigli.",
                    "snapshot": r,
                }
            )
        return plans, sha, before, hashlib.sha256(dumps(rows).encode()).hexdigest()
    finally:
        w.close()


def run(aplicar=False):
    plans, sha, before, rows_sha = preparar()
    kept = [p for p in plans if p["relacion"]]
    record = {
        "estado": "inspeccion",
        "fuente": str(FUENTE),
        "sha256Fuente": sha,
        "sha256Originales": rows_sha,
        "clasificadas": len(plans),
        "relacionadas": len(kept),
        "archivadas": len(plans) - len(kept),
        "netoVinculos": sum((Decimal(str(p["snapshot"]["Importe"])) for p in kept), Decimal(0)),
        "antes": before,
        "filas": plans,
    }
    OUT.mkdir(parents=True, exist_ok=True)
    log = OUT / f"alcance-{datetime.now():%Y%m%d-%H%M%S}.json"
    if not aplicar:
        log.write_text(dumps(record), encoding="utf-8")
        return {k: v for k, v in record.items() if k not in ("filas", "antes")} | {
            "informe": str(log)
        }
    record["backupVerificado"] = backup_verificado("condominio_alcance_038")
    record["estado"] = "respaldado"
    log.write_text(dumps(record), encoding="utf-8")
    with get_connection(readonly=False) as conn:
        c = conn.cursor()
        try:
            c.execute("SET XACT_ABORT ON; BEGIN TRANSACTION")
            if c.execute("SELECT DB_NAME()").fetchone()[0] != "WC":
                raise RuntimeError("Solo WC")
            c.execute(
                "SELECT * FROM dbo.MovimientosCajaEfectivo WITH (TABLOCKX,HOLDLOCK) WHERE Caja='CondominioLSC' ORDER BY IdMovimiento"
            )
            if hashlib.sha256(dumps(snapshot(c)).encode()).hexdigest() != rows_sha:
                raise RuntimeError("Cambió el libro; volver a inspeccionar")
            c.execute("""IF OBJECT_ID('dbo.AlcanceHistoricoCaja','U') IS NULL
                CREATE TABLE dbo.AlcanceHistoricoCaja (
                  IdMovimiento int NOT NULL PRIMARY KEY REFERENCES dbo.MovimientosCajaEfectivo(IdMovimiento),
                  RelacionGiamigli bit NOT NULL, Motivo nvarchar(500) NOT NULL,
                  Fuente nvarchar(500) NOT NULL, HashFuente char(64) NOT NULL,
                  Snapshot nvarchar(max) NOT NULL CHECK(ISJSON(Snapshot)=1),
                  Usuario nvarchar(60) NOT NULL, Fecha datetime2 NOT NULL DEFAULT SYSUTCDATETIME())""")
            for p in plans:
                prior = c.execute(
                    "SELECT RelacionGiamigli,Snapshot FROM dbo.AlcanceHistoricoCaja WHERE IdMovimiento=?",
                    p["id"],
                ).fetchone()
                if prior:
                    if bool(prior[0]) != p["relacion"] or prior[1] != dumps(p["snapshot"]):
                        raise RuntimeError(
                            "Clasificación existente diferente; requiere revisión explícita"
                        )
                    continue
                c.execute(
                    "INSERT INTO dbo.AlcanceHistoricoCaja (IdMovimiento,RelacionGiamigli,Motivo,Fuente,HashFuente,Snapshot,Usuario) VALUES (?,?,?,?,?,?,?)",
                    p["id"],
                    p["relacion"],
                    p["motivo"],
                    str(FUENTE) + " / Caja Efectivo",
                    sha,
                    dumps(p["snapshot"]),
                    "Codex: alcance autorizado Sergio 23/10/2026",
                )
            c.execute("""CREATE OR ALTER VIEW dbo.vw_MovimientosCajaEfectivoOperativos AS
                SELECT m.* FROM dbo.MovimientosCajaEfectivo m
                WHERE m.Caja<>'CondominioLSC' OR m.IdOperacionCaja IS NOT NULL OR EXISTS
                (SELECT 1 FROM dbo.AlcanceHistoricoCaja a WHERE a.IdMovimiento=m.IdMovimiento AND a.RelacionGiamigli=1)""")
            after = estado(c)
            if before != after:
                raise RuntimeError("Cambió otra caja/cuenta durante la limpieza")
            c.execute(
                "SELECT * FROM dbo.MovimientosCajaEfectivo WHERE Caja='CondominioLSC' ORDER BY IdMovimiento"
            )
            if hashlib.sha256(dumps(snapshot(c)).encode()).hexdigest() != rows_sha:
                raise RuntimeError("Se modificaron originales")
            c.execute(
                "SELECT COUNT(*),SUM(Importe) FROM dbo.vw_MovimientosCajaEfectivoOperativos WHERE Caja='CondominioLSC'"
            )
            n, total = c.fetchone()
            if n != len(kept) or total != record["netoVinculos"]:
                raise RuntimeError("Vista no reproduce el plan")
            c.execute("COMMIT TRANSACTION")
            record.update(estado="aplicado", despues=after)
        except Exception as exc:
            c.execute("IF @@TRANCOUNT>0 ROLLBACK TRANSACTION")
            record.update(estado="revertido", error=str(exc))
            log.write_text(dumps(record), encoding="utf-8")
            raise
    log.write_text(dumps(record), encoding="utf-8")
    return {k: v for k, v in record.items() if k not in ("filas", "antes", "despues")} | {
        "informe": str(log)
    }


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--aplicar", action="store_true")
    print(dumps(run(p.parse_args().aplicar)))
