"""Lee y normaliza las 6 hojas relevantes de `Cajas Giamigli.xlsx`
(027-migracion-cajas-giamigli). Solo parsea — no escribe nada en `WC`, eso
lo hacen los scripts `migrar_*.py`. Ver research.md §1/§4/§7 y
data-model.md para el detalle de columnas y reglas.

Estructura confirmada por inspección directa (research.md §1):

Hojas de socios (`Cuenta Sergio`, `Cuenta Lucy`, `Cuenta Cond LSC`,
`Cuenta Ceci`) — encabezado en fila 2, datos desde fila 3:
    A Fecha | B Proveedor/Servicio | C Detalle | D Nro. Documento
    E Debe AR$ | F Debe Kg. Carne | G Debe Dolares
    H Haber AR$ | I Haber Kg. Carne | J Haber Dolares
    K Saldos Kg. Carne | L Saldos Dolares | M Forma Pago | N Saldo

`Caja Efectivo Pesos` (Caja='GiamigliSA') — encabezado en fila 1, datos
desde fila 2:
    A Fecha | B Concepto | C Cuenta | D Razon Social | E PC
    F Nro. Documento | G Importe | H Saldo | I Recuento

`Caja chica campo` (Caja='CampoChica') — encabezado en fila 1, datos
desde fila 2:
    A Fecha | B Proveedor/Servicio | C Detalle | D Debe | E Haber
    F Forma Pago | G Saldo
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl

TOLERANCIA_CERO = 1e-9


@dataclass
class CasoARevisar:
    hoja: str
    numero_fila: int
    motivo: str
    datos_crudos: dict = field(default_factory=dict)

    def datos_crudos_json(self) -> str:
        return json.dumps(self.datos_crudos, default=str, ensure_ascii=False)


@dataclass
class FilaSocio:
    numero_fila: int
    fecha: date
    proveedor: str | None
    detalle: str | None
    tipo: str  # 'AsignacionGasto' (Debe) | 'Devolucion' (Haber)
    importe_pesos: float
    importe_usd: float
    importe_kg_carne: float
    forma_pago: str | None


@dataclass
class FilaCaja:
    numero_fila: int
    fecha: date
    concepto: str | None
    detalle: str | None
    importe: float  # con signo: positivo = ingreso, negativo = egreso
    cuenta: str | None = None
    forma_pago: str | None = None
    numero_documento: str | None = None


def _a_fecha(valor) -> date | None:
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return None


def _a_float_o_none(valor) -> float | None:
    if valor is None:
        return None
    try:
        numero = float(valor)
    except (TypeError, ValueError):
        return None
    return numero


def _no_cero(valor: float | None) -> bool:
    return valor is not None and abs(valor) > TOLERANCIA_CERO


def leer_hoja_socio(ruta_excel: str, nombre_hoja: str) -> tuple[list[FilaSocio], list[CasoARevisar]]:
    wb = openpyxl.load_workbook(ruta_excel, data_only=True, read_only=True)
    ws = wb[nombre_hoja]

    filas: list[FilaSocio] = []
    casos: list[CasoARevisar] = []

    for numero_fila, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        (
            fecha_cruda,
            proveedor,
            detalle,
            _nro_doc,
            debe_ars,
            debe_kg,
            debe_usd,
            haber_ars,
            haber_kg,
            haber_usd,
            _saldo_kg,
            _saldo_usd,
            forma_pago,
            *_resto,
        ) = (list(row) + [None] * 14)[:14]

        fecha = _a_fecha(fecha_cruda)
        debe_ars, debe_kg, debe_usd = (_a_float_o_none(x) for x in (debe_ars, debe_kg, debe_usd))
        haber_ars, haber_kg, haber_usd = (_a_float_o_none(x) for x in (haber_ars, haber_kg, haber_usd))

        tiene_algun_dato = any(
            v is not None
            for v in (proveedor, detalle, fecha, debe_ars, debe_kg, debe_usd, haber_ars, haber_kg, haber_usd, forma_pago)
        )
        if not tiene_algun_dato:
            continue  # fila completamente vacía: se ignora en silencio (research.md §4)

        tiene_debe = _no_cero(debe_ars) or _no_cero(debe_kg) or _no_cero(debe_usd)
        tiene_haber = _no_cero(haber_ars) or _no_cero(haber_kg) or _no_cero(haber_usd)

        datos_crudos = {
            "fecha": fecha_cruda,
            "proveedor": proveedor,
            "detalle": detalle,
            "debeARS": debe_ars,
            "debeKg": debe_kg,
            "debeUSD": debe_usd,
            "haberARS": haber_ars,
            "haberKg": haber_kg,
            "haberUSD": haber_usd,
            "formaPago": forma_pago,
        }

        if fecha is None:
            casos.append(CasoARevisar(nombre_hoja, numero_fila, "Sin fecha", datos_crudos))
            continue
        if not tiene_debe and not tiene_haber:
            casos.append(CasoARevisar(nombre_hoja, numero_fila, "Sin ningún importe", datos_crudos))
            continue
        if tiene_debe and tiene_haber:
            casos.append(CasoARevisar(nombre_hoja, numero_fila, "Debe y Haber con valores a la vez", datos_crudos))
            continue

        if tiene_debe:
            tipo = "AsignacionGasto"
            importe_pesos, importe_kg_carne, importe_usd = debe_ars or 0.0, debe_kg or 0.0, debe_usd or 0.0
        else:
            tipo = "Devolucion"
            importe_pesos, importe_kg_carne, importe_usd = haber_ars or 0.0, haber_kg or 0.0, haber_usd or 0.0

        filas.append(
            FilaSocio(
                numero_fila=numero_fila,
                fecha=fecha,
                proveedor=proveedor,
                detalle=detalle,
                tipo=tipo,
                importe_pesos=importe_pesos,
                importe_usd=importe_usd,
                importe_kg_carne=importe_kg_carne,
                forma_pago=forma_pago,
            )
        )

    return filas, casos


def leer_hoja_caja(ruta_excel: str, nombre_hoja: str, caja: str) -> tuple[list[FilaCaja], list[CasoARevisar]]:
    if caja not in ("GiamigliSA", "CampoChica"):
        raise ValueError(f"caja inválida: {caja!r}")

    wb = openpyxl.load_workbook(ruta_excel, data_only=True, read_only=True)
    ws = wb[nombre_hoja]

    filas: list[FilaCaja] = []
    casos: list[CasoARevisar] = []

    for numero_fila, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if caja == "GiamigliSA":
            fecha_cruda, concepto, cuenta, _razon_social, _pc, nro_doc, importe_crudo, *_resto = (
                list(row) + [None] * 8
            )[:8]
            detalle = None
            forma_pago = None
        else:
            fecha_cruda, concepto, detalle, debe, haber, forma_pago, *_resto = (list(row) + [None] * 7)[:7]
            cuenta = None
            nro_doc = None
            debe = _a_float_o_none(debe)
            haber = _a_float_o_none(haber)
            importe_crudo = None if debe is None and haber is None else (haber or 0.0) - (debe or 0.0)

        fecha = _a_fecha(fecha_cruda)
        importe = _a_float_o_none(importe_crudo)

        datos_crudos = {
            "fecha": fecha_cruda,
            "concepto": concepto,
            "detalle": detalle,
            "cuenta": cuenta,
            "formaPago": forma_pago,
            "numeroDocumento": nro_doc,
            "importe": importe,
        }

        tiene_algun_dato = any(v is not None for v in (fecha, concepto, detalle, importe, forma_pago))
        if not tiene_algun_dato:
            continue  # fila completamente vacía: se ignora en silencio (research.md §4)

        if fecha is None:
            casos.append(CasoARevisar(nombre_hoja, numero_fila, "Sin fecha", datos_crudos))
            continue
        if not _no_cero(importe):
            casos.append(CasoARevisar(nombre_hoja, numero_fila, "Sin ningún importe", datos_crudos))
            continue

        filas.append(
            FilaCaja(
                numero_fila=numero_fila,
                fecha=fecha,
                concepto=concepto,
                detalle=detalle,
                importe=importe,
                cuenta=cuenta,
                forma_pago=forma_pago,
                numero_documento=str(nro_doc) if nro_doc is not None else None,
            )
        )

    return filas, casos
