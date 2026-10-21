"""XXXXXXXXX (97, gastos sin proveedor identificado; cuenta informal): pagos en efectivo de la planilla
"Cajas Giamigli.xlsx" que faltaban en el sistema (Sergio 21/10/2026: los pagos en general se hacen en efectivo; no borrar
duplicados).

Para cada comprobante sin pago se toma un movimiento de la planilla del mismo importe (+-$1) y fecha (+-10 días) que sea de
la "Caja chica campo" o figure como XXXXXXXX en la caja de efectivo, y que no esté ya registrado en el sistema. Además la
yegua del 25/11/2021 ($107.000): caja de efectivo 19/11/2021 "Compra Yegua" $100.000 + "Viaje traslado Yegua" $7.000
(Julio Cabrera). Se registran como Pagos efectivo (Pago / Efectivo) en la caja correspondiente. Idempotente (marca en
[Nro Mov] = 90360000 + hoja*10000 + fila de la planilla; hoja 1 = Caja Efectivo Pesos, 2 = Caja chica campo).
"""
import openpyxl

from datetime import date

from scripts.nd_interna_tc_036 import crear
from src.db.connection import _assert_target_is_wc, execute_write_transaction, fetch_all, reconciliation_transaction
from src.features.auditoria_cuentas.bimonetaria import cargar_cuenta, invalidar
from src.features.vinculos.backup import backup_verificado

C = 97
PLANILLA = r'C:\Users\Sergio\Documents\La Herencia\Administracion y gestion\Cuentas a pagar\Cajas Giamigli.xlsx'
CAJA = {'CajaPesos': 'La Herencia', 'CajaCampo': 'CampoChica'}
HOJA = {'CajaPesos': 1, 'CajaCampo': 2}


def marca(a):
    return 90360000 + HOJA[a[0]] * 10000 + a[1]


def movimientos():
    wb = openpyxl.load_workbook(PLANILLA, data_only=True, read_only=True)
    out = []
    for n, r in enumerate(wb['Caja Efectivo Pesos'].iter_rows(min_row=2, values_only=True), start=2):
        if r[0] and hasattr(r[0], 'year') and isinstance(r[6], (int, float)):
            out.append(('CajaPesos', n, r[0], -r[6], str(r[1] or ''), str(r[3] or ''), r[2] or 'Blue'))
    for n, r in enumerate(wb['Caja chica campo'].iter_rows(min_row=2, values_only=True), start=2):
        if r[0] and hasattr(r[0], 'year') and isinstance(r[3], (int, float)):
            out.append(('CajaCampo', n, r[0], r[3], str(r[1] or ''), str(r[2] or ''), 'White'))
    return out


def main():
    _assert_target_is_wc()
    M = movimientos()
    S = fetch_all("SELECT IdPagoEfectivo i,Fecha f,[Importe imputado] a,[Nro Mov] m FROM dbo.[Pagos efectivo] WHERE IdContacto=?", (C,))
    hechos = {int(s['m']) for s in S if s['m']}
    D = fetch_all("SELECT c.IdDeuda i,c.Fecha f,c.[Nro Documento] n,t.GranTotal g FROM dbo.Compras c JOIN dbo.vw_Cns_Total_Compra t ON t.IdDeuda=c.IdDeuda "
                  "WHERE c.IdContacto=? ORDER BY c.Fecha", (C,))
    usados_s, usados_m, pagados = set(), set(), set()
    for d in D:  # comprobantes que ya tienen su pago en el sistema
        m = [s for s in S if s['i'] not in usados_s and abs(s['a'] - d['g']) < 1 and abs((s['f'] - d['f']).days) <= 10]
        if m:
            usados_s.add(m[0]['i']); pagados.add(d['i'])
    for a in M:  # movimientos de la planilla ya registrados
        if marca(a) in hechos:
            usados_m.add((a[0], a[1]))
        elif [s for s in S if s['i'] not in usados_s and abs(s['a'] - a[3]) < 0.02 and abs((s['f'] - a[2]).days) <= 5] and 'XXXX' in a[4] + a[5]:
            usados_m.add((a[0], a[1]))
    nuevos = []
    for d in D:
        if d['i'] in pagados:
            continue
        if d['n'] == 'S/N' and abs(d['g'] - 107000) < 0.01:
            cand = [a for a in M if a[0] == 'CajaPesos' and a[2].year == 2021 and 'Yegua' in a[4] and (a[0], a[1]) not in usados_m]
        else:
            cand = [a for a in M if (a[0], a[1]) not in usados_m and abs(a[3] - d['g']) < 1 and abs((a[2] - d['f']).days) <= 10
                    and (a[0] == 'CajaCampo' or 'XXXX' in a[5] or 'XXXX' in a[4]) and 'Caja Chica' not in a[4]]
            cand = cand[:1]
        for a in cand:
            usados_m.add((a[0], a[1]))
            nuevos.append((a, d))
    print(len(nuevos), 'pagos', round(sum(a[3] for a, _ in nuevos), 2))
    for a, d in nuevos:
        print(' ', a[0], a[2].date(), a[3], a[4][:25], '->', d['n'], d['f'].date())
    if not nuevos:
        saldo = round(cargar_cuenta(C, refrescar=True)['saldoPesos'], 2)
        if abs(saldo) >= 0.01:  # residuo de 15 años de gastos en efectivo: ajuste interno (cuenta informal, Sergio 21/10/2026)
            with reconciliation_transaction():
                print(crear(C, [(date(2026, 9, 30), saldo, '0059799', 'Ajuste cuenta informal')], prefijos=('AJD', 'AJC'),
                            texto='Ajuste cuenta informal', ajusta=False))
            invalidar()
            print('saldo', cargar_cuenta(C, refrescar=True)['saldoPesos'])
        return
    print(backup_verificado('xxxx-pagos-caja-036'))
    with reconciliation_transaction():
        execute_write_transaction([("INSERT INTO dbo.[Pagos efectivo] (IdContacto,IdTipoMovimiento,IdFormaPago,Fecha,Cuenta,Caja,[Importe imputado],[Nro Mov]) "
                                    "VALUES (?,1,3,?,?,?,?,?)", (C, a[2], a[6], CAJA[a[0]], a[3], marca(a))) for a, _ in nuevos])
    invalidar()
    print('saldo', cargar_cuenta(C, refrescar=True)['saldoPesos'])


if __name__ == '__main__': main()
