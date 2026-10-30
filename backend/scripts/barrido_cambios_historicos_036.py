"""036, 30/10/2026: barrido de todo lo histórico (fecha anterior al corte del 25/09/2026) que hoy en WC es distinto de la
base original LaHerencia (Access congelado, solo lectura): registros nuevos, borrados y modificados (importe, contacto,
fecha), más los asientos que solo existen en WC (ajustes de cuenta corriente y movimientos de socios cargados durante la
revisión). No cambia nada: genera una planilla para que Sergio revise. Uso (desde backend/):
python -m scripts.barrido_cambios_historicos_036 <salida.xlsx>
"""
from __future__ import annotations

import sys
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.db.connection import fetch_all

CORTE = "2026-09-25"

# (hoja, tabla, clave, fecha, contacto, importe-SQL, descripción-SQL, documento-SQL)
TABLAS = [
    ("Compras", "Compras", "IdDeuda", "Fecha", "IdContacto",
     "(SELECT SUM(d.Cantidad*d.[Precio Unitario]*(1+ISNULL(d.IVA,0)/100.0)) FROM {db}.dbo.Det_Compras d WHERE d.IdCompra=t.IdDeuda)"
     "+ISNULL(t.[Conceptos no gravados],0)+ISNULL(t.[Ingresos Brutos],0)+{perc}+ISNULL(t.[Gastos Varios],0)"
     "+ISNULL(t.[Ley de Sellos],0)+ISNULL(t.Comision,0)+ISNULL(t.Financiacion,0)+ISNULL(t.Guias,0)",
     "CONCAT(t.[Tipo documento],' ',t.Tipo,' ',t.[Nro Documento],' (',t.Moneda,')')", "t.[Documento Original]"),
    ("Caja (pagos efectivo)", "Pagos efectivo", "IdPagoEfectivo", "Fecha", "IdContacto", "t.[Importe imputado]",
     "CONCAT(t.Caja,' / ',t.Cuenta,' / tipo ',t.IdTipoMovimiento,' forma ',t.IdFormaPago)", "NULL"),
    ("Impuestos", "Impuestos", "IdImpuesto", "Fecha", "IdOrganismo", "t.Importe",
     "CONCAT(t.[Periodo liquidado],' | ',t.[Numero de documento])", "t.[Documento Original]"),
    ("Ajustes de ventas de granos", "Venta Granos_Ajustes", "IdAjuste", None, None, "t.Importe",
     "CONCAT('venta ',t.IdVenta,': ',t.Concepto)", "NULL"),
    ("Retenciones", "Retenciones", "IdRetencionSQL", "Fecha", "IdContacto", "t.Importe", "t.[Numero Certificado]", "t.[Numero Certificado]"),
    ("Retenciones IVA granos", "Retenciones IVA Granos", "IdRetencion", None, None,
     "ISNULL(t.[Importe Retenido],0)-ISNULL(t.[Importe Devolucion],0)",
     "CONCAT('liq ',t.IdLiquidacion,' cert ',t.[Nro Certificado],' devol ',t.[Importe Devolucion])", "t.Imagen"),
    ("Banco BNA (contacto)", "Movimientos BNA", "IdMovimientoBNA", "[Fecha / Hora Mov#]", "IdContacto", "t.Importe", "t.Concepto", "t.Concepto"),
]


def _filas(db: str, hoja, tabla, clave, fecha, contacto, importe, desc, doc) -> dict:
    f = f"t.{fecha}" if fecha and not fecha.startswith("[") else (f"t.{fecha}" if fecha else "NULL")
    c = f"t.{contacto}" if contacto else "NULL"
    sql = (f"SELECT t.[{clave}] AS id, {f} AS fecha, {c} AS contacto, {importe.format(db=db, perc="ISNULL(t.PercepcionIVA,0)" if db == "WC" else "0")} AS importe, "
           f"{desc} AS descr, {doc} AS doc FROM {db}.dbo.[{tabla}] t")
    return {r["id"]: r for r in fetch_all(sql, ())}


def _antes_del_corte(r) -> bool:
    return r["fecha"] is None or str(r["fecha"])[:10] < CORTE


def _num(v) -> float:
    return round(float(v or 0), 2)


def comparar() -> list[dict]:
    nombres = {r["i"]: r["n"] for r in fetch_all("SELECT IdContacto i,[Razon Social] n FROM dbo.Contactos", ())}
    out = []
    for hoja, tabla, clave, fecha, contacto, importe, desc, doc in TABLAS:
        wc = _filas("WC", hoja, tabla, clave, fecha, contacto, importe, desc, doc)
        ac = _filas("LaHerencia", hoja, tabla, clave, fecha, contacto, importe, desc, doc)
        for k in set(wc) | set(ac):
            a, w = ac.get(k), wc.get(k)
            base = w or a
            if not _antes_del_corte(base):
                continue
            cambio, detalle = None, []
            if a is None:
                cambio = "Nuevo (no estaba en Access)"
            elif w is None:
                cambio = "Borrado (estaba en Access)"
            else:
                if abs(_num(a["importe"]) - _num(w["importe"])) >= 0.01:
                    detalle.append(f"importe {_num(a['importe']):,.2f} → {_num(w['importe']):,.2f}")
                if a["contacto"] != w["contacto"]:
                    detalle.append(f"contacto {nombres.get(a['contacto'], a['contacto'])} → {nombres.get(w['contacto'], w['contacto'])}")
                if str(a["fecha"])[:10] != str(w["fecha"])[:10]:
                    detalle.append(f"fecha {str(a['fecha'])[:10]} → {str(w['fecha'])[:10]}")
                if detalle:
                    cambio = "Modificado"
            if not cambio:
                continue
            diferencia = _num(w["importe"] if w else 0) - _num(a["importe"] if a else 0)
            out.append({"hoja": hoja, "cambio": cambio, "id": k, "fecha": str(base["fecha"])[:10] if base["fecha"] else "",
                        "contacto": nombres.get(base["contacto"], base["contacto"] or ""), "importeAccess": _num(a["importe"]) if a else None,
                        "importeWC": _num(w["importe"]) if w else None, "diferencia": diferencia,
                        "descripcion": str(base["descr"] or ""), "detalle": "; ".join(detalle),
                        "documento": bool((w or {}).get("doc"))})
    # Solo en WC: ajustes de cuenta corriente y movimientos de socios de la revisión con fecha anterior al corte.
    for r in fetch_all("SELECT IdAjuste i,IdContacto c,Fecha f,Lado l,Importe m,Motivo mo,Usuario u FROM dbo.AjustesCuentaCorriente", ()):
        if str(r["f"])[:10] < CORTE:
            out.append({"hoja": "Ajustes de cuenta corriente", "cambio": "Nuevo (solo en el sistema)", "id": r["i"], "fecha": str(r["f"])[:10],
                        "contacto": nombres.get(r["c"], r["c"]), "importeAccess": None, "importeWC": _num(r["m"]),
                        "diferencia": _num(r["m"]) * (1 if r["l"] == "Deuda" else -1), "descripcion": f"{r['l']}: {r['mo']}",
                        "detalle": f"cargado por {r['u']}", "documento": False})
    for r in fetch_all("SELECT IdMovimiento i,IdSocio s,Tipo t,Importe m,Fecha f,Medio me,Motivo mo,Usuario u,Anulada a "
                       "FROM dbo.MovimientosCuentaSocio WHERE Usuario NOT LIKE 'migracion%'", ()):
        if str(r["f"])[:10] < CORTE and not r["a"]:
            out.append({"hoja": "Cuenta de socios", "cambio": "Nuevo (revisión)", "id": r["i"], "fecha": str(r["f"])[:10],
                        "contacto": {1: "Sergio", 2: "Lucy"}.get(r["s"], r["s"]), "importeAccess": None, "importeWC": _num(r["m"]),
                        "diferencia": _num(r["m"]), "descripcion": f"{r['t']} {r['me'] or ''}: {r['mo']}", "detalle": f"cargado por {r['u']}",
                        "documento": False})
    return clasificar(out)


def clasificar(filas: list[dict]) -> list[dict]:
    """Saca los renumerados (borrado + nuevo con misma fecha, contacto e importe: el mismo movimiento con otro número) y
    asigna prioridad: Alta = sin documento por $1.000 o más, borrados y cambios de importe de $1.000 o más; Baja = con
    documento, o diferencias de menos de $1; Media = el resto."""
    borrados = {}
    for f in filas:
        if f["cambio"].startswith("Borrado"):
            borrados.setdefault((f["hoja"], f["fecha"], f["contacto"], f["importeAccess"]), []).append(f)
    quitar = set()
    for f in filas:
        if f["cambio"].startswith("Nuevo (no estaba"):
            pares = borrados.get((f["hoja"], f["fecha"], f["contacto"], f["importeWC"]))
            if pares:
                quitar.add(id(pares.pop()))
                quitar.add(id(f))
    out = [f for f in filas if id(f) not in quitar]
    for f in out:
        d = abs(f["diferencia"])
        if d < 1:
            f["prioridad"] = "Baja"
        elif f["cambio"].startswith("Borrado") or (f["cambio"] == "Modificado" and d >= 1000) or (not f["documento"] and d >= 1000):
            f["prioridad"] = "Alta"
        elif f["documento"]:
            f["prioridad"] = "Baja"
        else:
            f["prioridad"] = "Media"
    return out


def escribir(filas: list[dict], salida: str) -> None:
    wb = Workbook()
    res = wb.active
    res.title = "Resumen"
    res.append(["Barrido de cambios históricos (fecha anterior al 25/09/2026) contra Access", "", "", ""])
    res.append(["Hoja", "Prioridad", "Cantidad", "Importe total afectado"])
    resumen = {}
    for f in filas:
        k = (f["hoja"], f["prioridad"])
        n, s = resumen.get(k, (0, 0.0))
        resumen[k] = (n + 1, s + abs(f["diferencia"]))
    for (h, c), (n, s) in sorted(resumen.items(), key=lambda x: -x[1][1]):
        res.append([h, c, n, round(s, 2)])
    cols = ["Prioridad", "Con documento", "Cambio", "Id", "Fecha", "Contacto/organismo", "Importe Access", "Importe hoy", "Diferencia", "Descripción", "Detalle",
            "Decisión de Sergio"]
    for hoja in dict.fromkeys(f["hoja"] for f in filas):
        ws = wb.create_sheet(hoja[:31])
        ws.append(cols)
        orden = {"Alta": 0, "Media": 1, "Baja": 2}
        for f in sorted((x for x in filas if x["hoja"] == hoja), key=lambda x: (orden[x["prioridad"]], -abs(x["diferencia"]))):
            ws.append([f["prioridad"], "Sí" if f["documento"] else "No", f["cambio"], f["id"], f["fecha"], f["contacto"], f["importeAccess"], f["importeWC"], f["diferencia"],
                       f["descripcion"], f["detalle"], ""])
    for ws in wb.worksheets:
        for c in ws[1 if ws.title != "Resumen" else 2]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2F5233")
        for i, ancho in enumerate((10, 10, 28, 12, 12, 34, 16, 16, 16, 60, 50, 30), 1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                if isinstance(c.value, float):
                    c.number_format = '#,##0.00;[Red]-#,##0.00'
        ws.freeze_panes = "A2" if ws.title != "Resumen" else "A3"
    wb.save(salida)


if __name__ == "__main__":
    filas = comparar()
    escribir(filas, sys.argv[1])
    resumen = {}
    for f in filas:
        k = (f["hoja"], f["prioridad"])
        n, s = resumen.get(k, (0, 0.0))
        resumen[k] = (n + 1, s + abs(f["diferencia"]))
    for (h, c), (n, s) in sorted(resumen.items()):
        print(f"{h:32} {c:8} {n:5} {s:18,.2f}")
