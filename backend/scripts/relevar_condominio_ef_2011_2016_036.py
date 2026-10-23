"""Matriz de investigación SOLO LECTURA. No crea pagos, vínculos ni asientos.

Coincidencia de fecha/importe = candidato, no conciliación automática.
"""

import csv
import hashlib
import json
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook
from src.features.cajas_efectivo.operaciones import read

FUENTE = Path(r"C:\Users\Sergio\Documents\Alquileres urbanos\Condominio LSC\Movimientos LSC.xlsx")
OUT = Path(__file__).resolve().parents[2] / "backups" / "condominio-ef-2011-2016-036"
DESDE = "2011-01-01"
HASTA = "2016-04-01"


def money(v):
    return Decimal(str(v)).quantize(Decimal(".01"), rounding=ROUND_HALF_UP)


def day(v):
    return date.fromisoformat(str(v)[:10])


def normalized(v):
    return "".join(
        c
        for c in unicodedata.normalize("NFKD", str(v or "").lower())
        if not unicodedata.combining(c)
    )


def candidates(rows, amount, when, days, sign=1):
    return [
        r
        for r in rows
        if money(r["importe"]) == money(amount) * sign
        and abs((day(r["fecha"]) - when).days) <= days
    ]


def run():
    source_sha = hashlib.sha256(FUENTE.read_bytes()).hexdigest()
    w = load_workbook(FUENTE, read_only=True, data_only=True)
    sources = []
    try:
        for i, v in enumerate(w["Movs"].iter_rows(max_col=10, values_only=True), 1):
            if (
                i < 2
                or not isinstance(v[0], datetime)
                or not DESDE <= v[0].date().isoformat() < HASTA
            ):
                continue
            if str(v[5]).strip() != "EF" or not any(
                "giamigli" in normalized(v[j]) for j in (2, 4, 7, 9)
            ):
                continue
            if not isinstance(v[8], (int, float)):
                raise ValueError(f"Importe fuente inválido en fila {i}")
            sources.append(
                {
                    "fila": i,
                    "fecha": v[0].date().isoformat(),
                    "sentido": v[1],
                    "referencia": v[2],
                    "numero": v[3],
                    "clasificacion": v[4],
                    "tipoCaja": v[5],
                    "rubro": v[6],
                    "responsable": v[7],
                    "importe": money(v[8]),
                    "comentario": v[9],
                }
            )
    finally:
        w.close()
    cash = read(
        "SELECT IdMovimiento id,Caja caja,Fecha fecha,Concepto concepto,Detalle detalle,Importe importe,IdPagoEfectivo pago,FilaPlanilla fila,IdOperacionCaja operacion FROM dbo.vw_MovimientosCajaEfectivoOperativos WHERE Fecha>=? AND Fecha<? AND Anulada=0 ORDER BY Fecha,IdMovimiento",
        (DESDE, HASTA),
    )
    payments = read(
        "SELECT p.IdPagoEfectivo id,p.Fecha fecha,p.Caja caja,p.[Importe imputado] importe,p.IdContacto contacto,c.[Razon Social] nombre FROM dbo.vw_PagosEfectivoVigentes p LEFT JOIN dbo.Contactos c ON c.IdContacto=p.IdContacto WHERE p.Fecha>=? AND p.Fecha<? AND p.[Importe imputado]>0 ORDER BY p.Fecha,p.IdPagoEfectivo",
        (DESDE, HASTA),
    )
    banks = read(
        "SELECT IdMovimientoBNA id,[Fecha / Hora Mov#] fecha,Concepto concepto,Importe importe,IdContacto contacto,IdCuentaBancaria cuenta FROM dbo.[Movimientos BNA] WHERE [Fecha / Hora Mov#]>=? AND [Fecha / Hora Mov#]<? AND Importe>0 ORDER BY [Fecha / Hora Mov#],IdMovimientoBNA",
        (DESDE, HASTA),
    )
    socios = read(
        "SELECT IdMovimiento id,Fecha fecha,Tipo tipo,Importe importe,Origen origen,IdOrigen idOrigen,Motivo motivo,ImporteUSD usd,ImporteKgCarne kg FROM dbo.MovimientosCuentaSocio WHERE IdSocio=3 AND Anulada=0 AND Fecha>=? AND Fecha<? ORDER BY Fecha,IdMovimiento",
        (DESDE, HASTA),
    )
    fichas = read(
        "SELECT IdContacto,Estado,Corte,SaldoAlCierre FROM dbo.RevisionFichas WHERE IdContacto IN (549,390,564,375)"
    )
    all_cash = read(
        "SELECT Caja,COUNT_BIG(*) filas,SUM(Importe) saldo FROM dbo.vw_MovimientosCajaEfectivoOperativos WHERE Anulada=0 GROUP BY Caja ORDER BY Caja"
    )
    index = defaultdict(list)
    for c in cash:
        if c["pago"]:
            index[c["pago"]].append(c["id"])
    matrix = []
    for s in sources:
        when = day(s["fecha"])
        amount = s["importe"]
        ref = normalized(s["referencia"])
        expense = s["clasificacion"] == "Compra Giamigli"
        row = {
            **s,
            "candidatosPago": [],
            "candidatosCaja": [],
            "candidatosBanco": [],
            "candidatosSocio": candidates(socios, amount, when, 5),
            "observaciones": [],
        }
        if expense:
            row["candidatosPago"] = candidates(payments, amount, when, 5)
            row["candidatosCaja"] = candidates(
                [x for x in cash if x["caja"] == "GiamigliSA"], amount, when, 5, -1
            )

            def words(value):
                return set(normalized(value).split()) - {"pago", "de", "del", "la", "el", "y"}

            row["discrepanciasProveedor"] = [
                p["id"]
                for p in row["candidatosPago"]
                if not (words(s["referencia"]) & words(p["nombre"]))
            ]
            if row["discrepanciasProveedor"]:
                row["observaciones"].append(
                    "Proveedor distinto entre fuente y sistema: no vincular ni reasignar automáticamente."
                )
            owner = [p for p in row["candidatosPago"] if p["caja"] == "Condominio"]
            linked = [p for p in owner if index[p["id"]]]
            row["estado"] = (
                "candidato_pago_condominio_con_caja"
                if linked
                else "candidato_pago_condominio"
                if owner
                else "solo_candidatos_caja_o_pago"
                if row["candidatosCaja"] or row["candidatosPago"]
                else "sin_contraparte_identificada"
            )
            row["pagoCondominioConCaja"] = [
                {"pago": p["id"], "cajas": index[p["id"]]} for p in linked
            ]
            legs = []
            for c in row["candidatosCaja"]:
                legs.extend(
                    x
                    for x in cash
                    if x["concepto"] == "Aporte Condominio LSC (pago directo)"
                    and f"egreso mov. {c['id']})" in (x["detalle"] or "")
                )
            row["aportesCompensatoriosExistentes"] = legs
            if row["candidatosPago"] or row["candidatosCaja"]:
                row["observaciones"].append(
                    "Fecha/importe solo generan candidatos: corroborar proveedor/comprobante y correspondencia uno a uno."
                )
        else:
            row["candidatosBanco"] = candidates(banks, amount, when, 4)
            row["candidatosCaja"] = candidates(
                [x for x in cash if x["caja"] == "GiamigliSA"], amount, when, 4
            )
            row["registroCondominioExistente"] = candidates(
                [x for x in cash if x["caja"] == "CondominioLSC"], amount, when, 4, -1
            )
            if "deposit" in ref:
                row["estado"] = (
                    "deposito_bancario_ya_asignado_549"
                    if any(b["contacto"] == 549 for b in row["candidatosBanco"])
                    else "deposito_bancario_a_verificar"
                )
                row["observaciones"].append(
                    "EF en origen es billete depositado al banco: no crear ingreso a caja Giamigli."
                )
            elif ref == "aporte efectivo":
                row["estado"] = (
                    "aporte_caja_existente_sin_socio_historico"
                    if row["candidatosCaja"]
                    else "aporte_caja_a_verificar"
                )
            elif row["registroCondominioExistente"]:
                row["estado"] = "ya_en_registro_condominio_2016"
                row["observaciones"].append(
                    "Se superpone con fuente desde abril2016: conservar un único movimiento físico."
                )
            else:
                row["estado"] = "retiro_a_investigar"
        if not row["candidatosSocio"]:
            row["observaciones"].append(
                "Sin asiento candidato en socio3 en este período; no significa automáticamente que falte un asiento: cotejar banco/cuenta549."
            )
        matrix.append(row)
    # Repeticiones de candidato entre filas impiden asignaciones automáticas.
    used = defaultdict(list)
    for r in matrix:
        for p in r["candidatosPago"]:
            used[p["id"]].append(r["fila"])
    shared = {str(k): v for k, v in used.items() if len(v) > 1}
    groups = defaultdict(lambda: {"filas": 0, "importe": Decimal(0)})
    for r in matrix:
        groups[r["estado"]]["filas"] += 1
        groups[r["estado"]]["importe"] += r["importe"]
    if hashlib.sha256(FUENTE.read_bytes()).hexdigest() != source_sha:
        raise RuntimeError("La fuente cambió durante el relevamiento")
    report = {
        "estado": "relevamiento_solo_lectura",
        "fuente": str(FUENTE),
        "sha256Fuente": source_sha,
        "desde": DESDE,
        "hastaExclusivo": HASTA,
        "filas": len(matrix),
        "grupos": dict(groups),
        "totalComprasEF": sum(
            (r["importe"] for r in matrix if r["clasificacion"] == "Compra Giamigli"), Decimal(0)
        ),
        "totalRetirosEF": sum(
            (r["importe"] for r in matrix if r["clasificacion"] == "Retiro Giamigli"), Decimal(0)
        ),
        "movimientosSocio3Periodo": len(socios),
        "fichas": fichas,
        "cajasAlRelevar": all_cash,
        "pagosCandidatosCompartidos": shared,
        "matriz": matrix,
        "advertencia": "Ninguna coincidencia constituye aplicación, cierre ni deuda nueva. No se modificaron WC ni Excel.",
    }
    OUT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    path = OUT / f"matriz-{stamp}.json"
    path.write_text(json.dumps(report, default=str, ensure_ascii=False, indent=2), encoding="utf-8")
    columns = [
        "fila",
        "fecha",
        "clasificacion",
        "referencia",
        "importe",
        "estado",
        "pagos",
        "cajas",
        "bancos",
        "socios",
        "observaciones",
    ]
    with (OUT / f"matriz-{stamp}.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        for r in matrix:
            writer.writerow(
                {
                    **{k: r[k] for k in columns[:6]},
                    "pagos": ";".join(str(x["id"]) for x in r["candidatosPago"]),
                    "cajas": ";".join(str(x["id"]) for x in r["candidatosCaja"]),
                    "bancos": ";".join(str(x["id"]) for x in r["candidatosBanco"]),
                    "socios": ";".join(str(x["id"]) for x in r["candidatosSocio"]),
                    "observaciones": " | ".join(r["observaciones"]),
                }
            )
    return {
        k: v
        for k, v in report.items()
        if k not in ("matriz", "pagosCandidatosCompartidos", "cajasAlRelevar", "fichas")
    } | {"informe": str(path), "candidatosCompartidos": len(shared)}


if __name__ == "__main__":
    print(json.dumps(run(), default=str, ensure_ascii=True))
