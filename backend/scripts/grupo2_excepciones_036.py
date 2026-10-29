"""036 — grupo 2: planilla de excepciones (facturas con diferencia >0,10% contra la tarjeta o sin vínculo), con la
causa probable según la regla del 29/10/2026. Solo lectura.
Uso (desde backend/): python -m scripts.grupo2_excepciones_036 <salida.xlsx>
"""

from __future__ import annotations

import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from src.db.connection import fetch_all
from scripts.grupo2_redondeo_tarjeta_036 import casos as _  # noqa: F401  (misma definición de cuentas)


def _pesos(x: float) -> str:
    return f"$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


if __name__ == "__main__":
    cuentas = [r["ct"] for r in fetch_all(
        "SELECT IdContacto AS ct FROM dbo.vw_MovimientosCuenta_Base GROUP BY IdContacto "
        "HAVING SUM(CASE WHEN Origen IN ('Tarjetas','Tarjeta sin imputar') THEN 1 ELSE 0 END) > 0 "
        "AND SUM(CASE WHEN Origen IN ('Banco Nacion','Galicia','Pagos efectivo','Pagos Valores Recibidos') AND Credito > 0 THEN 1 ELSE 0 END) = 0 "
        "AND SUM(CASE WHEN Origen = 'Compras' AND Deuda > 0 THEN 1 ELSE 0 END) > 0", ())]
    lista = ",".join(map(str, cuentas))
    vinc = defaultdict(float)
    for r in fetch_all("SELECT IdCompra AS i, ImporteImputado AS x FROM dbo.Tarjetas_Resumenes_Lineas_Compras", ()):
        vinc[r["i"]] += float(r["x"] or 0)
    docs = fetch_all(
        "SELECT c.IdContacto AS ct, k.[Razon Social] AS r, c.IdDeuda AS i, c.Fecha AS f, c.[Nro Documento] AS n, c.[Tipo documento] AS td, "
        "c.Moneda AS m, v.ImporteDocumento AS v FROM dbo.Compras c JOIN dbo.vw_Compras_ImporteDocumento v ON v.IdDeuda = c.IdDeuda "
        f"JOIN dbo.Contactos k ON k.IdContacto = c.IdContacto WHERE c.IdContacto IN ({lista})", ())
    sueltas = fetch_all(f"SELECT IdContacto AS ct, Fecha AS f, Credito AS x, Origen AS o FROM dbo.vw_MovimientosCuenta_Base "
                        f"WHERE IdContacto IN ({lista}) AND Origen IN ('Tarjetas','Tarjeta sin imputar') AND Credito > 0", ())
    # compras particulares (regla 07/10/2026): la factura descuenta la parte personal y un asiento "Particular L/S" del
    # mismo día compensa lo que la tarjeta pagó de más
    particular = defaultdict(float)
    for r in fetch_all(f"SELECT IdContacto AS ct, Fecha AS f, Credito AS x FROM dbo.vw_MovimientosCuenta_Base WHERE IdContacto IN ({lista}) "
                       "AND Origen = 'Pagos efectivo' AND Credito < 0", ()):
        particular[(r["ct"], r["f"].date())] += -float(r["x"])
    filas, pct_por_cuenta = [], defaultdict(list)
    for d in docs:
        t = float(d["v"] or 0)
        if (d["td"] and "Cr" in d["td"]) or t <= 0:
            continue
        x = round(vinc.get(d["i"], 0.0), 2)
        if x > t and particular.get((d["ct"], d["f"].date())):
            x = round(x - particular[(d["ct"], d["f"].date())], 2)   # parte particular compensada con su asiento
        if x and abs(t - x) / t <= 0.001:
            continue
        pct = (t - x) / t * 100 if x else None
        if pct is not None:
            pct_por_cuenta[d["ct"]].append(round(pct, 2))
        filas.append((d, t, x, pct))
    salida = []
    for d, t, x, pct in filas:
        if d["m"] == "Dolares":
            causa = "Factura en US$: diferencia de cambio (regla TC del comprobante)"
        elif not x:
            iguales = [s for s in sueltas if s["ct"] == d["ct"] and abs(float(s["x"]) - t) <= max(0.01, t * 0.001)
                       and abs((s["f"] - d["f"]).days) <= 90]
            causa = (f"Sin vínculo: hay un consumo de tarjeta de {_pesos(float(iguales[0]['x']))} del {iguales[0]['f']:%d/%m/%Y} sin vincular"
                     if iguales else "Sin vínculo y sin consumo de tarjeta que coincida: ¿pagada por otro medio o falta el consumo?")
        else:
            repetido = Counter(pct_por_cuenta[d["ct"]])[round(pct, 2)] >= 2
            if t < 50 and x > 10 * t:
                causa = "Factura cargada sin importe (o casi): corregir el total con el comprobante"
            elif x > t and any(abs(x - k * t) <= 0.02 * k for k in (2, 3)):
                causa = f"Consumo vinculado {round(x / t)} veces a la misma factura: quitar el vínculo duplicado"
            elif x > t:
                causa = "Consumo de tarjeta mayor que la factura: falta otra factura del mismo consumo o parte fue compra particular"
            elif repetido:
                causa = f"Descuento repetido del {pct:.2f}% en este proveedor"
            elif abs(pct) <= 1 and float(x).is_integer():
                causa = "Redondeo en el pago (importe redondo)"
            elif abs(pct) <= 1:
                causa = "Entre 0,10% y 1% sin causa clara: revisar con el comprobante"
            else:
                causa = "Más del 1%: pago parcial, descuento o falta otro pago; revisar"
        salida.append([d["r"], d["n"], d["f"], _pesos(t), _pesos(x) if x else "(sin vínculo)",
                       f"{pct:.2f}%" if pct is not None else "", causa])
    salida.sort(key=lambda r: (r[0], r[2]))
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Grupo 2 - excepciones"
    hoja.append(["Proveedor", "Factura", "Fecha", "Total factura", "Cobrado por tarjeta", "Diferencia", "Causa probable"])
    for c in hoja[1]:
        c.font = Font(bold=True)
    for r in salida:
        hoja.append(r)
        hoja.cell(hoja.max_row, 3).number_format = "DD/MM/YYYY"
    for col, w in zip("ABCDEFG", (30, 18, 11, 15, 18, 10, 80)):
        hoja.column_dimensions[col].width = w
    for fila in hoja.iter_rows(min_row=2):
        fila[6].alignment = Alignment(wrap_text=True, vertical="top")
    libro.save(sys.argv[1])
    print(len(salida), "facturas")
    for k, v in Counter(r[6].split(":")[0].split(" del ")[0] for r in salida).most_common():
        print(f"  {v:3}  {k}")
