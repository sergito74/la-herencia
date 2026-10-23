"""Caja física USD: inspección por defecto; --aplicar respalda y carga historia una vez."""

import argparse
import hashlib
import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from src.db.connection import get_connection, _assert_target_is_wc
from src.features.vinculos.backup import backup_verificado

LIBRO = Path(
    r"C:\Users\Sergio\Documents\La Herencia\Administracion y gestion\Cuentas a pagar\Cajas Giamigli.xlsx"
)
CAJA = "GiamigliUSD"


def fuente():
    content = LIBRO.read_bytes()
    w = load_workbook(LIBRO, read_only=True, data_only=True)
    try:
        s = w["Caja Efectivo Dolares"]
        rows = []
        running = Decimal(0)
        for i, v in enumerate(s.iter_rows(min_row=2, max_col=9, values_only=True), 2):
            if not any(x is not None for x in v):
                continue
            if (
                not isinstance(v[0], datetime)
                or not isinstance(v[6], (int, float))
                or v[2] not in ("A", "B")
            ):
                raise ValueError(f"Fila fuente no verificable: {i}")
            amount = Decimal(str(v[6])).quantize(Decimal(".01"))
            running += amount
            if v[7] is None or running != Decimal(str(v[7])).quantize(Decimal(".01")):
                raise ValueError(f"Saldo fuente inconsistente: {i}")
            rows.append((i, v[0], str(v[1] or ""), str(v[3] or ""), amount, v[2]))
        if len(rows) != 6 or running != Decimal("360.00"):
            raise ValueError("La fuente cambió respecto del relevamiento; revisar antes de cargar.")
        return rows, hashlib.sha256(content).hexdigest()
    finally:
        w.close()


def saldos(c):
    return [
        list(r)
        for r in c.execute(
            "SELECT Caja,COUNT_BIG(*),SUM(Importe) FROM dbo.MovimientosCajaEfectivo WHERE Caja<>? GROUP BY Caja ORDER BY Caja",
            CAJA,
        ).fetchall()
    ]


def run(aplicar=False):
    _assert_target_is_wc()
    rows, sha = fuente()
    with get_connection() as conn:
        c = conn.cursor()
        if c.execute("SELECT DB_NAME()").fetchone()[0] != "WC":
            raise RuntimeError("Solo WC")
        before = saldos(c)
        constraints = {}
        for table in ("MovimientosCajaEfectivo", "ArqueosCajaEfectivo"):
            constraints[table] = [
                tuple(r)
                for r in c.execute(
                    "SELECT name,definition FROM sys.check_constraints WHERE parent_object_id=OBJECT_ID(?) AND CHARINDEX('[Caja]',definition)>0",
                    f"dbo.{table}",
                ).fetchall()
            ]
            if len(constraints[table]) != 1:
                raise RuntimeError(f"CHECK de caja inesperado: {table}")
    record = {
        "fuente": str(LIBRO),
        "sha256": sha,
        "filas": len(rows),
        "saldoUSD": "360.00",
        "saldosPesosAntes": before,
        "constraintsAntes": constraints,
        "estado": "inspeccion",
    }
    if not aplicar:
        return record
    record["backupVerificado"] = backup_verificado("caja_dolares_040")
    folder = Path(__file__).resolve().parents[2] / "backups" / "caja-dolares-040"
    folder.mkdir(parents=True, exist_ok=True)
    log = folder / f"migracion-{datetime.now():%Y%m%d-%H%M%S}.json"

    def save():
        log.write_text(
            json.dumps(record, default=str, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    record["estado"] = "respaldado"
    save()
    with get_connection(readonly=False) as conn:
        c = conn.cursor()
        try:
            c.execute("SET XACT_ABORT ON; BEGIN TRANSACTION")
            c.execute(
                "SELECT COUNT(*) FROM dbo.MovimientosCajaEfectivo WITH (TABLOCKX,HOLDLOCK)"
            ).fetchone()
            for table, items in constraints.items():
                name = items[0][0].replace("]", "]]")
                c.execute(f"ALTER TABLE dbo.[{table}] DROP CONSTRAINT [{name}]")
                c.execute(
                    f"ALTER TABLE dbo.[{table}] WITH CHECK ADD CONSTRAINT [CK_{table}_Caja040] CHECK (Caja IN ('GiamigliSA','CampoChica','CondominioLSC','GiamigliUSD'))"
                )
            inserted = 0
            for fila, fecha, concepto, detalle, importe, cuenta in rows:
                prior = c.execute(
                    "SELECT Fecha,Concepto,Detalle,Importe,Cuenta FROM dbo.MovimientosCajaEfectivo WHERE Caja=? AND FilaPlanilla=?",
                    CAJA,
                    fila,
                ).fetchall()
                if prior:
                    if (
                        len(prior) != 1
                        or str(prior[0][0])[:10] != fecha.date().isoformat()
                        or tuple(prior[0][1:]) != (concepto, detalle, importe, cuenta)
                    ):
                        raise ValueError(f"Fila histórica modificada: {fila}")
                    continue
                c.execute(
                    "INSERT INTO dbo.MovimientosCajaEfectivo (Caja,Fecha,Concepto,Detalle,Importe,Cuenta,Usuario,FilaPlanilla,Tipo) VALUES (?,?,?,?,?,?,?,?,?)",
                    CAJA,
                    fecha,
                    concepto,
                    detalle,
                    importe,
                    cuenta,
                    "Migracion040: Caja Efectivo Dolares",
                    fila,
                    "ingreso_dolares" if importe > 0 else "egreso_dolares",
                )
                inserted += 1
            after = saldos(c)
            if after != before:
                raise RuntimeError("Cambió una caja en pesos durante la migración")
            record.update(estado="aplicado", insertadas=inserted, saldosPesosDespues=after)
            c.execute("COMMIT TRANSACTION")
        except Exception as exc:
            c.execute("IF @@TRANCOUNT > 0 ROLLBACK TRANSACTION")
            record.update(estado="revertido", error=str(exc))
            save()
            raise
    save()
    record["informe"] = str(log)
    return record


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--aplicar", action="store_true")
    print(json.dumps(run(p.parse_args().aplicar), default=str, ensure_ascii=True))
