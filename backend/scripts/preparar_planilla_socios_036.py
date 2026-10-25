"""Lectura y plan de sincronización de las cuatro cuentas, sin editar Excel/WC."""

import json
import hashlib
import shutil
from bisect import bisect_right
from datetime import date, datetime
from decimal import Decimal as D
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
FOLDER = ROOT / "backups/socios-inml-5050-20261025-036"
EXCEL = Path("C:/Users/Sergio/Documents/La Herencia/Administracion y gestion/Cuentas a pagar/Cajas Giamigli.xlsx")
HOJAS = {1: "Cuenta Sergio", 2: "Cuenta Lucy", 3: "Cuenta Cond LSC", 4: "Cuenta Ceci"}


def run():
    data = json.loads((FOLDER / "sistema.json").read_text(encoding="utf-8"))
    w = openpyxl.load_workbook(EXCEL, data_only=True)
    plan = []
    for socio, hoja in HOJAS.items():
        regs = {r["IdMovimiento"]: r for r in data["movimientos"] if r["IdSocio"] == socio and not r["Anulada"]}
        usados, enlaces, sin = set(), [], []
        s = w[hoja]
        ultima = 2
        for row in s.iter_rows(min_row=3):
            f, debe, haber = row[0].value, row[4].value, row[7].value
            if not isinstance(f, (date, datetime)):
                continue
            ultima = max(ultima, row[0].row)
            # Después de la primera sincronización manda la identidad estable,
            # incluso para el mutuo sin importe nominal en pesos.
            id_previo = s.cell(row[0].row, 17).value
            if isinstance(id_previo, int) and id_previo in regs:
                if id_previo in usados:
                    raise ValueError(f"Identidad duplicada en {hoja}: {id_previo}")
                usados.add(id_previo)
                enlaces.append({"fila": row[0].row, "id": id_previo, "tipo": regs[id_previo]["Tipo"]})
                continue
            for tipo, importe in (("AsignacionGasto", debe), ("Devolucion", haber)):
                if not isinstance(importe, (int, float)) or abs(importe) < .0001:
                    continue
                candidates = [r for r in regs.values() if r["IdMovimiento"] not in usados and r["Fecha"][:10] == f.strftime("%Y-%m-%d")
                              and r["Tipo"] == tipo and abs(D(r["Importe"]) - D(str(importe))) <= D(".01")]
                # Samsung: diferencia nominal de $0,99 ya documentada en el traspaso.
                if not candidates and socio == 1 and abs(D(str(importe)) - D("220000")) < D(".01"):
                    candidates = [r for r in regs.values() if r["IdMovimiento"] == 4 and r["IdMovimiento"] not in usados]
                # Sancor octubre: fecha real de cobro 05/10 ya corregida en 036.
                if not candidates and socio == 2 and row[0].row == 506 and importe == 50259.78:
                    candidates = [r for r in regs.values() if r["IdMovimiento"] == 2879 and r["IdMovimiento"] not in usados]
                if not candidates:
                    sin.append({"fila": row[0].row, "fecha": f.isoformat(), "tipo": tipo, "importe": importe,
                                "proveedor": row[1].value, "detalle": row[2].value})
                    continue
                texto = " ".join(str(row[i].value or "") for i in (1, 2, 12)).lower()
                chosen = max(candidates, key=lambda r: SequenceMatcher(None, texto, (r["Motivo"] or "").lower()).ratio())
                usados.add(chosen["IdMovimiento"])
                enlaces.append({"fila": row[0].row, "id": chosen["IdMovimiento"], "tipo": tipo})
        plan.append({"socio": socio, "hoja": hoja, "ultimaFila": ultima, "enlaces": enlaces, "sinEnlace": sin,
                     "agregar": [r["IdMovimiento"] for r in regs.values() if r["IdMovimiento"] not in usados]})
    (FOLDER / "plan-excel.json").write_text(json.dumps(plan, default=str, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps([{**{k: v for k, v in p.items() if k not in ("enlaces", "agregar")},
                       "enlaces": len(p["enlaces"]), "agregar": len(p["agregar"])} for p in plan], ensure_ascii=False, indent=2))
    preparar_cambios(data, plan)


def preparar_cambios(data, plan):
    """Valores y fórmulas para el autor JS; openpyxl se usa solo para leer."""
    w = openpyxl.load_workbook(EXCEL)
    d = {r["Fecha"][:10]: float(r["Prom_billete"]) for r in data["dolares"]}
    n = {r["Fecha"][:10]: float(r["Indice"]) for r in data["indices"]}
    ds, ns = sorted(d), sorted(n)
    movimientos = {r["IdMovimiento"]: r for r in data["movimientos"]}
    cambios = []
    for p in plan:
        s = w[p["hoja"]]
        valores, formulas = {}, {}
        entradas = [(e["fila"], movimientos[e["id"]]) for e in p["enlaces"]]
        entradas += [(p["ultimaFila"] + i + 1, movimientos[id_]) for i, id_ in enumerate(p["agregar"])]
        ultima = max(f for f, _ in entradas)
        # Bloquear novedades documentales inesperadas, sin neutralizarlas por monto.
        for no in p["sinEnlace"]:
            if (p["socio"], no["fila"]) not in ((1, 346), (2, 302)):
                raise ValueError(f"Fila nueva sin resolver: {p['hoja']} {no}")
        filas = {f: r for f, r in entradas}
        for f, r in entradas:
            fecha = r["Fecha"][:10]
            fd, fn = ds[bisect_right(ds, fecha) - 1], ns[bisect_right(ns, fecha) - 1]
            agregado = f > p["ultimaFila"]
            if agregado:
                valores.update({f"B{f}": HOJAS[p["socio"]].removeprefix("Cuenta "), f"C{f}": r["Motivo"] or r["Origen"] or "Asiento sistema",
                                f"M{f}": r["Medio"] or "", f"D{f}": f"Asiento {r['IdMovimiento']}"})
            elif r["IdMovimiento"] in (4, 2879):
                original = s[f"C{f}"].value or ""
                valores[f"C{f}"] = str(original) + (" [nominal del sistema $219.999,01; planilla anterior $220.000]" if r["IdMovimiento"] == 4
                                                        else " [fecha de cobro corregida al 05/10/2026]")
            # Número de serie Excel real, sin cadenas de fecha en cálculos.
            valores[f"A{f}"] = (date.fromisoformat(fecha) - date(1899, 12, 30)).days
            debe = r["Tipo"] == "AsignacionGasto"
            valores.update({f"E{f}": float(r["Importe"]) if debe else 0, f"H{f}": 0 if debe else float(r["Importe"]),
                            f"O{f}": d[fd], f"P{f}": n[fn], f"Q{f}": r["IdMovimiento"],
                            f"R{f}": (date.fromisoformat(fd) - date(1899, 12, 30)).days,
                            f"S{f}": (date.fromisoformat(fn) - date(1899, 12, 30)).days,
                            f"T{f}": "Mutuo2018" if r["Origen"] == "Mutuo2018" else "50/50",
                            f"U{f}": "Vigente"})
            if r["Origen"] == "Mutuo2018":
                valores.update({f"F{f}": 0, f"I{f}": 0,
                                f"G{f}": float(r["ImporteUSD"]) if debe else 0,
                                f"J{f}": 0 if debe else float(r["ImporteUSD"])})
            else:
                for col, impcol, pricecol, dec in (("F", "E", "P", 3), ("I", "H", "P", 3), ("G", "E", "O", 4), ("J", "H", "O", 4)):
                    formulas[f"{col}{f}"] = f"=ROUND({impcol}{f}*50%/{pricecol}{f},{dec})"
        for f in range(3, ultima + 1):
            if f not in filas:
                # Conserva texto histórico; excluye solo los dos casos ya resueltos.
                if (p["socio"], f) in ((1, 346), (2, 302)):
                    imp = "$485.400; corresponde al Condominio" if p["socio"] == 1 else "Debe $40 / Haber $10; cambio neto $30 bajo umbral, no registrado"
                    valores.update({f"C{f}": f"EXCLUIDO del saldo: {imp}. Registro original en respaldo 25/10/2026.", f"U{f}": "Excluido documentado"})
                for col in ("E", "F", "G", "H", "I", "J"):
                    valores[f"{col}{f}"] = 0
            # Las filas vacías y el mutuo no deben cortar el saldo acumulado.
            prev = f - 1
            mutuo = f in filas and filas[f]["Origen"] == "Mutuo2018"
            formulas[f"K{f}"] = (f"=K{prev}" if f > 3 else "=0") if mutuo else (f"=K{prev}-F{f}+I{f}" if f > 3 else f"=-F{f}+I{f}")
            formulas[f"L{f}"] = (f"=L{prev}" if f > 3 else "=0") if mutuo else (f"=L{prev}-G{f}+J{f}" if f > 3 else f"=-G{f}+J{f}")
            formulas[f"N{f}"] = f"=N{prev}+H{f}-E{f}" if f > 3 else f"=H{f}-E{f}"
        # Q=identidad estable para futuras sincronizaciones; no modifica la base.
        valores.update({"W9": d[ds[-1]], "W10": n[ns[-1]], "V9": "BNA actual $/US$", "V10": "INMAG actual $/kg", "Q1": "Corte 25/10/2026",
                        "O2": "BNA $/US$", "P2": "INML/INMAG $/kg", "Q2": "Asiento sistema",
                        "R2": "Fecha BNA", "S2": "Fecha índice", "T2": "Valuación", "U2": "Estado",
                        "V4": "Fuente novillo", "W4": "https://www.mercadoagroganadero.com.ar/dll/hacienda2.dll/haciinfo000011",
                        "V5": "Fuente dólar", "W5": "WC: Dolar BNA.Prom_billete (última cotización disponible a la fecha)",
                        "V6": "Criterio", "W6": "50% USD y 50% kg en todas las fechas; INML Liniers continuado por INMAG",
                        "V7": "Cotización actual", "W7": f"BNA {ds[-1]}; INMAG {ns[-1]}"})
        formulas["C1"] = f"=K{ultima}*$W$10+L{ultima}*$W$9"
        if p["socio"] == 2:
            valores.update({"V1": "Mutuo2018 US$", "V2": "Mutuo2018 AR$"})
            formulas.update({"W1": f'=SUMIF(T3:T{ultima},"Mutuo2018",J3:J{ultima})-SUMIF(T3:T{ultima},"Mutuo2018",G3:G{ultima})',
                             "W2": "=W1*$W$9"})
        cambios.append({"hoja": p["hoja"], "valores": valores, "formulas": formulas,
                        "ultimaFila": ultima, "agregados": len(p["agregar"]), "enlaces": entradas})
    # No incluir datos duplicados en el archivo de comandos; ya están en sistema.json.
    for c in cambios:
        c["enlaces"] = [{"fila": f, "id": r["IdMovimiento"]} for f, r in c["enlaces"]]
    (FOLDER / "cambios-excel.json").write_text(json.dumps(cambios, ensure_ascii=False), encoding="utf-8")
    backup = FOLDER / "Cajas Giamigli-antes-5050.xlsx"
    sha = hashlib.sha256(EXCEL.read_bytes()).hexdigest()
    if backup.exists() and hashlib.sha256(backup.read_bytes()).hexdigest() != sha:
        backup = FOLDER / f"Cajas Giamigli-antes-5050-{sha[:12]}.xlsx"
    if not backup.exists():
        shutil.copy2(EXCEL, backup)
    if hashlib.sha256(backup.read_bytes()).hexdigest() != sha:
        raise ValueError("Falló verificación de respaldo Excel")
    (FOLDER / "respaldo-excel.json").write_text(json.dumps({"original": str(EXCEL), "backup": str(backup), "sha256": sha}), encoding="utf-8")


if __name__ == "__main__":
    run()
