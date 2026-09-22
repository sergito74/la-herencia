"""Regresión de columnas y formatos de las planillas, sin conexión SQL."""

from io import BytesIO

from openpyxl import load_workbook

from src.features.resultado_cultivo import exportacion


def test_exportacion_con_filas_y_formatos(monkeypatch):
    item = dict.fromkeys([
        "superficieSembrada", "superficieCosechada", "superficiePicada", "rinde",
        "costoTotalPesos", "costoTotalDolares", "costoPorHectareaSembradaPesos",
        "costoPorHectareaSembradaDolares", "costoPorHectareaCosechadaPesos",
        "costoPorHectareaCosechadaDolares", "ventaNetaPesos", "ventaNetaDolares",
        "margenBrutoPesos", "margenBrutoDolares", "rentabilidadPesos", "rentabilidadDolares",
    ], None)
    item.update(idCultivo=1, cultivo="Soja", campania="2026/2027", costoTotalPesos=100, rentabilidadPesos=0.5)
    monkeypatch.setattr(exportacion.resultado, "detalle_costos", lambda *_: [{
        "concepto": "Seguro", "rubro": "Seguros", "montoPesos": 100, "montoDolares": 10,
        "origen": "Seguro", "idCompra": None, "idDetalleCompra": None, "idOrdenTrabajo": None,
    }])
    wb = load_workbook(BytesIO(exportacion._crear_libro([item], 32)))
    assert wb.sheetnames == ["Resultado", "Detalle de costos"]
    assert wb["Resultado"].max_column == 19
    assert wb["Resultado"]["G2"].value == 100
    assert wb["Resultado"]["G2"].number_format == exportacion._MONEDA
    assert wb["Resultado"]["Q2"].value == 0.5
    assert wb["Resultado"]["Q2"].number_format == "0.00%"
    assert wb["Detalle de costos"]["E2"].number_format == exportacion._MONEDA
    assert wb["Detalle de costos"]["F2"].number_format == exportacion._DOLARES
