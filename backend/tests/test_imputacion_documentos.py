"""Tests del informe de imputación por documento comercial (017, pedido del
usuario 2026-09-25)."""

from __future__ import annotations

from src.features.imputacion import exportacion, repository


def _linea(id_detalle, fracciones=None):
    return {
        "idDetalleCompra": id_detalle,
        "producto": "Producto de prueba",
        "cantidad": 10.0,
        "unidad": "LTS",
        "precioUnitario": 5.0,
        "campaniaManual": "2025/2026",
        "centroCostoManual": "Agricultura",
        "rubroManual": "Herbicidas",
        "fracciones": fracciones or [],
    }


def test_informe_documentos_xlsx_genera_filas_agrupadas_y_totales(monkeypatch):
    documentos = [
        {
            "idCompra": 1,
            "fecha": "2026-01-01",
            "idContacto": 10,
            "proveedor": "Proveedor S.A.",
            "tipoDocumento": "Factura",
            "numeroDocumento": "0001-00000001",
            "moneda": "Pesos",
        }
    ]
    lineas = [
        _linea(
            100,
            fracciones=[
                {"cultivo": "Trigo", "campania": "2025/2026", "centroCosto": None, "esGanaderia": False, "importe": 123.45, "estado": "Aprobada"},
                {"cultivo": None, "campania": None, "centroCosto": None, "esGanaderia": None, "importe": 50.0, "estado": "Aprobada"},
            ],
        )
    ]
    monkeypatch.setattr(repository, "listar_documentos_con_imputacion", lambda *a, **k: (documentos, 1))
    monkeypatch.setattr(repository, "lineas_con_imputacion_batch", lambda ids: {1: lineas})

    contenido = exportacion.informe_documentos_xlsx(None, None, None)

    assert contenido[:2] == b"PK"  # firma de un .xlsx (zip) válido

    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(contenido))
    ws = wb.active
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    # línea + 2 fracciones + total del documento + total general = 5 filas
    assert len(filas) == 5
    assert filas[3][6] == "Total Factura 0001-00000001"
    assert filas[3][7] == 173.45
    assert filas[4][6] == "TOTAL GENERAL"
    assert filas[4][7] == 173.45
    # las 2 filas de fracciones quedan plegadas (outline_level=1)
    assert ws.row_dimensions[3].outline_level == 1
    assert ws.row_dimensions[4].outline_level == 1


def test_etiqueta_destino_prioriza_cultivo_luego_centro_luego_stock():
    assert exportacion._etiqueta_destino({"cultivo": "Maiz", "campania": "2026/2027"}) == "Maiz / 2026/2027"
    assert exportacion._etiqueta_destino({"cultivo": None, "centroCosto": "Adm. General"}) == "Adm. General"
    assert exportacion._etiqueta_destino({"cultivo": None, "centroCosto": None, "esGanaderia": True}) == "Ganadería"
    assert exportacion._etiqueta_destino({"cultivo": None, "centroCosto": None, "esGanaderia": None}) == "En stock sin consumir"


def test_lineas_con_imputacion_batch_una_sola_query_por_tabla(monkeypatch):
    """Regresión N+1 (hallazgo de revisión SQL Server, 2026-09-25): traer las
    líneas de varios documentos debe ser 2 queries, no 2×N."""
    llamadas = []

    def fake_fetch_all(sql, params=()):
        llamadas.append(sql)
        if "FROM dbo.Det_Compras dc" in sql:
            return [
                {"idCompra": 1, "idDetalleCompra": 100, "producto": "A", "cantidad": 1.0, "unidad": "LTS",
                 "precioUnitario": 1.0, "campaniaManual": None, "centroCostoManual": None, "rubroManual": None},
                {"idCompra": 2, "idDetalleCompra": 200, "producto": "B", "cantidad": 1.0, "unidad": "LTS",
                 "precioUnitario": 1.0, "campaniaManual": None, "centroCostoManual": None, "rubroManual": None},
            ]
        return []

    monkeypatch.setattr(repository, "fetch_all", fake_fetch_all)

    resultado = repository.lineas_con_imputacion_batch((1, 2))

    assert len(llamadas) == 2  # 1 para líneas, 1 para fracciones — sin importar cuántos documentos
    assert set(resultado.keys()) == {1, 2}


def test_informe_documentos_xlsx_soporta_importes_decimal(monkeypatch):
    """Regresión real (WC, 2026-09-25): pyodbc devuelve `Importe` como
    `decimal.Decimal`, no `float` — sumarlo sin convertir contra el
    acumulador float rompía la exportación con TypeError."""
    from decimal import Decimal

    documentos = [
        {
            "idCompra": 1,
            "fecha": "2026-01-01",
            "idContacto": 10,
            "proveedor": "Proveedor S.A.",
            "tipoDocumento": "Factura",
            "numeroDocumento": "0001-00000001",
            "moneda": "Pesos",
        }
    ]
    lineas = [
        _linea(
            100,
            fracciones=[
                {"cultivo": "Trigo", "campania": "2025/2026", "centroCosto": None, "esGanaderia": False, "importe": Decimal("123.45"), "estado": "Aprobada"},
            ],
        )
    ]
    monkeypatch.setattr(repository, "listar_documentos_con_imputacion", lambda *a, **k: (documentos, 1))
    monkeypatch.setattr(repository, "lineas_con_imputacion_batch", lambda ids: {1: lineas})

    contenido = exportacion.informe_documentos_xlsx(None, None, None)

    assert contenido[:2] == b"PK"
