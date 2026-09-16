"""Contract tests for POST /api/tesoreria/excel/validar (FR-007/008/009).

The endpoint wiring is tested with fixtures (mocking `excel_import`), same
approach as the rest of this suite. `excel_import`'s actual parsing logic
(header detection, Argentine amount parsing, Galicia column mapping) is
covered separately by unit tests below, since generating a real legacy
`.xls` binary requires `xlwt` (not part of this project's dependencies —
only the read-only `xlrd` is, per research.md), while `.xlsx` can be built
directly with `openpyxl`.
"""

from __future__ import annotations

import io

import httpx
import openpyxl
import pytest

from src.features.tesoreria import excel_import
from src.main import app


@pytest.fixture
def client():
    return httpx.ASGITransport(app=app)


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.mark.anyio
async def test_validar_excel_bna_valido(client, monkeypatch):
    fixture = {
        "medioDetectado": "bna",
        "valido": True,
        "errores": [],
        "movimientosPrevisualizados": [
            {
                "fecha": "2026-01-19",
                "comprobante": "81415",
                "concepto": "PLAZO FIJO",
                "importe": 5107397.26,
                "saldo": 3230108.83,
            }
        ],
    }
    monkeypatch.setattr(excel_import, "validar_y_previsualizar", lambda name, content: fixture)

    async with httpx.AsyncClient(transport=client, base_url="http://test") as ac:
        response = await ac.post(
            "/api/tesoreria/excel/validar",
            files={"archivo": ("resumen.xls", b"contenido", "application/vnd.ms-excel")},
        )

    assert response.status_code == 200
    body = response.json()
    assert body["medioDetectado"] == "bna"
    assert body["valido"] is True


@pytest.mark.anyio
async def test_validar_excel_invalido(client, monkeypatch):
    fixture = {
        "medioDetectado": None,
        "valido": False,
        "errores": ["No se encontró la fila de encabezado esperada"],
        "movimientosPrevisualizados": [],
    }
    monkeypatch.setattr(excel_import, "validar_y_previsualizar", lambda name, content: fixture)

    async with httpx.AsyncClient(transport=client, base_url="http://test") as ac:
        response = await ac.post(
            "/api/tesoreria/excel/validar",
            files={"archivo": ("desconocido.xls", b"contenido", "application/vnd.ms-excel")},
        )

    # Invalid structure is reported in the body, not as an HTTP error (FR-009).
    assert response.status_code == 200
    body = response.json()
    assert body["valido"] is False
    assert body["movimientosPrevisualizados"] == []


@pytest.mark.anyio
async def test_validar_excel_rejects_get(client):
    async with httpx.AsyncClient(transport=client, base_url="http://test") as ac:
        response = await ac.get("/api/tesoreria/excel/validar")
    assert response.status_code in (404, 405)


def test_parse_importe_arg_formato_argentino():
    assert excel_import._parse_importe_arg("$ 5.107.397,26") == 5107397.26
    assert excel_import._parse_importe_arg("$ -9,00") == -9.00
    assert excel_import._parse_importe_arg(None) is None


def test_procesar_galicia_real_shape():
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("Movimientos")
    sheet.append(
        [
            "Fecha",
            "Descripción",
            "Origen",
            "Débitos",
            "Créditos",
            "Grupo de Conceptos",
            "Concepto",
            "Número de Terminal",
            "Observaciones Cliente",
            "Número de Comprobante",
            "Leyendas Adicionales 1",
            "Leyendas Adicionales 2",
            "Leyendas Adicionales 3",
            "Leyendas Adicionales 4",
            "Tipo de Movimiento",
            "Saldo",
        ]
    )
    from datetime import datetime

    sheet.append(
        [
            datetime(2026, 8, 20),
            "Trf Inmed Proveed",
            "Home Banking",
            46044.04,
            0,
            "Transferencias",
            "Transferencia",
            None,
            None,
            "57808209",
            "Jauregui Y Morales",
            "30545419110",
            "FACTURAS",
            None,
            "Débito",
            774943.40,
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)

    resultado = excel_import.validar_y_previsualizar("resumen.xlsx", buffer.getvalue())

    assert resultado["medioDetectado"] == "galicia"
    assert resultado["valido"] is True
    mov = resultado["movimientosPrevisualizados"][0]
    assert mov["descripcion"] == "Trf Inmed Proveed"
    assert mov["debitos"] == 46044.04
    assert mov["leyendas"] == ["Jauregui Y Morales", "30545419110", "FACTURAS", None]


def test_procesar_galicia_encabezado_invalido():
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("Movimientos")
    sheet.append(["Columna Rara", "Otra"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    resultado = excel_import.validar_y_previsualizar("resumen.xlsx", buffer.getvalue())

    assert resultado["valido"] is False
    assert resultado["medioDetectado"] is None
    assert resultado["movimientosPrevisualizados"] == []
