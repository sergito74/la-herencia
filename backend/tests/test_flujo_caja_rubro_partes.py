"""030 — flujo de caja por rubro: reparto por lo aplicado, traspasos entre
cuentas propias, cotización del día, agregación, saldos y exportación.
Funciones puras o con monkeypatch; nunca toca WC."""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from src.features.flujo_caja import atribucion, clasificacion, cotizacion, exportacion, repository


def ap(tipo, id_doc, importe):
    return {"tipoDocumento": tipo, "idDocumentoAplicado": id_doc, "importeAplicado": importe}


@pytest.fixture
def ventas(monkeypatch):
    monkeypatch.setattr(atribucion, "_rubro_de_venta", lambda tipo, i: {1: "Venta Soja", 2: "Venta Terneros"}[i])


# --- Reparto por importe aplicado (T002) ---

def test_reparte_por_lo_aplicado_entre_dos_rubros(ventas):
    partes = atribucion.partes_desde_aplicaciones([ap("VentaGranos", 1, 700), ap("VentaHacienda", 2, 300)], 1000, date(2026, 5, 1))
    assert [(p["rubro"], p["importe"]) for p in partes] == [("Venta Soja", 700), ("Venta Terneros", 300)]


def test_remanente_va_a_pendiente_y_antes_del_corte_a_historico(ventas):
    partes = atribucion.partes_desde_aplicaciones([ap("VentaGranos", 1, 600)], 1000, date(2026, 5, 1))
    assert partes[-1] == {"rubro": "Pendiente de aplicar", "centroCosto": None, "importe": 400, "documentoAplicado": None}
    viejo = atribucion.partes_desde_aplicaciones([ap("VentaGranos", 1, 600)], 1000, date(2015, 8, 31))
    assert viejo[-1]["rubro"] == "Histórico sin aplicar"


def test_compra_con_renglones_de_dos_rubros_se_reparte_por_importe(monkeypatch):
    monkeypatch.setattr(atribucion, "_rubros_de_compra", lambda i: [
        {"rubro": "Semillas", "centroCosto": "Agricultura", "peso": 750},
        {"rubro": "Combustible", "centroCosto": "Maquinaria", "peso": 250},
    ])
    partes = atribucion.partes_desde_aplicaciones([ap("CompraDeuda", 9, 1000)], -1000, date(2026, 5, 1))
    assert sorted((p["rubro"], p["centroCosto"], p["importe"]) for p in partes) == [
        ("Combustible", "Maquinaria", 250), ("Semillas", "Agricultura", 750)]


def test_suma_de_partes_es_exacta_y_sobreaplicado_se_escala(ventas):
    partes = atribucion.partes_desde_aplicaciones([ap("VentaGranos", 1, 333.333), ap("VentaHacienda", 2, 333.333),
                                                   ap("VentaGranos", 1, 333.337)], 1000, date(2026, 5, 1))
    assert round(sum(p["importe"] for p in partes), 2) == 1000
    escalado = atribucion.partes_desde_aplicaciones([ap("VentaGranos", 1, 1500)], 1000, date(2026, 5, 1))
    assert [p["importe"] for p in escalado] == [1000]


def test_sin_aplicaciones_devuelve_none():
    assert atribucion.partes_desde_aplicaciones([], 100, date(2026, 5, 1)) is None


# --- Traspasos entre cuentas propias (T004) ---

def test_tipo_interno():
    assert clasificacion.tipo_interno("Galicia", -100, None, "Inversiones") == "Colocación FIMA"
    assert clasificacion.tipo_interno("Galicia", 100, None, "Inversiones") == "Rescate FIMA"
    assert clasificacion.tipo_interno("BNA", -100, "TRANSF MIS TIT 30712114602") == "Traspaso entre bancos"
    assert clasificacion.tipo_interno("Galicia", -100, "Pago proveedor", "Proveedores") is None


def mov(banco, fecha, importe, tipo=None):
    return {"banco": banco, "fecha": fecha, "importe": importe, "esInterno": tipo is not None, "tipoInterno": tipo}


def test_empareja_lado_galicia_a_dos_dias():
    ms = [mov("BNA", date(2026, 3, 1), -500, "Traspaso entre bancos"), mov("Galicia", date(2026, 3, 3), 500)]
    clasificacion.emparejar_traspasos(ms)
    assert ms[1]["esInterno"] and ms[1]["tipoInterno"] == "Traspaso entre bancos"
    assert not ms[0].get("sinContraparte")


def test_sin_pareja_a_cuatro_dias_queda_sin_contraparte():
    ms = [mov("BNA", date(2026, 3, 1), -500, "Traspaso entre bancos"), mov("Galicia", date(2026, 3, 5), 500)]
    clasificacion.emparejar_traspasos(ms)
    assert ms[0]["sinContraparte"] and not ms[1]["esInterno"]


def test_dos_candidatos_se_empareja_solo_el_mas_cercano():
    ms = [mov("BNA", date(2026, 3, 1), -500, "Traspaso entre bancos"),
          mov("Galicia", date(2026, 3, 3), 500), mov("Galicia", date(2026, 3, 2), 500)]
    clasificacion.emparejar_traspasos(ms)
    assert [m["esInterno"] for m in ms[1:]] == [False, True]


# --- Cotización del día (T005) ---

def test_cotizacion_dia_exacto_fallback_y_ausencia():
    serie = {date(2026, 3, 2): 1000.0}
    assert cotizacion.cotizacion_del_dia(serie, date(2026, 3, 2)) == (1000.0, date(2026, 3, 2))
    assert cotizacion.cotizacion_del_dia(serie, date(2026, 3, 5)) == (1000.0, date(2026, 3, 2))
    assert cotizacion.cotizacion_del_dia(serie, date(2026, 3, 10)) is None


# --- Agregación, granularidad y USD (T016) ---

def parte(fecha, seccion, rubro, importe, centro=None, cuenta="Galicia CC", esFima=False):
    return {"fecha": fecha, "seccion": seccion, "rubro": rubro, "centroCosto": centro if seccion == "egresos" else None,
            "importeArs": importe, "cuenta": cuenta, "esFima": esFima}


PARTES = [
    parte(date(2026, 1, 10), "ingresos", "Venta Soja", 1000),
    parte(date(2026, 2, 20), "egresos", "Sueldos", -300, "Personal"),
    parte(date(2026, 3, 5), "egresos", "Pendiente de aplicar", -50, "Sin centro de costos"),
    parte(date(2026, 3, 6), "internos", "Colocación FIMA", -200, esFima=True),
]


def test_total_del_rango_igual_en_todas_las_granularidades():
    totales = set()
    for g in ("semanal", "mensual", "trimestral", "anual"):
        per = repository.periodos_del_rango(date(2026, 1, 1), date(2026, 3, 31), g)
        a = repository.agregar_por_rubro(PARTES, g, per)
        totales.add(round(sum(a["netoOperativoPorPeriodo"].values()) + sum(a["internos"]["totalPorPeriodo"].values()), 2))
    assert totales == {450.0}


def test_internos_siempre_presentes_y_fuera_del_neto_operativo():
    per = repository.periodos_del_rango(date(2026, 1, 1), date(2026, 3, 31), "mensual")
    a = repository.agregar_por_rubro(PARTES, "mensual", per)
    assert [r["rubro"] for r in a["internos"]["rubros"]] == ["Colocación FIMA", "Rescate FIMA", "Traspaso entre bancos"]
    assert a["netoOperativoPorPeriodo"] == {"2026-01": 1000, "2026-02": -300, "2026-03": -50}
    assert per == ["2026-01", "2026-02", "2026-03"]


def test_usd_convierte_por_dia_y_reporta_lo_que_no_tiene_cotizacion():
    partes = [dict(p) for p in PARTES]
    repository.convertir_partes_a_usd(partes, {date(2026, 1, 10): 1000.0, date(2026, 2, 20): 1200.0})
    per = repository.periodos_del_rango(date(2026, 1, 1), date(2026, 3, 31), "mensual")
    a = repository.agregar_por_rubro(partes, "mensual", per, "USD")
    assert a["ingresos"]["rubros"][0]["valores"]["2026-01"] == 1.0
    assert a["egresos"]["totalPorPeriodo"]["2026-02"] == -0.25
    assert {(s["periodo"], s["rubro"]) for s in a["sinTipoCambio"]} == {("2026-03", "Pendiente de aplicar"), ("2026-03", "Colocación FIMA")}


def test_colocacion_fima_pasa_de_galicia_al_fima_sin_cambiar_el_total():
    per = ["2026-03"]
    finales = repository.saldos_finales({"Nación": 0, "Galicia CC": 1000, "Galicia Fondo FIMA": 0},
                                        [PARTES[3]], "mensual", per)
    assert finales["Galicia CC"]["2026-03"] == 800 and finales["Galicia Fondo FIMA"]["2026-03"] == 200


# --- Exportación (T022) ---

def test_excel_tiene_importes_numericos_y_las_filas_de_la_vista(monkeypatch):
    per = ["2026-01", "2026-02"]
    datos = {
        "moneda": "ARS", "periodos": per,
        "saldoInicial": {"cuentas": [{"cuenta": "Nación", "importe": 10.0}], "total": 10.0},
        "ingresos": {"rubros": [{"rubro": "Venta Soja", "valores": {"2026-01": 1000.0, "2026-02": 0.0}, "total": 1000.0}],
                     "totalPorPeriodo": {"2026-01": 1000.0, "2026-02": 0.0}},
        "egresos": {"centrosCosto": [], "totalPorPeriodo": {"2026-01": 0.0, "2026-02": 0.0}},
        "netoOperativoPorPeriodo": {"2026-01": 1000.0, "2026-02": 0.0},
        "internos": {"rubros": [], "totalPorPeriodo": {}},
        "saldoFinalPorPeriodo": {"2026-01": 1010.0, "2026-02": 1010.0},
        "saldoFinalPorCuenta": {"Nación": {"2026-01": 1010.0, "2026-02": 1010.0}},
        "sinTipoCambio": [],
    }
    monkeypatch.setattr(repository, "flujo_por_rubro", lambda *a: datos)
    wb = load_workbook(BytesIO(exportacion.flujo_por_rubro_xlsx(date(2026, 1, 1), date(2026, 2, 28), "mensual", "ARS")))
    filas = {r[0]: r[1:] for r in wb.active.iter_rows(values_only=True) if r and r[0]}
    assert filas["Venta Soja"][:3] == (1000.0, 0.0, 1000.0)
    assert isinstance(filas["Total saldo final"][0], (int, float))


# --- 031: la fuente unificada alimenta el reparto (T007) ---

def test_debito_de_resumen_se_reparte_por_rubros_de_sus_compras(monkeypatch):
    from src.features.vinculos import cadenas

    monkeypatch.setattr(atribucion, "_rubros_de_compra", lambda i: [{"rubro": {10: "Semillas", 11: "Combustible"}[i],
                                                                     "centroCosto": "Agricultura", "peso": 1}])
    vinculos = cadenas.construir_vinculos({
        "aplicaciones": [], "tesoreria": [], "backfill": [], "valores": [], "movimientos": {}, "documentos": {},
        "lineasCompras": [{"idLinea": 1, "idCompra": 10, "importe": 300}, {"idLinea": 2, "idCompra": 11, "importe": 100}],
        "lineas": {1: {"idResumen": 7}, 2: {"idResumen": 7}},
        "pagosResumen": [{"idResumen": 7, "importe": 500, "origen": "bna", "idMovimiento": 99}]})
    aplic = cadenas.documentos_de_movimiento(vinculos)[("bna", 99)]
    partes = atribucion.partes_desde_aplicaciones(aplic, -500, date(2026, 5, 1))
    assert sorted((p["rubro"], p["importe"]) for p in partes) == [("Combustible", 100), ("Pendiente de aplicar", 100), ("Semillas", 300)]
    assert round(sum(p["importe"] for p in partes), 2) == 500
    assert {p["documentoAplicado"]["via"] for p in partes if p["documentoAplicado"]} == {"tarjeta"}


def test_impuesto_y_sueldo_caen_en_su_rubro(monkeypatch):
    monkeypatch.setattr(atribucion, "_rubro_de_impuesto", lambda i, memo: "Ingresos Brutos")
    partes = atribucion.partes_desde_aplicaciones(
        [{"tipoDocumento": "Impuesto", "idDocumentoAplicado": 1, "importeAplicado": 60, "via": "tesoreria"},
         {"tipoDocumento": "Remuneracion", "idDocumentoAplicado": 2, "importeAplicado": 40, "via": "tesoreria"}], -100, date(2026, 5, 1))
    assert [(p["rubro"], p["centroCosto"], p["importe"]) for p in partes] == [
        ("Ingresos Brutos", "Impuestos", 60), ("Sueldos", "Personal", 40)]
