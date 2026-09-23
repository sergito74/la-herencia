import io

import openpyxl
from fastapi.testclient import TestClient

from src.auth.tokens import crear_token

from src.features.tesoreria import exportacion, repository
from src.main import app

client = TestClient(app)
client.cookies.set(
    "la_herencia_session",
    crear_token(id_usuario=0, rol="Administrador"),
)  # 016-autenticacion: la API ahora exige sesión; estos tests preexistentes
# simulan un Administrador para no cambiar su comportamiento (SC-003).


def test_valores_propios_xlsx_incluye_comentarios(monkeypatch):
    monkeypatch.setattr(
        repository,
        "get_movimientos",
        lambda *a, **k: (
            [
                {
                    "idValor": 1, "numeroCheque": 123.0, "fechaEmision": None, "fechaVencimiento": None,
                    "importe": 1000.0, "cobrado": "S", "fechaCobro": None, "numeroCuenta": "16400292/80",
                    "comentarios": "Proveedor X",
                }
            ],
            1,
        ),
    )
    contenido = exportacion.valores_propios_xlsx(None, None)
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    ws = wb["Valores propios"]
    assert ws["A1"].value == "N° Cheque"
    assert ws["H1"].value == "Comentarios"
    assert ws["H2"].value == "Proveedor X"


def test_endpoint_exportar_valores_propios(monkeypatch):
    monkeypatch.setattr(exportacion, "valores_propios_xlsx", lambda *a, **k: b"contenido")
    r = client.get("/api/tesoreria/valores-propios/exportar")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
