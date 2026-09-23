import io

import openpyxl
from fastapi.testclient import TestClient

from src.auth.tokens import crear_token

from src.features.cuentas_corrientes import exportacion, repository
from src.main import app

client = TestClient(app)
client.cookies.set(
    "la_herencia_session",
    crear_token(id_usuario=0, rol="Administrador"),
)  # 016-autenticacion: la API ahora exige sesión; estos tests preexistentes
# simulan un Administrador para no cambiar su comportamiento (SC-003).


def test_get_saldos_todos_ordena_por_razon_social_por_defecto(monkeypatch):
    filas = []

    def fake_fetch_all(sql, params=()):
        filas.append(sql)
        return [{"idContacto": 1, "razonSocial": "B", "saldoParcial": 0.0}]

    monkeypatch.setattr(repository, "fetch_all", fake_fetch_all)
    resultado = repository.get_saldos_todos("razonSocial")
    assert resultado == [{"idContacto": 1, "razonSocial": "B", "saldoParcial": 0.0}]
    assert "ROW_NUMBER()" in filas[0]
    assert "[Razon Social]" in filas[0].split("ORDER BY")[-1]


def test_get_saldos_todos_orden_por_saldo_pone_mayor_deuda_primero(monkeypatch):
    capturado = {}

    def fake_fetch_all(sql, params=()):
        capturado["sql"] = sql
        return []

    monkeypatch.setattr(repository, "fetch_all", fake_fetch_all)
    repository.get_saldos_todos("saldo")
    assert "SaldoParcial ASC" in capturado["sql"].split("ORDER BY")[-1]


def test_endpoint_saldos_incluye_saldo_cero(monkeypatch):
    monkeypatch.setattr(
        repository,
        "get_saldos_todos",
        lambda orden: [
            {"idContacto": 1, "razonSocial": "Con deuda", "saldoParcial": -100.0},
            {"idContacto": 2, "razonSocial": "Saldado", "saldoParcial": 0.0},
        ],
    )
    r = client.get("/api/cuentas-corrientes/saldos")
    assert r.status_code == 200
    items = r.json()["items"]
    assert len(items) == 2
    assert any(i["saldoParcial"] == 0.0 for i in items)


def test_saldos_xlsx_genera_un_archivo_valido(monkeypatch):
    monkeypatch.setattr(
        repository,
        "get_saldos_todos",
        lambda orden: [
            {"idContacto": 1, "razonSocial": "Proveedor A", "saldoParcial": -1234.56},
            {"idContacto": 2, "razonSocial": "Proveedor B", "saldoParcial": 0.0},
        ],
    )
    contenido = exportacion.saldos_xlsx("razonSocial")
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    ws = wb["Saldos"]
    assert ws["A1"].value == "Razón Social"
    assert ws["A2"].value == "Proveedor A"
    assert ws["B2"].value == -1234.56


def test_cuenta_corriente_xlsx_sin_movimientos_no_falla(monkeypatch):
    monkeypatch.setattr(repository, "get_movimientos", lambda *a, **k: ([], 0))
    monkeypatch.setattr(repository, "get_saldo", lambda id_contacto: None)
    contenido = exportacion.cuenta_corriente_xlsx(999, None, None)
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    ws = wb["Movimientos"]
    assert ws["A1"].value == "Fecha"
    assert ws.max_row == 2  # solo encabezado + fila de saldo


def test_exportar_endpoints_devuelven_xlsx(monkeypatch):
    monkeypatch.setattr(exportacion, "saldos_xlsx", lambda orden: b"contenido")
    r = client.get("/api/cuentas-corrientes/saldos/exportar")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")

    monkeypatch.setattr(exportacion, "cuenta_corriente_xlsx", lambda *a, **k: b"contenido")
    r2 = client.get("/api/cuentas-corrientes/contactos/1/exportar")
    assert r2.status_code == 200
