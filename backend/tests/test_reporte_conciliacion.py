import io
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.features.tarjetas_resumenes import reporte_conciliacion as rc
from src.main import app

client = TestClient(app)


def _datos():
    cargos = {clave: 0 for _, clave, _ in rc.CARGOS}
    resumen = {"idResumen": 1, "tarjeta": "Visa", "codigo": "R1", "fechaCierre": date(2021, 12, 2), "fechaVencimiento": date(2021, 12, 13), **cargos, "impuestoSellos": 50}
    base = {"proveedor": "Proveedor SA", "cuit": "30111111112", "estado": None, "motivo": None, "detalleMotivo": None, "importeDiferencia": None}
    lineas = [
        {**base, "idLinea": 10, "idResumen": 1, "fecha": date(2021, 11, 1), "detalle": "Consumo A", "importe": 1000},
        {**base, "idLinea": 11, "idResumen": 1, "fecha": date(2021, 11, 2), "detalle": "Consumo B", "importe": 500, "estado": "SinDocumento", "motivo": "Impuesto"},
        {**base, "idLinea": 12, "idResumen": 1, "fecha": date(2021, 11, 3), "detalle": "Consumo C", "importe": 300},
        {**base, "idLinea": 13, "idResumen": 1, "fecha": date(2021, 11, 4), "detalle": "Consumo D", "importe": 200, "estado": "DiferenciaAceptada", "motivo": "Redondeo", "importeDiferencia": 2},
    ]
    doc = {"idCompra": 1, "fecha": date(2021, 11, 1), "tipo": "Factura", "numero": "0001-1", "moneda": "Dolares", "tipoDeCambio": 1200,
           "importeOriginal": 0.5, "compraParticular": 0, "ajustaTipoCambio": 0, "neto": 0.4, "iva": 0.1, "otros": 0,
           "proveedor": "Proveedor SA", "cuit": "30111111112"}
    nota = {**doc, "idCompra": 2, "tipo": "Nota de Débito", "numero": "0001-2", "moneda": "Pesos", "tipoDeCambio": 1, "importeOriginal": 400,
            "ajustaTipoCambio": 1, "neto": 330.58, "iva": 69.42}
    vinculos = [
        {**doc, "idLinea": 10, "imputado": 600},
        {**nota, "idLinea": 10, "imputado": 400},
        {**doc, "idLinea": 13, "imputado": 198},
    ]
    pagos = [{"idResumen": 1, "fecha": date(2021, 12, 13), "importe": 2050, "origen": "Galicia"}]
    return {"resumenes": [resumen], "lineas": lineas, "vinculos": vinculos, "pagos": pagos, "tarjeta": "Visa"}


def _libro():
    return rc.construir_libro(_datos(), {"tarjeta": "Visa", "desde": date(2021, 12, 1), "hasta": None})


def test_hojas_y_resumen_con_estados():
    wb = _libro()
    assert wb.sheetnames == ["Resúmenes", "Conciliación", "Pagos", "Ayuda"]  # todos tienen CUIT: sin hoja extra
    ws = wb["Resúmenes"]
    fila = [c.value for c in ws[2]]
    assert fila[4] == 2000  # total consumos
    assert fila[5] == 50  # impuesto de sellos
    titulos = [c.value for c in ws[1]]
    i = titulos.index("Conciliados")
    # 1 conciliada, 1 con diferencia, 1 sin documento, 1 pendiente
    assert fila[i : i + 5] == [1, 1, 1, 1] or fila[i : i + 5][:4] == [1, 1, 1, 1]
    assert fila[-1] == "Pendiente de conciliar"
    assert "SUM" in str(ws["E3"].value)  # fila de totales con fórmulas


def test_detalle_una_fila_por_documento_y_estados_por_linea():
    ws = _libro()["Conciliación"]
    filas = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    por_linea = {}
    for f in filas:
        por_linea.setdefault(f[3], []).append(f)
    assert len(por_linea[10]) == 2  # factura + nota de ajuste
    assert [f[6] for f in por_linea[10]] == [1000, None]  # el importe del consumo solo en la primera fila
    assert por_linea[10][0][7] == "Conciliada"
    assert por_linea[10][1][24] == "Nota de ajuste de tipo de cambio"
    assert por_linea[11][0][7] == "Sin documento / no aplica" and por_linea[11][0][8] == "Impuesto"
    assert por_linea[11][0][13] is None  # sin documento
    assert por_linea[12][0][7] == "Pendiente"
    assert por_linea[13][0][7] == "Conciliada con diferencia aceptada" and por_linea[13][0][10] == 2


def test_documento_en_dolares_se_pesifica_y_lleva_su_formato():
    ws = _libro()["Conciliación"]
    fila = next(r for r in ws.iter_rows(min_row=2) if r[3].value == 10)
    assert fila[16].value == "Dólares" and fila[17].value == 1200
    assert fila[22].value == 600  # 0,50 us$ x 1200
    assert "us$" in fila[21].number_format and "$" in fila[22].number_format


def test_pagos_y_ayuda():
    wb = _libro()
    fila = [c.value for c in wb["Pagos"][2]]
    assert fila[4] == 2050 and fila[5] == "Banco Galicia"
    textos = " ".join(str(c.value) for r in wb["Ayuda"].iter_rows() for c in r if c.value)
    assert "Visa" in textos and "01/12/2021" in textos


def test_endpoint_descarga_un_xlsx_valido(monkeypatch):
    monkeypatch.setattr(rc, "obtener_datos", lambda *a, **k: _datos())
    r = client.get("/api/tarjetas-resumenes/reporte-conciliacion?idTarjeta=4&fechaCierreDesde=2021-12-01")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in r.headers["content-disposition"] and ".xlsx" in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content)).sheetnames[0] == "Resúmenes"


def test_hoja_de_proveedores_sin_cuit():
    datos = _datos()
    for v in datos["vinculos"]:
        v["cuit"] = None
    wb = rc.construir_libro(datos, {"tarjeta": None, "desde": None, "hasta": None})
    assert "Proveedores sin CUIT" in wb.sheetnames
    fila = [c.value for c in wb["Proveedores sin CUIT"][2]]
    assert fila == ["Proveedor SA", 3]
    textos = " ".join(str(c.value) for r in wb["Ayuda"].iter_rows() for c in r if c.value)
    assert "no tienen CUIT" in textos
